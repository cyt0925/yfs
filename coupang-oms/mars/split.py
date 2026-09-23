"""① 拆單 ＋ ② 回填 EIP 採購單號／約倉時間。

拆單條件（Alice 規格第 1 步）：同一張酷澎 PO、同一個到貨日、同一個到貨倉之下，再依
  品類（商品總表 B：Chocolate→CHO、Gum→GUM、Petcare→PET）
  單位（整合表 O：箱／盒／包；SPLIT_BY_UNIT 關掉就不分）
  中標（商品總表 O 中盒貼標：V＝需貼中標）
分成一份一份的拆單表。一份拆單表＝一張 EIP 採購單（一個 EIP 採購單號）。

對商品總表：先用 (永豐料號, 單位)；對不到再用整合表的「報價備註」（組出商品，例如 M60019810 的報價備註是
M10254053，商品總表裡是後者）；再對不到才退回「同料號別的單位」並標出來。都對不到的品項不進拆單表，畫面列出來。
下採料號用對到的那個商品總表料號（組出商品就是報價備註那個）。

箱數＝出貨數量 ÷ 整合表箱入數（Alice：EIP 一律用 Integration 的 Q 欄箱數、單位是箱）。出貨數量用 ② 訂單明細
現在的數字（人改過的照改過的）。箱數不是整數的不能產 EIP 採購表，擋下來講清楚是哪幾個。
整合表箱入數跟商品總表同單位的箱入數不一樣時只提醒（例如 M10403852 盒：整合表 48、商品總表 4），箱數照整合表。

存檔：按「產出拆單表」時把每份的品項存下來（mst_mars_splits.items_json）。還沒填 EIP 採購單號的，下次產出會更新；
填了之後就不再動它——那份已經送去 EIP 了，訂單之後有變只在畫面上標「EIP 送出後訂單有變」，要重拆就先把單號清掉。
"""
import datetime as _dt
import zipfile

import purchase as _purchase  # 採購表轉換：瑪氏的 EIP 採購表範本與備註格式照它的（mars_template.xls）

from .common import *  # noqa: F401,F403
from .products import load_products
from .po import load_settings, po_filename, po_missing, warehouse_rows

ITEM_KEYS = ("po_number", "sku_id", "yf_sku", "purchase_code", "unit", "qty_ship", "box_file", "cases")


def _valid_date(s):
    try:
        _dt.date.fromisoformat(s)
        return True
    except (TypeError, ValueError):
        return False


def _mars_orders(conn, date_from, date_to):
    cfg = _line_groups()
    rows = _rows(conn.execute(
        "SELECT * FROM mst_orders WHERE delivery_date != '' AND delivery_date >= ? AND delivery_date <= ? "
        "ORDER BY delivery_date, po_number, sku_id", (date_from, date_to)))
    return [o for o in rows if _group_of(o["line"], cfg) == LINE]


def _match(o, by_key, by_code):
    """(商品總表那列, 怎麼對到的, 提醒) """
    yf, unit, quote = o["yf_sku"] or "", o["unit"] or "", o.get("quote_note") or ""
    if (yf, unit) in by_key:
        return by_key[(yf, unit)], "料號", ""
    if quote and quote != yf and (quote, unit) in by_key:
        return by_key[(quote, unit)], "報價備註", f"整合表料號 {yf} 在商品總表找不到，用報價備註 {quote} 對到"
    for code, how in ((yf, "料號"), (quote, "報價備註")):
        if code and code in by_code:
            units = by_code[code]
            p = units.get("箱") or next(iter(units.values()))
            return p, how, f"商品總表沒有 {code} 的「{unit}」這個單位，用「{p['unit']}」那列的資料"
    return None, "", ""


def build_items(conn, date_from, date_to):
    """把期間內的瑪氏訂單逐列對商品總表，回傳 (items, unmatched, zero)。"""
    by_key, by_code = load_products(conn)
    items, unmatched, zero = [], [], 0
    for o in _mars_orders(conn, date_from, date_to):
        qty = o["qty_ship"]
        if not qty:
            zero += 1
            continue
        box = o["box_size_file"]
        cases = round(qty / box, 4) if box else None
        p, via, note = _match(o, by_key, by_code)
        base = {"po_number": o["po_number"], "sku_id": o["sku_id"], "barcode": o["barcode"], "yf_sku": o["yf_sku"],
                "quote_note": o.get("quote_note") or "", "product_name": o["product_name"], "brand": o["brand"],
                "unit": o["unit"], "qty_ship": qty, "qty_coupang": o["qty_coupang"], "box_file": box, "cases": cases,
                "delivery_date": o["delivery_date"], "warehouse": o["warehouse"], "address": o.get("address") or "",
                "order_type": o["order_type"], "line": o["line"], "qty_overridden": bool(o["qty_ship_overridden"])}
        if p is None:
            unmatched.append(base)
            continue
        issues = []
        if note:
            issues.append(note)
        if not box:
            issues.append("整合表箱入數是空的，算不出箱數")
        elif cases is not None and abs(cases - round(cases)) > 1e-6:
            issues.append(f"箱數 {qty}÷{box}＝{cases:g} 不是整數")
        if box and p["unit"] == (o["unit"] or "") and p["box_qty"] and abs(p["box_qty"] - box) > 1e-6:
            issues.append(f"箱入數：整合表 {box:g}、商品總表 {p['box_qty']:g}（箱數照整合表）")
        base.update({
            "purchase_code": p["yf_sku"], "via": via, "mars_code": p["mars_code"], "category_name": p["category"],
            "category": cat_code(p["category"]), "mars_name": p["name"], "price": p["price"],
            "inner_per_case": p["inner_per_case"], "pcs_per_case": p["pcs_per_case"], "mars_unit": p["unit"],
            "mars_box_qty": p["box_qty"], "label": p["inner_label"], "note": p["note"], "po_case_note": p["po_case_note"],
            "issues": issues,
            "blocking": (not box) or (cases is not None and abs(cases - round(cases)) > 1e-6) or not cat_code(p["category"]),
        })
        if not base["category"]:
            issues.append(f"商品總表的 Category 是「{p['category'] or '空白'}」，分不出品類")
        items.append(base)
    return items, unmatched, zero


def split_key(it):
    unit = it["unit"] if SPLIT_BY_UNIT else "箱"
    return "|".join([it["po_number"], it["delivery_date"], it["warehouse"] or "", it["category"] or "?", unit or "", it["label"]])


def make_filename(g, seq=1):
    unit = g["unit"] if SPLIT_BY_UNIT else "箱"
    return (f"訂單系統拆單表_{g['po_number']}_{g['delivery_date'].replace('-', '')}_{g['warehouse']}_"
            f"{g['category'] or '品類不明'}_{unit}_{label_text(g['label'])}_{seq:02d}.xlsx")


def group_items(items):
    groups = collections.OrderedDict()
    for it in sorted(items, key=lambda x: (x["delivery_date"], x["po_number"], x["warehouse"] or "", x["category"] or "", x["unit"] or "", x["label"], x["sku_id"])):
        k = split_key(it)
        g = groups.setdefault(k, {"split_key": k, "po_number": it["po_number"], "delivery_date": it["delivery_date"],
                                  "warehouse": it["warehouse"] or "", "category": it["category"] or "",
                                  "unit": it["unit"] if SPLIT_BY_UNIT else "箱", "label": it["label"], "items": []})
        g["items"].append(it)
    for g in groups.values():
        g["item_count"] = len(g["items"])
        g["cases_total"] = round(sum(i["cases"] or 0 for i in g["items"]), 4)
        g["blocking"] = [f"{i['yf_sku']} {i['product_name']}：{'、'.join(i['issues'])}" for i in g["items"] if i["blocking"]]
        g["filename"] = make_filename(g)
    return list(groups.values())


def _sig(items):
    """比對「存下來的」跟「現在的」品項用：只看會影響 EIP 的欄位。"""
    return sorted((i["po_number"], i["sku_id"], i.get("purchase_code") or "", i["unit"] or "", i["qty_ship"], i["box_file"]) for i in items)


def _diff_text(old, new):
    o = {(i["sku_id"]): i for i in old}; n = {(i["sku_id"]): i for i in new}
    out = []
    for k in n.keys() - o.keys():
        out.append(f"多了 {n[k]['yf_sku']}（{n[k]['qty_ship']}）")
    for k in o.keys() - n.keys():
        out.append(f"少了 {o[k]['yf_sku']}")
    for k in o.keys() & n.keys():
        if o[k]["qty_ship"] != n[k]["qty_ship"]:
            out.append(f"{n[k]['yf_sku']} 出貨 {o[k]['qty_ship']}→{n[k]['qty_ship']}")
    return "；".join(out)


def _saved(conn, date_from, date_to):
    return {s["split_key"]: s for s in _rows(conn.execute(
        "SELECT * FROM mst_mars_splits WHERE delivery_date >= ? AND delivery_date <= ?", (date_from, date_to)))}


def _view(date_from, date_to):
    conn = get_conn()
    try:
        items, unmatched, zero = build_items(conn, date_from, date_to)
        groups = group_items(items)
        saved = _saved(conn, date_from, date_to)
        n_products = conn.execute("SELECT COUNT(*) AS n FROM mst_mars_products").fetchone()["n"]
        po_settings = load_settings(conn)
        whs = {w["code"]: w for w in warehouse_rows(conn, po_settings)}
    finally:
        conn.close()
    out = []
    for g in groups:
        s = saved.pop(g["split_key"], None)
        row = {k: g[k] for k in ("split_key", "po_number", "delivery_date", "warehouse", "category", "unit", "label",
                                 "item_count", "cases_total", "blocking")}
        row["items"] = g["items"]
        if s is None:
            row.update({"id": None, "status": "new", "filename": g["filename"], "eip_po": "", "slot_time": ""})
        else:
            snap = json.loads(s["items_json"] or "[]")
            same = _sig(snap) == _sig(g["items"])
            if s["eip_po"]:
                status = "filled" if same else "changed_after_eip"
            else:
                status = "generated" if same else "changed"
            row.update({"id": s["id"], "status": status, "filename": s["filename"], "eip_po": s["eip_po"],
                        "slot_time": s["slot_time"], "generated_at": s["updated_at"] or s["created_at"],
                        "diff": "" if same else _diff_text(snap, g["items"])})
            if s["eip_po"]:          # 送出去的是存下來的那份
                row["items"] = snap; row["item_count"] = s["item_count"]; row["cases_total"] = s["cases_total"]
        out.append(row)
    for s in saved.values():          # 存過、但現在訂單已經拆不出這份了
        snap = json.loads(s["items_json"] or "[]")
        out.append({"id": s["id"], "split_key": s["split_key"], "po_number": s["po_number"], "delivery_date": s["delivery_date"],
                    "warehouse": s["warehouse"], "category": s["category"], "unit": s["unit"], "label": s["label"],
                    "item_count": s["item_count"], "cases_total": s["cases_total"], "blocking": [], "items": snap,
                    "status": "gone", "filename": s["filename"], "eip_po": s["eip_po"], "slot_time": s["slot_time"],
                    "diff": "現在的訂單已經沒有這份（品項被拿掉、改期或改倉）"})
    for r in out:                     # ③ 瑪氏採購單：這份差什麼才產得出來（空＝可以按）
        wh = whs.get(r["warehouse"]) or {"missing": ["地址", "電話", "ship-to"]}
        r["po_missing"] = po_missing(r, r["items"], wh, po_settings) if r["id"] else ["先按「產出拆單表」"]
        r["po_filename"] = po_filename(r)
    out.sort(key=lambda r: (r["delivery_date"], r["po_number"], r["warehouse"], r["category"], r["unit"], r["label"]))
    wh_missing = sorted(code for code, w in whs.items() if w["missing"] and any(r["warehouse"] == code for r in out))
    return {"splits": out, "unmatched": unmatched, "zero_rows": zero, "has_products": n_products > 0, "wh_missing": wh_missing,
            "summary": {"pos": len({r["po_number"] for r in out}), "files": len(out),
                        "cases": round(sum(r["cases_total"] or 0 for r in out), 4),
                        "items": sum(r["item_count"] for r in out), "unmatched": len(unmatched),
                        "filled": sum(1 for r in out if r["eip_po"]),
                        "po_ready": sum(1 for r in out if not r["po_missing"]),
                        "warnings": sum(1 for r in out for i in r["items"] if i.get("issues"))}}


def _range_args(src):
    date_from, date_to = norm_text(src.get("from")), norm_text(src.get("to"))
    if not (_valid_date(date_from) and _valid_date(date_to)):
        return None, None
    return tuple(sorted([date_from, date_to]))


@mars_bp.route("/api/mars/splits")
def api_mars_splits():
    date_from, date_to = _range_args(request.args)
    if not date_from:
        return jsonify({"error": "請選到貨日的起訖（像 2026-09-21）。"}), 400
    v = _view(date_from, date_to)
    v.update({"from": date_from, "to": date_to, "split_by_unit": SPLIT_BY_UNIT})
    return jsonify(v)


# ── 產檔 ─────────────────────────────────────────────────────────────────────
SPLIT_COLUMNS = [   # 拆單表的欄：前面是整合表（訂單），後面是商品總表帶進來的（Alice：A、B、E、F、H、I、L、O、Q、R）
    ("PO單號", "po_number"), ("到貨倉別", "warehouse"), ("地址", "address"), ("交付日期", "delivery_date"),
    ("SKU ID", "sku_id"), ("條碼(國條)", "barcode"), ("永豐料號", "yf_sku"), ("報價備註", "quote_note"),
    ("品牌", "brand"), ("品名", "product_name"), ("出貨數量", "qty_ship"), ("單位", "unit"), ("箱入數", "box_file"),
    ("箱數", "cases"), ("下採料號", "purchase_code"),
    ("瑪氏貨號", "mars_code"), ("Category", "category_name"), ("品名(永豐建檔)", "mars_name"), ("系統價格(未稅)", "price"),
    ("每箱中盒數", "inner_per_case"), ("每箱產品數(最小單位)", "pcs_per_case"), ("單位類別", "mars_unit"),
    ("中盒貼標", "label"), ("備註", "note"), ("採購單箱備註", "po_case_note"),
]


def split_workbook(s, items):
    """一份拆單表：表頭一列＋品項＋合計，最後兩欄是 EIP 採購單號、約倉時間（回填後重新下載就有）。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "拆單表"
    head = [h for h, _ in SPLIT_COLUMNS] + ["EIP採購單號", "約倉時間"]
    ws.append(head)
    for it in items:
        ws.append([it.get(k) for _, k in SPLIT_COLUMNS] + [s.get("eip_po") or "", s.get("slot_time") or ""])
    ws.append([])
    total = ["合計", "", "", "", "", "", "", "", "", f"{len(items)} 個品項", sum(i["qty_ship"] or 0 for i in items), "", "",
             round(sum(i["cases"] or 0 for i in items), 4)]
    ws.append(total)
    fill_o = PatternFill("solid", fgColor="DCE3FF"); fill_m = PatternFill("solid", fgColor="D9F7F1")
    n_order = 15
    for c in ws[1]:
        c.font = Font(bold=True); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = fill_o if c.column <= n_order else fill_m
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            if c.column in (1, 5, 6, 7, 8, 15, 16):
                c.number_format = "@"
    widths = {1: 16, 2: 9, 3: 30, 4: 11, 5: 17, 6: 16, 7: 13, 8: 13, 9: 14, 10: 46, 11: 9, 12: 6, 13: 7, 14: 7, 15: 13,
              16: 11, 17: 10, 18: 40, 19: 12, 20: 9, 21: 11, 22: 8, 23: 8, 24: 30, 25: 20, 26: 14, 27: 18}
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def eip_rows(s, items):
    """EIP 上傳用採購表的資料列：同一個下採料號加總箱數，單位一律箱。備註照採購表轉換瑪氏的格式。"""
    d = _dt.date.fromisoformat(s["delivery_date"])
    remark = _purchase.build_remark("mars", [s["po_number"]], f"{d.month:02d}{d.day:02d}", s["warehouse"])
    qty = collections.OrderedDict()
    for it in items:
        qty[it["purchase_code"]] = qty.get(it["purchase_code"], 0) + (it["cases"] or 0)
    return [{"material_no": code, "qty": int(round(q)), "remark": remark} for code, q in qty.items()], remark


def eip_file(s, items):
    rows, remark = eip_rows(s, items)
    return _purchase.fill_template("mars", rows, remark)


def _stem(filename):
    return filename[:-5] if filename.endswith(".xlsx") else filename


@mars_bp.route("/api/mars/splits/generate", methods=["POST"])
def api_mars_generate():
    """把期間內的拆單存起來、打包下載：每份一個拆單表（.xlsx）＋一個 EIP 上傳用採購表（.xls）。"""
    payload = request.get_json(silent=True) or {}
    date_from, date_to = _range_args(payload)
    if not date_from:
        return jsonify({"error": "請選到貨日的起訖。"}), 400
    operator = _operator()
    conn = get_conn()
    try:
        if conn.execute("SELECT COUNT(*) AS n FROM mst_mars_products").fetchone()["n"] == 0:
            return jsonify({"error": "還沒上傳瑪氏商品總表，拆不出品類和中標。"}), 400
        items, unmatched, _zero = build_items(conn, date_from, date_to)
        groups = group_items(items)
        if not groups:
            return jsonify({"error": "這段期間沒有瑪氏的訂單可以拆。"}), 400
        blocking = [b for g in groups for b in g["blocking"]]
        if blocking:
            return jsonify({"error": "有品項的箱數算不出整數或分不出品類，EIP 採購表不能這樣送。先處理這幾個：",
                            "details": blocking[:30]}), 400
        if unmatched and not payload.get("ack_unmatched"):
            return jsonify({"error": f"有 {len(unmatched)} 個品項對不到瑪氏商品總表，不會進拆單表。", "needs_ack": True,
                            "details": [f"{u['po_number']} {u['yf_sku']} {u['product_name']}" for u in unmatched[:30]]}), 409
        saved = _saved(conn, date_from, date_to)
        stamp = now(); files = []
        for g in groups:
            s = saved.pop(g["split_key"], None)
            snap = json.dumps(g["items"], ensure_ascii=False)
            if s is None:
                cur = conn.execute(
                    """INSERT INTO mst_mars_splits (split_key, po_number, delivery_date, warehouse, category, unit, label, seq,
                       filename, items_json, item_count, cases_total, created_by, created_at, updated_by, updated_at)
                       VALUES (?,?,?,?,?,?,?,1,?,?,?,?,?,?,?,?)""",
                    (g["split_key"], g["po_number"], g["delivery_date"], g["warehouse"], g["category"], g["unit"], g["label"],
                     g["filename"], snap, g["item_count"], g["cases_total"], operator, stamp, operator, stamp))
                s = {"id": cur.lastrowid, "eip_po": "", "slot_time": "", "filename": g["filename"], **g}
                use_items = g["items"]
            elif s["eip_po"]:
                use_items = json.loads(s["items_json"] or "[]")      # 已送 EIP：照當時那份
            else:
                conn.execute("UPDATE mst_mars_splits SET items_json = ?, item_count = ?, cases_total = ?, updated_by = ?, updated_at = ? WHERE id = ?",
                             (snap, g["item_count"], g["cases_total"], operator, stamp, s["id"]))
                use_items = g["items"]
            files.append((s, use_items))
        for s in saved.values():          # 拆不出來了、也還沒送 EIP 的舊拆單：刪掉免得有人拿去用
            if not s["eip_po"]:
                conn.execute("DELETE FROM mst_mars_splits WHERE id = ?", (s["id"],))
        _log(conn, LINE, "", "", "", "mars_split", "瑪氏拆單", "", f"{len(files)} 份",
             operator, "manual", f"到貨日 {date_from}～{date_to}")
        conn.commit()
    finally:
        conn.close()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for s, its in files:
            stem = _stem(s["filename"])
            zf.writestr(f"{stem}.xlsx", split_workbook(s, its))
            zf.writestr(f"{stem}_EIP上傳.xls", eip_file(s, its))
    buf.seek(0)
    name = f"瑪氏拆單_{date_from.replace('-', '')}" + ("" if date_from == date_to else f"-{date_to.replace('-', '')}") + ".zip"
    return send_file(buf, as_attachment=True, download_name=name, mimetype="application/zip")


@mars_bp.route("/api/mars/splits/<int:split_id>/file")
def api_mars_split_file(split_id):
    """單份重新下載（照存下來的那份）：kind=split 拆單表、kind=eip EIP 上傳用採購表。"""
    kind = request.args.get("kind", "split")
    conn = get_conn()
    try:
        s = _row(conn.execute("SELECT * FROM mst_mars_splits WHERE id = ?", (split_id,)))
    finally:
        conn.close()
    if s is None:
        return jsonify({"error": "找不到這份拆單表，可能已經重拆過了。"}), 404
    items = json.loads(s["items_json"] or "[]")
    if kind == "eip":
        return send_file(io.BytesIO(eip_file(s, items)), as_attachment=True, download_name=f"{_stem(s['filename'])}_EIP上傳.xls",
                         mimetype="application/vnd.ms-excel")
    return send_file(io.BytesIO(split_workbook(s, items)), as_attachment=True, download_name=s["filename"],
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@mars_bp.route("/api/mars/splits/<int:split_id>", methods=["PUT"])
def api_mars_split_update(split_id):
    """回填 EIP 採購單號（PO＋9 碼，空白＝清掉）、約倉時間。一個 EIP 採購單號只能對一份拆單表。"""
    payload = request.get_json(silent=True) or {}
    operator = _operator()
    conn = get_conn()
    try:
        s = _row(conn.execute("SELECT * FROM mst_mars_splits WHERE id = ?", (split_id,)))
        if s is None:
            return jsonify({"error": "找不到這份拆單表，請重新整理。"}), 404
        sets, vals = [], []
        if "eip_po" in payload:
            eip = norm_text(payload.get("eip_po")).upper().replace(" ", "")
            if eip and not EIP_PO_RE.match(eip):
                return jsonify({"error": f"EIP 採購單號要像 PO202609004（PO 加 9 個數字），你填的是「{eip}」。"}), 400
            if eip:
                dup = _row(conn.execute("SELECT filename FROM mst_mars_splits WHERE eip_po = ? AND id != ?", (eip, split_id)))
                if dup:
                    return jsonify({"error": f"{eip} 已經填在「{dup['filename']}」，一個採購單號只能對一份拆單表。"}), 400
            if eip != (s["eip_po"] or ""):
                sets.append("eip_po = ?"); vals.append(eip)
                _log(conn, LINE, s["po_number"], "", "", "mars_eip_po", "EIP 採購單號", s["eip_po"], eip, operator, "manual", s["filename"])
        if "slot_time" in payload:
            slot = str(payload.get("slot_time") or "").strip()[:80]     # 照人打的存，不轉全形半形（「（1台車）」要原樣進採購單）
            if slot != (s["slot_time"] or ""):
                sets.append("slot_time = ?"); vals.append(slot)
                _log(conn, LINE, s["po_number"], "", "", "mars_slot_time", "約倉時間", s["slot_time"], slot, operator, "manual", s["filename"])
        if sets:
            sets += ["updated_by = ?", "updated_at = ?"]; vals += [operator, now()]
            conn.execute(f"UPDATE mst_mars_splits SET {', '.join(sets)} WHERE id = ?", vals + [split_id])
            conn.commit()
        fresh = _row(conn.execute("SELECT id, eip_po, slot_time, filename FROM mst_mars_splits WHERE id = ?", (split_id,)))
    finally:
        conn.close()
    return jsonify({"ok": True, "split": fresh})


__all__ = ["build_items", "group_items", "split_key", "make_filename", "split_workbook", "eip_rows",
           "api_mars_splits", "api_mars_generate", "api_mars_split_file", "api_mars_split_update"]
