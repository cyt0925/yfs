"""拆單底稿填 EIP 採購單號（Jerry 2026-10-14）。

OP 自己做的拆單底稿（「拆解訂單」那一頁，一列一個品項）有一欄要一列一列手填 EIP 採購單號
（表頭寫「想改倉庫」，實際放的是 PO202610362 這種）。系統裡其實有答案：拆單表每一份回填過 EIP 採購單號，
也存了每份有哪些（酷澎 PO, SKU ID）。所以讓 OP 把底稿丟進來，系統逐列對、填好、吐回同一份檔案。

規矩（Jerry 確認的示意）：
- 吐回原本那份檔，每一頁、格式都不動，只填那一欄。
- 對到一個 EIP 採購單號：空的就填；本來填的一樣就不動；本來填的不一樣不蓋掉，塗紅提醒。
- 系統裡還沒回填、對到好幾份、找不到這個單號＋SKU：留空，塗黃。
- 同一個單號＋SKU 改過到貨日會有新、舊兩份拆單（舊的那份已回填就留著），所以先用底稿的「交期」挑同一天的那份。
"""
import datetime as _dt
import os
import re as _re

from openpyxl.styles import PatternFill

from .common import *  # noqa: F401,F403

EIP_HEADS = ("想改倉庫", "EIP採購單號", "採購單號")    # 這一欄放 EIP 採購單號（底稿的表頭寫「想改倉庫」）
SKU_HEADS = ("SKUNo.", "SKUNO.", "SKUID", "SKU")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="FFC7CE")
ST_TEXT = {"fill": "已填入", "same": "原本已填、與系統相同", "diff": "原本填的與系統不同（未覆蓋）",
           "pending": "系統裡還沒回填 EIP 採購單號", "multi": "系統裡對到多份 EIP 採購單", "missing": "系統裡找不到這個單號＋SKU"}


def _txt(v):
    """格子內容轉成比對用的字串：13000000621842.0 → 13000000621842。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _head(v):
    return str(v or "").replace(" ", "").replace("\n", "").replace("　", "")


def find_sheet(wb):
    """回 (工作表, 表頭列, {欄位: 欄號})；優先「拆解訂單」，其他頁有「單號」和 SKU 欄也行。找不到回 None。"""
    for ws in sorted(wb.worksheets, key=lambda w: 0 if w.title == "拆解訂單" else 1):
        for r in range(1, 11):
            heads = {}
            for c in range(1, min(ws.max_column or 0, 60) + 1):
                h = _head(ws.cell(r, c).value)
                if h and h not in heads:
                    heads[h] = c
            sku = next((heads[h] for h in SKU_HEADS if h in heads), None)
            if "單號" not in heads or not sku:
                continue
            eip = next((heads[h] for h in EIP_HEADS if h in heads), None)
            if eip is None and "倉庫" in heads:
                eip = heads["倉庫"] + 1
            if eip:
                return ws, r, {"po": heads["單號"], "sku": sku, "eip": eip, "yf": heads.get("永豐料號"), "name": heads.get("品名"),
                               "date": heads.get("交期") or heads.get("到貨日")}
    return None


def _date(v):
    """交期格子 → YYYY-MM-DD；讀不出來回空字串。"""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    m = _re.match(r"^\s*(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", str(v or ""))
    return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else ""


def system_map(conn):
    """(酷澎 PO, SKU ID) → [(EIP 採購單號（空＝那份還沒回填）, 到貨日)]，一份拆單一筆。"""
    m = {}
    for s in _rows(conn.execute("SELECT eip_po, delivery_date, items_json FROM mst_mars_splits")):
        for it in json.loads(s["items_json"] or "[]"):
            k = (_txt(it.get("po_number")), _txt(it.get("sku_id")))
            m.setdefault(k, []).append(((s["eip_po"] or "").strip(), s["delivery_date"] or ""))
    known = {(_txt(r["po_number"]), _txt(r["sku_id"])) for r in _rows(conn.execute("SELECT po_number, sku_id FROM mst_orders"))}
    return m, known


def fill(wb, conn):
    """直接改 wb，回 {"sheet", "rows": [...], "counts": {...}}；底稿格式不對丟 ValueError。"""
    found = find_sheet(wb)
    if not found:
        raise ValueError("找不到拆單底稿：需要有一頁的表頭含「單號」、「SKU No.」，以及放 EIP 採購單號的那一欄（表頭「想改倉庫」或「EIP 採購單號」）。")
    ws, hr, col = found
    m, known = system_map(conn)
    rows, counts = [], {k: 0 for k in ST_TEXT}
    for r in range(hr + 1, (ws.max_row or hr) + 1):
        po, sku = _txt(ws.cell(r, col["po"]).value), _txt(ws.cell(r, col["sku"]).value)
        if not po and not sku:
            continue
        cell = ws.cell(r, col["eip"])
        now_v = _txt(cell.value)
        got = m.get((po, sku), [])
        day = _date(ws.cell(r, col["date"]).value) if col["date"] else ""
        if day and any(d == day for _e, d in got):          # 改過到貨日：只看同一天的那份
            got = [g for g in got if g[1] == day]
        filled = {e for e, _d in got if e}
        if len(filled) == 1:
            eip = next(iter(filled))
            if not now_v:
                cell.value, st = eip, "fill"
            elif now_v == eip:
                st = "same"
            else:
                cell.fill, st = RED, "diff"
        elif len(filled) > 1:
            eip, st = "、".join(sorted(filled)), "multi"
        else:
            eip, st = "", ("pending" if got or (po, sku) in known else "missing")
        if st in ("pending", "multi", "missing") and not now_v:
            cell.fill = YELLOW
        counts[st] += 1
        rows.append({"row": r, "po": po, "sku": sku, "yf": _txt(ws.cell(r, col["yf"]).value) if col["yf"] else "",
                     "name": _txt(ws.cell(r, col["name"]).value)[:40] if col["name"] else "", "was": now_v, "eip": eip,
                     "status": st, "text": ST_TEXT[st]})
    if not rows:
        raise ValueError(f"「{ws.title}」這一頁沒有品項。")
    return {"sheet": ws.title, "rows": rows, "counts": counts}


def _load(up):
    if up is None or not up.filename:
        raise ValueError("沒有收到檔案。")
    if not up.filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("請上傳 .xlsx 檔（舊版 .xls 請先另存成 .xlsx）。")
    try:
        return openpyxl.load_workbook(io.BytesIO(up.read()), keep_vba=up.filename.lower().endswith(".xlsm"))
    except Exception as exc:  # noqa: BLE001
        raise ValueError("無法開啟此 Excel 檔，請確認檔案格式。") from exc


@mars_bp.route("/api/mars/fill_eip/preview", methods=["POST"])
def api_fill_eip_preview():
    """先看：幾列會填、幾列對不到。不存任何東西。"""
    conn = get_conn()
    try:
        out = fill(_load(request.files.get("file")), conn)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    finally:
        conn.close()
    return jsonify(out)


@mars_bp.route("/api/mars/fill_eip", methods=["POST"])
def api_fill_eip():
    """填好下載：同一份檔案，只改 EIP 採購單號那一欄。"""
    up = request.files.get("file")
    conn = get_conn()
    try:
        wb = _load(up)
        out = fill(wb, conn)
        c = out["counts"]
        _log(conn, LINE, "", "", "", "mars_fill_eip", "拆單底稿填 EIP 採購單號", "",
             f"填入 {c['fill']} 列、對不到 {c['pending'] + c['multi'] + c['missing']} 列、與系統不同 {c['diff']} 列", _operator(), "manual", up.filename)
        conn.commit()
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    finally:
        conn.close()
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    stem, ext = os.path.splitext(up.filename)
    resp = send_file(buf, as_attachment=True, download_name=f"{stem}_已填EIP採購單號{ext}",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["X-Mars-Filled"] = str(out["counts"]["fill"])
    return resp


__all__ = ["find_sheet", "fill", "system_map", "api_fill_eip_preview", "api_fill_eip"]
