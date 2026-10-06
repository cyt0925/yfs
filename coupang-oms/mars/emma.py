"""④ EMMA 匯入檔（酷澎訂單匯入）：取代原本的拆單表（Alice 2026-09-30）。

格式照 Kate 系統匯出的 Coupang_PO_Order（工作表「酷澎訂單匯入」）：
  A 出貨備註＝那份的 EIP 採購單號（還沒填就空著，畫面上提醒）   B 訂單編號＝酷澎 PO
  C 收件人、F 客編、G BillTo、H ShipTo、J 付款方式＝固定字（採購單設定裡可改）   I 倉庫別 空白
  D 收件人手機＝倉庫資料的電話轉成 +886-02-55927598／+886-0911556291   E 收件地址＝訂單上的地址
  K 料號＝下採料號（報價備註）  L 單位 箱  M 數量＝箱數  N 單價＝酷澎下單單價(含稅)×箱入數  O 金額小計＝M×N
  最後一列 M 欄 SUBTOTAL 加總。同一張 PO 同一個下採料號合成一列。
真檔 0921 對過：單價 129×6＝774、70×48＝3360，小計都吻合。

一份拆單表一個檔，也能勾幾份合併成一個檔、或整段期間全部合併（EMMA 一次匯一個檔）。
檔名（Jerry 定的）：酷澎訂單匯入_1002交貨-TAO3、TAO6、TAO17.xlsx；單份加 PO、品類、中標免得撞名。
"""
import datetime as _dt

from .common import *  # noqa: F401,F403
from .po import load_settings, warehouse_rows

EMMA_HEADERS = ["出貨備註", "訂單編號", "收件人", "收件人手機", "收件地址", "客編", "BillTo", "ShipTo", "倉庫別", "付款方式",
                "料號", "單位", "數量", "單價", "金額小計"]
EMMA_WIDTHS = [10.5, 16, 19, 17.2, 19, 8.1, 6.1, 7.2, 8.3, 12.8, 11.6, 6.1, 6.1, 6.1, 10.5]


def intl_phone(phone):
    """02-5592-7598 → +886-02-55927598；0911-556-291 → +886-0911556291（照 Kate 系統匯出的寫法）。看不懂照原樣。"""
    raw = str(phone or "").strip()
    if not raw:
        return ""
    if raw.startswith("+"):
        return raw
    parts = [p for p in re.split(r"[\s\-()]+", raw) if p]
    digits = "".join(parts)
    if not digits.isdigit():
        return raw
    if digits.startswith("09"):
        return f"+886-{digits}"
    if len(parts) >= 2 and parts[0].startswith("0"):
        return f"+886-{parts[0]}-{''.join(parts[1:])}"
    return f"+886-{digits}"


def _unit_prices(conn, items):
    """酷澎下單單價(含稅)從 ② 訂單抓（舊的拆單快照沒存這欄）。"""
    keys = {(i["po_number"], i["sku_id"]) for i in items}
    out = {}
    for po, sku in keys:
        r = conn.execute("SELECT unit_price FROM mst_orders WHERE po_number = ? AND sku_id = ?", (po, sku)).fetchone()
        if r is not None:
            out[(po, sku)] = r["unit_price"]
    return out


def emma_rows(conn, splits, settings=None, whs=None):
    """splits：[(拆單表列, 品項 list)]。回傳 (rows, warnings)。"""
    settings = settings or load_settings(conn)
    whs = whs if whs is not None else {w["code"]: w for w in warehouse_rows(conn, settings)}
    all_items = [i for _, its in splits for i in its]
    prices = _unit_prices(conn, all_items)
    rows, warnings = [], []
    merged = collections.OrderedDict()
    for s, items in splits:
        wh = whs.get(s["warehouse"]) or {}
        phone = intl_phone(wh.get("phone") or "")
        if not phone:
            warnings.append(f"{s['warehouse']} 的電話沒填，收件人手機留空")
        if not s.get("eip_po"):
            warnings.append(f"{s['filename']}：EIP 採購單號還沒填，出貨備註留空")
        for it in items:
            code = it.get("purchase_code") or it.get("yf_sku") or ""
            price = it.get("unit_price")
            if price is None:
                price = prices.get((it["po_number"], it["sku_id"]))
            box = it.get("box_file") or 0
            unit_price = round(price * box, 2) if price is not None and box else None
            key = (s["id"], it["po_number"], code)
            r = merged.get(key)
            if r is None:
                r = merged[key] = {"remark": s.get("eip_po") or "", "po_number": it["po_number"], "recipient": settings["emma_recipient"],
                                   "phone": phone, "address": it.get("address") or wh.get("address") or "",
                                   "customer": settings["emma_customer"], "billto": settings["emma_billto"], "shipto": settings["emma_shipto"],
                                   "warehouse": "", "payment": settings["emma_payment"], "code": code, "unit": "箱", "qty": 0,
                                   "unit_price": unit_price}
            r["qty"] += it.get("cases") or 0
            if r["unit_price"] is None and unit_price is not None:
                r["unit_price"] = unit_price
    for r in merged.values():
        r["qty"] = int(round(r["qty"])) if abs(r["qty"] - round(r["qty"])) < 1e-6 else r["qty"]
        if r["unit_price"] is None:
            warnings.append(f"{r['po_number']} {r['code']}：② 訂單沒有酷澎下單單價，單價留空")
            r["subtotal"] = None
        else:
            if float(r["unit_price"]).is_integer():
                r["unit_price"] = int(r["unit_price"])
            r["subtotal"] = round(r["qty"] * r["unit_price"], 2)
            if float(r["subtotal"]).is_integer():
                r["subtotal"] = int(r["subtotal"])
        rows.append(r)
    return rows, sorted(set(warnings), key=warnings.index)


def emma_workbook(rows):
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "酷澎訂單匯入"
    ws.append(EMMA_HEADERS)
    for r in rows:
        ws.append([r["remark"], r["po_number"], r["recipient"], r["phone"], r["address"], r["customer"], r["billto"], r["shipto"],
                   r["warehouse"] or None, r["payment"], r["code"], r["unit"], r["qty"], r["unit_price"], r["subtotal"]])
    last = len(rows) + 1
    ws.cell(last + 1, 13).value = f"=SUBTOTAL(9,M2:M{last})"
    for i, w in enumerate(EMMA_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=last):
        row[1].number_format = "@"; row[10].number_format = "@"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def emma_filename(splits):
    """酷澎訂單匯入_1002交貨-TAO3、TAO6、TAO17.xlsx；單份多帶 PO、品類、中標（依單位拆時再加原單位）。"""
    dates = sorted({s["delivery_date"] for s, _ in splits})
    whs = []
    for s, _ in splits:
        if s["warehouse"] not in whs:
            whs.append(s["warehouse"])
    d = "、".join(f"{int(x[5:7]):02d}{int(x[8:10]):02d}" for x in dates)
    tail = ""
    if len(splits) == 1:
        s = splits[0][0]
        unit = f"({s['unit']})" if SPLIT_BY_UNIT and s["unit"] and s["unit"] != "箱" else ""
        tail = f"_{s['po_number']}_{s['category'] or '品類不明'}_{label_text(s['label'])}{unit}"
    return f"酷澎訂單匯入_{d}交貨-{'、'.join(whs)}{tail}.xlsx"


def emma_file(conn, splits, settings=None, whs=None):
    rows, warnings = emma_rows(conn, splits, settings, whs)
    return emma_workbook(rows), warnings


def _load_splits(conn, ids):
    out = []
    for sid in ids:
        s = _row(conn.execute("SELECT * FROM mst_mars_splits WHERE id = ?", (sid,)))
        if s is not None:
            out.append((s, json.loads(s["items_json"] or "[]")))
    out.sort(key=lambda x: (x[0]["delivery_date"], x[0]["po_number"], x[0]["warehouse"], x[0]["category"], x[0]["unit"], x[0]["label"]))
    return out


def _send_emma(conn, splits, ack):
    missing = [s["filename"] for s, _ in splits if not s.get("eip_po")]
    if missing and not ack:
        return jsonify({"error": f"有 {len(missing)} 份還沒填 EIP 採購單號，出貨備註會空著。", "needs_ack": True, "details": missing[:30]}), 409
    data, warnings = emma_file(conn, splits)
    resp = send_file(io.BytesIO(data), as_attachment=True, download_name=emma_filename(splits),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["X-Mars-Emma-Rows"] = str(len(splits))
    resp.headers["X-Mars-Emma-Warnings"] = str(len(warnings))
    return resp


@mars_bp.route("/api/mars/splits/<int:split_id>/emma")
def api_split_emma(split_id):
    """一份拆單表 → 一個 EMMA 匯入檔。EIP 單號沒填要 ack=1 才給（畫面會先問）。"""
    conn = get_conn()
    try:
        splits = _load_splits(conn, [split_id])
        if not splits:
            return jsonify({"error": "找不到這份拆單表，可能已經重拆過了。"}), 404
        return _send_emma(conn, splits, request.args.get("ack") == "1")
    finally:
        conn.close()


@mars_bp.route("/api/mars/emma", methods=["POST"])
def api_emma_merge():
    """勾選的幾份（ids）或整段到貨日（from／to）合併成一個 EMMA 匯入檔。"""
    payload = request.get_json(silent=True) or {}
    conn = get_conn()
    try:
        ids = payload.get("ids")
        if ids:
            ids = [int(x) for x in ids]
        else:
            d1, d2 = norm_text(payload.get("from")), norm_text(payload.get("to"))
            try:
                _dt.date.fromisoformat(d1); _dt.date.fromisoformat(d2)
            except ValueError:
                return jsonify({"error": "請選到貨日的起訖，或勾選要合併的拆單表。"}), 400
            d1, d2 = sorted([d1, d2])
            ids = [r["id"] for r in _rows(conn.execute("SELECT id FROM mst_mars_splits WHERE delivery_date >= ? AND delivery_date <= ?", (d1, d2)))]
        splits = _load_splits(conn, ids)
        if not splits:
            return jsonify({"error": "沒有可以合併的拆單表，先按「下載 EIP 採購單」，或那一列的按鈕。"}), 400
        return _send_emma(conn, splits, bool(payload.get("ack_missing_eip")))
    finally:
        conn.close()


__all__ = ["EMMA_HEADERS", "intl_phone", "emma_rows", "emma_workbook", "emma_filename", "emma_file", "api_split_emma", "api_emma_merge"]
