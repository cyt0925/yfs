"""④ 改單統計：每個月幾張 PO、幾張被改過、改了幾次、為什麼。主管要拿這個講「改單都是成本」。

名詞怎麼算（同事問起就照這裡講）：
- PO 歸哪個月：看它目前的交貨日（多個日期取最早那天）。沒排日期的另外一列「未排日期」。
- 被改過：這張 PO 在變動紀錄（mst_logs）裡有至少一次「改單事件」。
- 改單事件：同一張 PO、同一個來源、同一時間的變動算一次。因為整張 PO 改期會一個品項記一筆，
  一次動作記 20 筆不能算 20 次。我們改的看到「秒」（一次儲存同一秒寫完，兩次儲存分得開）；
  酷澎改的看到「分鐘」（匯入一批寫紀錄會跨秒）。
  - 酷澎改的（source=import）：整合表出貨數量／整合表交貨日變了、或品項從檔案消失。
    不看 qty_ship「隨整合表更新」那筆——它跟 qty_file_ship 是同一件事，算兩次會重複。
  - 我們改的（source=manual）：人手改出貨數量／交貨日。「恢復為整合表…」是把人改的撤掉，不算。
- 改了什麼：一次改單可能同時改數量和交期，所以「改數量」「改交期」「品項被拿掉」三類各自算次數，加起來會
  大於改單次數，這是正常的。改數量再分下修／上修（每個品項比新舊值；一次裡有下有上就算「有上有下」）。
- 原因：只有我們改的才有（缺貨／沒車／酷澎要求／其他）；2026-09-22 以前的紀錄沒有原因，列「未填」。
- 明細：每一次改單一列（時間、PO、誰、來源、改了什麼、原因），畫面上點任何數字都能看到它背後是哪幾次；
  Excel 也有「明細」分頁，主管想自己做樞紐分析就用那一頁。
- 測試：原因記成「測試」的改單（同事在正式站試功能時勾的）預設不算，統計頁「含測試」勾了才算；
  排除掉的次數另外回 test_events，畫面上寫一句「另有 N 次測試沒算」讓人知道有這回事。
- 估計工時：每次改單大約幾分鐘由主管給（畫面上填），系統只是乘出來，不是系統決定的。
資料從系統開始用（2026-09 中）才有，之前 Excel 時代的不在裡面。

2026-10-19 改版（Ivan 要追「原本下多少、最後出多少、為什麼改」，Jerry 看過示意圖）：
- 兩個系統合起來算：同一份訂單彙總表會匯進酷澎訂單管理系統（edit_logs，小真）和商品主檔自動化（mst_logs，Alice、妮可），
  兩邊都記「酷澎改了什麼」。每張 PO 每個品項、每個欄位（下單數量、訂單彙總表出貨數量、到貨日、倉別、品項移除）把兩邊看過的值
  照時間排成一條，值真的變了才算一次；誰先匯入就用誰看到的時間。所以同一個改動不會算兩次，只匯進其中一邊的也不會漏。
- 我方修改兩邊都算（各自是不同的動作）：商品主檔改出貨數量／到貨日、酷澎訂單管理改出貨數量／到貨日／倉別。
- 原因：每一次改單都要有原因，酷澎修改也要（很多是我們請酷澎改的）。原因只用大項（REASONS，草稿等 Alice、Ivan 定）；
  改的時候已經選了的（商品主檔、瑪氏）照舊對應過來；沒有的列「待補」，在改單統計頁勾好幾筆一起補，存在 mst_change_reasons。
- 原始下單→目前出貨用金額（數量 × 下單單價）：一張單有箱、盒、包，數量直接加起來沒有意思。
  原始＝兩邊最早看到的下單數量；目前＝兩邊最近更新的那筆出貨數量（品項移除算 0）。
- 品項移除之後又重新出現：重新出現那次也算一次酷澎修改，算「改數量（上修）」，金額差＝被拿掉前的數量 × 單價，
  跟被拿掉那次的金額差抵掉（東西回來了，目前出貨也有算它）。
- 畫面上的金額（Jerry 2026-10-19）：每月統計有原始下單、目前出貨、少出兩段；改單原因表有每個原因的金額差。
  上面「原始下單→目前出貨」那張數字卡、待補原因清單的金額差拿掉了（還不確定要不要）。
- 少出拆兩段（Jerry 2026-10-19「這怎算的」→「拆成兩段」）：
  進系統前就少了＝第一次匯入時訂單彙總表的出貨數量就比下單數量少（在 Kate 的系統先改的，沒有改單紀錄、沒地方補原因）；
  進系統後才改＝之後的變動（有改單紀錄、有原因）。第一次的出貨數量＝兩邊出貨數量最早那筆紀錄改之前的值，沒紀錄就是現在的值。
- 酷澎拆單（原訂單欄有值、只有酷澎訂單管理有存）：原單被拿掉的品項如果在拆出來的新單裡，新單那項接回原單算，
  不算少出也不算新的下單（不然原單少出、新單又多一筆下單，少出會算太多）。兩張單都要在統計期間內才接。

2026-10-21 改版（Jerry 看過示意圖「OK 照這樣做」）：修改次數改成以 PO 修改紀錄（rounds.py 的 po_rounds）為準，一輪處理完按一次 OK 算一次。
- 曾改單的 PO＝至少有一筆修改紀錄；修改次數（OK）分這輪在哪個階段結束（第一／二／三階段）、做了什麼（有下修／有改交期倉別／有拆單，一輪可複選所以加總會大於次數）、
  主因（一輪可勾好幾個，照小真的清單分類別列，沒用到的不列）。還沒記 OK 的 PO（上一個 OK 之後有改動）取代原本的「待補原因」。
- 上面那套「系統抓到的變動」（兩個系統合起來算的改單事件）不再當次數，每月統計留一欄灰字「系統抓到的變動次數」（每個月都有，2026-10-20 以前的月份也看得到），
  點了看明細；PO 張數、原始下單→目前出貨金額照舊每個月都算。舊的六個草稿大項原因（mst_change_reasons）不再列、不再補，表留著。
- 修改紀錄從 2026-10-20（rounds.SINCE）起才有，之前的月份修改次數會是 0，畫面上會寫。
"""
from .common import *  # noqa: F401,F403 — 共用工具、Flask、db、openpyxl 都從這裡來

import rounds

UNDATED = "未排日期"
OLD_REASON = {"缺貨": "原廠／產銷缺貨", "沒車": "運能（沒車）"}     # 商品主檔改單時選的舊原因，只在「系統抓到的變動」明細裡帶著
CAT_LABEL = {"qty": "數量下修", "date": "改交期／倉別", "split": "拆單", "other": "其他"}
STAGE_KEYS = [k for k, _a, _b in rounds.STAGES]
KIND_KEYS = [k for k, _v in rounds.KINDS]
SYS_LABEL = {"c": "酷澎訂單管理", "m": "商品主檔"}
_RESTORE_PREFIXES = ("恢復為整合表", "恢復為訂單彙總表")   # 舊紀錄寫「整合表」，2026-10-07 起寫「訂單彙總表」


def _stats_range():
    """回 (起月, 迄月, 月份清單)；all=1 是全部期間（Jerry 2026-10-19「我也要可以顯示全部的」），回 (None, None, ["全部"])。"""
    if request.args.get("all") in ("1", "true"):
        return None, None, ["全部"]
    month = request.args.get("month") or _this_month()
    month_to = request.args.get("month_to") or month
    if not _valid_month(month) or not _valid_month(month_to):
        return None, None, None
    months = _month_span(month, month_to)
    return months[0], months[-1], months


def _po_month(dates):
    real = sorted(d for d in dates if d)
    return real[0][:7] if real else UNDATED


KIND_LABEL = {"qty": "改數量", "date": "改到貨日", "wh": "改倉別", "gone": "品項移除"}
DIR_LABEL = {"down": "下修", "up": "上修", "mixed": "上修與下修", "": ""}
CANON_LABEL = {"order_qty": "下單數量", "file_ship": "訂單彙總表出貨數量", "ship": "出貨數量"}


def _num(v):
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def _nv(canon, v):
    """比對用的值：數字 100／100.0 一樣、日期只看到日、倉別大寫。"""
    if v is None:
        return ""
    t = str(v).strip()
    if canon in ("order_qty", "file_ship", "ship"):
        n = _num(t)
        return "" if n is None else (str(int(n)) if n == int(n) else str(n))
    if canon == "date":
        return t[:10]
    if canon == "wh":
        return t.upper()
    return t


_KEPT = ("酷澎檔案仍是舊值", "整合表出貨數量已變")    # 人工調整過、匯入不覆蓋：檔案的新值記在 old_value


def _canon_import(sys, l):
    """一筆匯入紀錄 → (欄位, 檔案上的新值, 改之前的值)；不是要算的欄位回 None。sys：c＝酷澎訂單管理、m＝商品主檔。"""
    f, note = l["field"], l.get("note") or ""
    kept = note.startswith(_KEPT)
    old = "" if kept else l["old_value"]
    if sys == "m":
        if f == "qty_coupang":
            return "order_qty", l["new_value"], old
        if f == "qty_file_ship":
            return "file_ship", l["new_value"], old
        if f == "delivery_date_file":
            return "date", l["new_value"], old
        if f == "warehouse":
            return "wh", l["new_value"], old
        if f == "qty_ship" and "已沒有這個品項" in note:
            return "gone", "gone", l["old_value"]
        if f == "missing_in_file":
            return "gone", "back", ""
        return None
    if f == "qty_coupang":
        return "order_qty", l["new_value"], old
    if f == "delivery_date":
        return "date", (l["old_value"] if kept else l["new_value"]), old
    if f == "warehouse":
        return "wh", (l["old_value"] if kept else l["new_value"]), old
    if f == "qty_ship":
        if note.startswith("隨整合表出貨數量連動"):
            return "file_ship", l["new_value"], old
        if note.startswith("整合表出貨數量已變"):
            return "file_ship", l["old_value"], ""
        if note.startswith("酷澎後台已移除"):
            return "gone", "gone", l["old_value"]
        return None                                         # 隨下單數量連動：下單數量那筆已經算了
    if f == "removed_from_coupang":
        return "gone", "back", ""
    return None


def _new_event(key, po, source, sys, when, operator):
    return {"key": key, "po": po, "source": source, "sys": sys, "when": when, "operator": operator, "reason": "", "note": "",
            "kinds": set(), "qty_dir": set(), "_qty": collections.defaultdict(list), "_date": [], "_wh": [], "_gone": 0, "_back": 0,
            "_delta": {}, "delta": 0.0}


def _add_qty(e, canon, sku, old, new, price):
    e["kinds"].add("qty"); e["_qty"][canon].append((old, new))
    a, b = _num(old), _num(new)
    if a is not None and b is not None and a != b:
        e["qty_dir"].add("down" if b < a else "up")
        # 同一個品項同一次：下單數量優先，其次訂單彙總表出貨數量、出貨數量（不重複算）
        rank = {"order_qty": 0, "file_ship": 1, "ship": 2}[canon]
        cur = e["_delta"].get(sku)
        if cur is None or rank < cur[0]:
            e["_delta"][sku] = (rank, (b - a) * (price or 0))


def _finish_events(ev):
    for e in ev.values():
        e["qty_dir"] = ("mixed" if len(e["qty_dir"]) > 1 else next(iter(e["qty_dir"]), ""))
        parts = []
        for canon in ("order_qty", "file_ship", "ship"):
            pairs = e["_qty"].get(canon) or []
            if pairs:
                ex = "、".join(f"{_short(a)}→{_short(b)}" for a, b in pairs[:3]) + ("…" if len(pairs) > 3 else "")
                parts.append(f"{CANON_LABEL[canon]} {len(pairs)} 項{'（' + DIR_LABEL[e['qty_dir']] + '）' if e['qty_dir'] else ''}：{ex}")
        if e["_date"]:
            pairs = sorted({(a, b) for a, b in e["_date"]})
            parts.append(f"到貨日 {'、'.join(f'{_short(a)}→{_short(b)}' for a, b in pairs[:2])}（{len(e['_date'])} 項）")
        if e["_wh"]:
            pairs = sorted({(a or "—", b) for a, b in e["_wh"]})
            parts.append(f"倉別 {'、'.join(f'{a}→{b}' for a, b in pairs[:2])}（{len(e['_wh'])} 項）")
        if e["_gone"]:
            parts.append(f"品項移除 {e['_gone']} 項")
        if e["_back"]:
            parts.append(f"品項重新出現 {e['_back']} 項")
        e["summary"] = "；".join(parts)
        e["kind_labels"] = "、".join(KIND_LABEL[k] for k in ("qty", "date", "wh", "gone") if k in e["kinds"])
        e["delta"] = round(e["delta"] + sum(v for _r, v in e["_delta"].values()), 2)
        for k in ("_qty", "_date", "_wh", "_gone", "_back", "_delta"):
            e.pop(k)
    return ev


def _import_events(logs, prices):
    """兩個系統的匯入紀錄 → 酷澎修改的改單事件（同一個改動只算一次，見檔頭）。另外回原始下單數量 {(po, sku): 值}。"""
    obs = collections.defaultdict(list)
    for sys, l in logs:
        c = _canon_import(sys, l)
        if c:
            canon, val, old = c
            obs[(l["po_number"], l["sku_id"], canon)].append((l["changed_at"] or "", sys, _nv(canon, val), _nv(canon, old), l))
    ev, first_qty = {}, {}
    for (po, sku, canon), pts in obs.items():
        pts.sort(key=lambda p: (p[0], p[1]))
        prev, gone_q = None, None
        for t, sys, v, old, l in pts:
            if prev is not None and v == prev:
                continue                                   # 另一個系統後來才看到同一個值：不重複算
            if v == "back" and prev is None:
                continue                                   # 沒看到被拿掉就說重新出現：沒有東西可以接回來
            before = prev if prev is not None else old
            if canon == "order_qty" and (po, sku) not in first_qty and before != "":
                first_qty[(po, sku)] = _num(before)
            key = f"import|{sys}|{po}|{t[:16]}"
            e = ev.setdefault(key, _new_event(key, po, "import", sys, t, l["operator"]))
            price = prices.get((po, sku))
            if canon in ("order_qty", "file_ship"):
                _add_qty(e, canon, sku, before, v, price)
            elif canon == "date":
                e["kinds"].add("date"); e["_date"].append((before, v))
            elif canon == "wh":
                e["kinds"].add("wh"); e["_wh"].append((before, v))
            elif v == "gone":
                e["kinds"].add("gone"); e["_gone"] += 1
                q = gone_q = _num(old)
                if q:
                    e["delta"] -= q * (price or 0)
            else:
                # 品項重新出現：出貨從 0 回到被拿掉前的數量，算改數量（上修），金額差跟被拿掉那次抵掉（Jerry 2026-10-19 CI 抓到：
                # 拿掉、重新出現跨了一分鐘，重新出現那次變成「改了什麼」空白、金額差 0 的改單）
                e["_back"] += 1; e["kinds"].add("qty"); e["qty_dir"].add("up")
                if gone_q and sku not in e["_delta"]:
                    e["_delta"][sku] = (1, gone_q * (price or 0))
            prev = v
    return ev, first_qty


def _manual_events(mlogs, clogs, prices):
    """我方修改：商品主檔（有原因）、酷澎訂單管理（沒有原因，要補）。人改的時間看到秒。"""
    ev = {}
    for sys, l in [("m", x) for x in mlogs] + [("c", x) for x in clogs]:
        stamp = l["changed_at"] or ""
        reason = (l.get("reason") or "") if sys == "m" else ""
        key = f"manual|{sys}|{l['po_number']}|{stamp[:19]}" + (f"|{reason}" if reason else "")
        e = ev.setdefault(key, _new_event(key, l["po_number"], "manual", sys, stamp, l["operator"]))
        if not e["reason"]:
            e["reason"] = reason
        note = l.get("note") or ""
        if note and note not in e["note"]:
            e["note"] = f"{e['note']}；{note}" if e["note"] else note
        if sys == "m" and note.startswith("勇信"):         # 瑪氏出貨的勇信缺貨也記在商品主檔的歷程；跟首頁「最近的動作」一樣分開標（Jerry 2026-10-19）
            e["label"] = "瑪氏出貨"
        f = l["field"]
        if f == "qty_ship":
            _add_qty(e, "ship", l["sku_id"], _nv("ship", l["old_value"]), _nv("ship", l["new_value"]), prices.get((l["po_number"], l["sku_id"])))
        elif f == "delivery_date":
            e["kinds"].add("date"); e["_date"].append((_nv("date", l["old_value"]), _nv("date", l["new_value"])))
        elif f == "warehouse":
            e["kinds"].add("wh"); e["_wh"].append((_nv("wh", l["old_value"]), _nv("wh", l["new_value"])))
    return ev


def _event_row(e, p):
    """給 API／Excel 用的一列。"""
    return {"key": e["key"], "when": e["when"], "po_number": e["po"], "month": p["month"], "lines": "／".join(sorted(p["lines"])),
            "source": e["source"], "source_label": "酷澎修改" if e["source"] == "import" else "我方修改",
            "system": e.get("label") or SYS_LABEL[e["sys"]], "operator": e["operator"], "kinds": sorted(e["kinds"]), "kind_labels": e["kind_labels"],
            "qty_dir": e["qty_dir"], "summary": e["summary"], "delta": e["delta"],
            "reason": e["final_reason"], "note": e["note"]}


def _include_test():
    return request.args.get("include_test") in ("1", "true")


def _item_rows(conn):
    """兩個系統的品項（同一個 PO＋SKU 兩邊都有就兩筆）。"""
    out = []
    for r in _rows(conn.execute("SELECT po_number, sku_id, line, delivery_date, qty_coupang, qty_ship, qty_file_ship, unit_price, "
                                "missing_in_file AS gone, first_seen_at, updated_at FROM mst_orders")):
        r["sys"] = "m"; r["parent_po"] = ""; out.append(r)
    for r in _rows(conn.execute("SELECT po_number, sku_id, line, delivery_date, qty_coupang, qty_ship, unit_price, parent_po, "
                                "removed_from_coupang AS gone, first_seen_at, updated_at FROM orders")):
        r["sys"] = "c"; r["qty_file_ship"] = None; out.append(r)
    return out


def _first_ship(conn, keep):
    """每個品項第一次看到的出貨數量：兩個系統出貨數量最早那筆紀錄「改之前」的值（沒紀錄的品項不在裡面＝從來沒變過）。"""
    logs = []
    for sys_, sql in (("m", "SELECT po_number, sku_id, field, old_value, new_value, note, changed_at FROM mst_logs "
                            "WHERE po_number != '' AND field IN ('qty_file_ship', 'qty_ship')"),
                      ("c", "SELECT po_number, sku_id, field, old_value, new_value, note, changed_at FROM edit_logs "
                            "WHERE po_number != '' AND field = 'qty_ship'")):
        logs += [(l["changed_at"] or "", sys_, l) for l in _rows(conn.execute(sql)) if l["po_number"] in keep]
    out = {}
    for t, _s, l in sorted(logs, key=lambda x: (x[0], x[1])):
        k = (l["po_number"], l["sku_id"])
        if k in out:
            continue
        before = l["new_value"] if (l["note"] or "").startswith(_KEPT) else l["old_value"]   # 人工調整過的那種：現在的值記在 new_value
        n = _num(before)
        if n is not None:
            out[k] = n
    return out


def _keep_pos(conn, month_from, month_to, line_group):
    """範圍內的 PO：{po: {month, lines, skus, items, dates, delivery_date}}，另回每個品項的列和單價。"""
    cfg = _line_groups()
    pos, by_item, prices = {}, collections.defaultdict(list), {}
    for r in _item_rows(conn):
        p = pos.setdefault(r["po_number"], {"dates": set(), "lines": set(), "skus": set()})
        p["dates"].add(r["delivery_date"] or ""); p["lines"].add(_group_of(r["line"], cfg)); p["skus"].add(r["sku_id"])
        by_item[(r["po_number"], r["sku_id"])].append(r)
        if r["unit_price"] is not None:
            prices.setdefault((r["po_number"], r["sku_id"]), _num(r["unit_price"]) or 0)
    span = None if month_from is None else set(_month_span(month_from, month_to))     # None＝全部期間
    keep = {}
    for po, p in pos.items():
        m = _po_month(p["dates"])
        if span is not None and m != UNDATED and m not in span:
            continue
        if line_group and line_group not in p["lines"]:
            continue
        p["month"] = m; p["items"] = len(p["skus"]); p["delivery_date"] = min((d for d in p["dates"] if d), default="")
        keep[po] = p
    return keep, by_item, prices


def _collect(conn, month_from, month_to, line_group, include_test=False):
    """範圍內的 PO（keep）與系統抓到的變動（events，兩個系統合起來算的改單事件）。
    回傳的第三個值是被排除的「測試」次數（include_test=True 時為 0）。"""
    keep, by_item, prices = _keep_pos(conn, month_from, month_to, line_group)
    if not keep:
        return keep, {}, 0
    restore = tuple(f"{x}%" for x in _RESTORE_PREFIXES)
    mlogs = _rows(conn.execute(
        """SELECT po_number, sku_id, source, field, old_value, new_value, operator, note, reason, changed_at FROM mst_logs
           WHERE po_number != '' AND (
                 (source = 'import' AND field IN ('qty_coupang', 'qty_file_ship', 'delivery_date_file', 'warehouse', 'missing_in_file'))
              OR (source = 'import' AND field = 'qty_ship' AND note LIKE ?)
              OR (source = 'manual' AND field IN ('qty_ship', 'delivery_date') AND note NOT LIKE ? AND note NOT LIKE ?))""",
        ("%已沒有這個品項%",) + restore))
    clogs = _rows(conn.execute(
        """SELECT po_number, sku_id, source, field, old_value, new_value, operator, note, changed_at FROM edit_logs
           WHERE po_number != '' AND (
                 (source = 'import' AND field IN ('qty_coupang', 'delivery_date', 'warehouse', 'qty_ship', 'removed_from_coupang'))
              OR (source = 'manual' AND field IN ('qty_ship', 'delivery_date', 'warehouse') AND note NOT LIKE ?))""", ("恢復%",)))
    mlogs = [l for l in mlogs if l["po_number"] in keep]; clogs = [l for l in clogs if l["po_number"] in keep]
    imp, first_qty = _import_events([("m", l) for l in mlogs if l["source"] == "import"] + [("c", l) for l in clogs if l["source"] == "import"], prices)
    man = _manual_events([l for l in mlogs if l["source"] == "manual"], [l for l in clogs if l["source"] == "manual"], prices)
    events = _finish_events({**imp, **man})
    for e in events.values():
        e["final_reason"] = OLD_REASON.get(e["reason"], e["reason"])      # 商品主檔改單時選的，只在明細帶著；統計不再列原因
    # 原始下單、第一次看到的出貨、目前出貨（數量）：每個品項看兩邊
    first_ship = _first_ship(conn, keep)
    amt = {}
    for po, p in keep.items():
        for sku in p["skus"]:
            rows = by_item[(po, sku)]
            orig = first_qty.get((po, sku))
            if orig is None:
                first = min(rows, key=lambda r: r["first_seen_at"] or "9")
                orig = _num(first["qty_coupang"]) or 0
            last = max(rows, key=lambda r: r["updated_at"] or "")
            cur = 0 if last["gone"] else _num(last["qty_ship"] if last["qty_ship"] is not None else (last["qty_file_ship"] if last["qty_file_ship"] is not None else last["qty_coupang"])) or 0
            fs = first_ship.get((po, sku))
            amt[(po, sku)] = {"price": prices.get((po, sku)) or 0, "orig": orig, "first": cur if fs is None else fs, "cur": cur,
                              "gone": bool(last["gone"]), "parent": next((r["parent_po"] for r in rows if r.get("parent_po")), ""), "moved": 0.0}
    # 酷澎拆單：原單拿掉的品項在新單裡 → 新單那項接回原單算（見檔頭）
    for (po, sku), a in amt.items():
        par = amt.get((a["parent"], sku)) if a["parent"] and a["parent"] != po else None
        if par is None or not par["gone"]:
            continue
        par["cur"] = a["cur"]; par["moved"] = a["orig"] * a["price"]
        a["orig"] = a["first"] = a["cur"] = 0
    for po, p in keep.items():
        p["orig_amount"] = p["cur_amount"] = p["pre_amount"] = p["post_amount"] = p["moved_amount"] = 0.0
        for sku in p["skus"]:
            a = amt[(po, sku)]
            p["orig_amount"] += a["orig"] * a["price"]; p["cur_amount"] += a["cur"] * a["price"]
            p["pre_amount"] += (a["first"] - a["orig"]) * a["price"]; p["post_amount"] += (a["cur"] - a["first"]) * a["price"]
            p["moved_amount"] += a["moved"]
    if include_test:
        return keep, events, 0
    tests = [k for k, e in events.items() if e["final_reason"] == TEST_REASON]
    for k in tests:
        events.pop(k)
    return keep, events, len(tests)


def _reason_cat(stage, kinds, reason):
    """一個主因歸哪個類別（主因表照類別分組）：先看這輪勾到的類別裡有沒有，沒有再照清單順序找。
    像「產能不及」同時在第二階段下修和拆單的清單裡，勾了下修就算下修。"""
    if "qty" in kinds and reason in rounds.REASONS_QTY.get(stage, []):
        return "qty"
    if "date" in kinds and reason in rounds.REASONS_DATE:
        return "date"
    if "split" in kinds and reason in rounds.REASONS_SPLIT:
        return "split"
    for cat, lst in (("qty", sum(rounds.REASONS_QTY.values(), [])), ("date", rounds.REASONS_DATE), ("split", rounds.REASONS_SPLIT)):
        if reason in lst:
            return cat
    return "other"


def _round_rows(conn, keep):
    """範圍內 PO 的修改紀錄（沒刪的），一輪一列，最新的在上面。"""
    rows = [rounds._row(r) for r in _rows(conn.execute(
        "SELECT id, po_number, ok_at, stage, kinds, reason, note, operator FROM po_rounds WHERE deleted_at = '' ORDER BY ok_at DESC, id DESC"))
        if r["po_number"] in keep]
    for r in rows:
        p = keep[r["po_number"]]
        r["month"] = p["month"]; r["lines"] = "／".join(sorted(p["lines"]))
        r["stage_short"] = rounds.STAGE_SHORT.get(r["stage"], r["stage"])
        r["cats"] = {x: _reason_cat(r["stage"], r["kinds"], x) for x in r["reasons"]}
    return rows


def _pending_rows(conn, keep):
    """改過但還沒記 OK 的 PO（在範圍內的），到貨日近的在前面。"""
    out = []
    for po in rounds.pending_pos(conn):
        p = keep.get(po)
        if p is None:
            continue
        info = rounds.pending_info(conn, po)
        out.append({"po_number": po, "month": p["month"], "lines": "／".join(sorted(p["lines"])), "delivery_date": p["delivery_date"],
                    "last_ok": info["last_ok"], "changes": info["changes"], "more": info["more"]})
    out.sort(key=lambda r: (r["delivery_date"] or "9", r["po_number"]))
    return out


def _build_stats(conn, month_from, month_to, line_group, include_test=False):
    keep, events, n_test = _collect(conn, month_from, month_to, line_group, include_test)
    base = {"month_from": month_from, "month_to": month_to, "test_events": n_test, "since": rounds.SINCE,
            "stage_labels": {k: rounds.STAGE_LABEL[k] for k in STAGE_KEYS}, "kind_labels": dict(rounds.KINDS), "cat_labels": CAT_LABEL}
    if not keep:
        return {**base, "months": [], "total": _finish_month(_empty_month("合計")), "reasons": [], "stages": _stage_rows(_empty_month("合計")),
                "top_pos": [], "pending": []}

    per_po = {po: {"rounds": 0, "events": 0, "stages": collections.Counter(), "kinds": collections.Counter(), "reasons": collections.Counter(),
                   "list": []} for po in keep}
    for e in events.values():
        per_po[e["po"]]["events"] += 1
    for r in reversed(_round_rows(conn, keep)):          # 舊的在前，top_pos 的「各輪」照時間排
        s = per_po[r["po_number"]]
        s["rounds"] += 1; s["stages"][r["stage"]] += 1
        for k in r["kinds"]:
            s["kinds"][k] += 1
        for x in r["reasons"]:
            s["reasons"][(r["cats"][x], x)] += 1
        s["list"].append({"stage": r["stage"], "stage_short": r["stage_short"], "kinds": r["kinds"], "kind_labels": r["kind_labels"],
                          "reasons": r["reasons"], "ok_at": r["ok_at"][:10]})
    pending = _pending_rows(conn, keep)
    pend_by_month = collections.Counter(r["month"] for r in pending)

    months = {}
    order = (_month_span(month_from, month_to) if month_from else sorted({p["month"] for p in keep.values()} - {UNDATED})) + [UNDATED]

    def add(m, s, p):
        m["pos"] += 1; m["items"] += p["items"]; m["_skus"].update(p["skus"])
        for k in ("orig_amount", "cur_amount", "pre_amount", "post_amount", "moved_amount"):
            m[k] += p[k]
        if s["rounds"]:
            m["changed_pos"] += 1
        m["rounds"] += s["rounds"]; m["events"] += s["events"]
        for k, v in s["stages"].items():
            m["stages"][k] = m["stages"].get(k, 0) + v
        for k, v in s["kinds"].items():
            m["kinds"][k] = m["kinds"].get(k, 0) + v
        for k, v in s["reasons"].items():
            m["reasons"][k] = m["reasons"].get(k, 0) + v
    total = _empty_month("合計")
    for po, p in keep.items():
        add(months.setdefault(p["month"], _empty_month(p["month"])), per_po[po], p)
        add(total, per_po[po], p)
    # 各階段 × 做了什麼（只算合計，畫面上那張表不分月）
    stage_kinds = {k: collections.Counter() for k in STAGE_KEYS}
    for s in per_po.values():
        for r in s["list"]:
            for k in r["kinds"]:
                stage_kinds[r["stage"]][k] += 1
    total["_stage_kinds"] = stage_kinds
    for m in months.values():
        m["pending"] = pend_by_month.get(m["month"], 0)
    total["pending"] = len(pending)
    out_months = [_finish_month(months[m]) for m in order if m in months]
    total = _finish_month(total)

    reasons = sorted(({"cat": cat, "cat_label": CAT_LABEL[cat], "reason": x, "n": n} for (cat, x), n in total["_reasons"].items()),
                     key=lambda r: (["qty", "date", "split", "other"].index(r["cat"]), -r["n"], r["reason"]))
    total.pop("_reasons")
    for m in out_months:
        m.pop("_reasons", None)

    top = sorted(((po, keep[po], per_po[po]) for po in keep if per_po[po]["rounds"]), key=lambda x: (-x[2]["rounds"], x[0]))[:15]
    top_pos = []
    for po, p, s in top:
        seen = []
        for r in s["list"]:
            seen += [x for x in r["reasons"] if x not in seen]
        top_pos.append({"po_number": po, "month": p["month"], "lines": "／".join(sorted(p["lines"])), "items": p["items"],
                        "rounds": s["rounds"], "list": s["list"], "reasons": "、".join(seen)})
    return {**base, "months": out_months, "total": total, "reasons": reasons, "stages": _stage_rows(total), "top_pos": top_pos, "pending": pending}


def _stage_rows(total):
    sk = total.get("_stage_kinds") or {}
    rows = [{"key": k, "label": rounds.STAGE_LABEL[k], "short": rounds.STAGE_SHORT[k], "rounds": total["stages"].get(k, 0),
             "kinds": {x: sk.get(k, {}).get(x, 0) for x in KIND_KEYS}} for k in STAGE_KEYS]
    total.pop("_stage_kinds", None)
    return rows


def _event_rows(keep, events):
    rows = [_event_row(e, keep[e["po"]]) for e in events.values()]
    rows.sort(key=lambda r: (r["when"], r["po_number"]), reverse=True)
    return rows


def _empty_month(label):
    return {"month": label, "pos": 0, "items": 0, "_skus": set(), "changed_pos": 0, "rounds": 0, "events": 0, "pending": 0,
            "stages": {}, "kinds": {}, "reasons": {},
            "orig_amount": 0.0, "cur_amount": 0.0, "pre_amount": 0.0, "post_amount": 0.0, "moved_amount": 0.0}


def _finish_month(m):
    m["skus"] = len(m.pop("_skus"))
    m["changed_pct"] = round(m["changed_pos"] * 100 / m["pos"]) if m["pos"] else 0
    m["per_changed_po"] = round(m["rounds"] / m["changed_pos"], 1) if m["changed_pos"] else 0
    m["stages"] = {k: m["stages"].get(k, 0) for k in STAGE_KEYS}
    m["kinds"] = {k: m["kinds"].get(k, 0) for k in KIND_KEYS}
    m["_reasons"] = m.pop("reasons")
    for k in ("orig_amount", "cur_amount", "pre_amount", "post_amount", "moved_amount"):
        m[k] = round(m[k], 2)
    m["short_amount"] = round(m["cur_amount"] - m["orig_amount"], 2)
    m["short_pct"] = round((m["orig_amount"] - m["cur_amount"]) * 100 / m["orig_amount"], 1) if m["orig_amount"] else 0
    return m


@master_bp.route("/api/master/stats")
def api_stats():
    month_from, month_to, months = _stats_range()
    if not months:
        return jsonify({"error": "月份格式錯誤，請使用 2026-09 格式。"}), 400
    line = norm_text(request.args.get("line"))
    conn = get_conn()
    try:
        data = _build_stats(conn, month_from, month_to, line, _include_test())
    finally:
        conn.close()
    data["line"] = line
    return jsonify(data)


@master_bp.route("/api/master/stats/rounds")
def api_stats_rounds():
    """修改紀錄明細：一輪一列，畫面上點修改次數時撈的。篩選在前端做。"""
    month_from, month_to, months = _stats_range()
    if not months:
        return jsonify({"error": "月份格式錯誤，請使用 2026-09 格式。"}), 400
    line = norm_text(request.args.get("line"))
    conn = get_conn()
    try:
        keep, _by, _pr = _keep_pos(conn, month_from, month_to, line)
        rows = _round_rows(conn, keep)
    finally:
        conn.close()
    return jsonify({"rounds": rows, "month_from": month_from, "month_to": month_to})


@master_bp.route("/api/master/stats/events")
def api_stats_events():
    """系統抓到的變動：一次一列，畫面上點灰字那欄時撈的。篩選在前端做（一段期間頂多幾百幾千筆）。"""
    month_from, month_to, months = _stats_range()
    if not months:
        return jsonify({"error": "月份格式錯誤，請使用 2026-09 格式。"}), 400
    line = norm_text(request.args.get("line"))
    conn = get_conn()
    try:
        keep, events, n_test = _collect(conn, month_from, month_to, line, _include_test())
    finally:
        conn.close()
    return jsonify({"events": _event_rows(keep, events), "month_from": month_from, "month_to": month_to, "test_events": n_test})


@master_bp.route("/api/master/stats/export")
def api_stats_export():
    month_from, month_to, months = _stats_range()
    if not months:
        return jsonify({"error": "月份格式錯誤，請使用 2026-09 格式。"}), 400
    line = norm_text(request.args.get("line"))
    minutes = norm_int(request.args.get("minutes")) or 0
    conn = get_conn()
    try:
        d = _build_stats(conn, month_from, month_to, line, _include_test())
        keep, events, n_test = _collect(conn, month_from, month_to, line, _include_test())
        rrows = _round_rows(conn, keep)
    finally:
        conn.close()
    from openpyxl.styles import Alignment, Font, PatternFill
    bold = Font(bold=True); head = PatternFill("solid", fgColor="DBEAFE"); tot = PatternFill("solid", fgColor="F1F5F9")
    grey = Font(color="6B7280")
    wb = openpyxl.Workbook()
    kl = dict(rounds.KINDS)

    ws = wb.active; ws.title = "每月統計"
    hdr = ["月份", "PO 張數", "品項數", "不同 SKU 數", "曾改單的 PO", "曾改單比例", "修改次數（OK）", "第一階段", "第二階段", "第三階段",
           "有下修", "有改交期／倉別", "有拆單", "系統抓到的變動次數", "還沒記 OK 的 PO",
           "原始下單金額", "目前出貨金額", "少出金額", "少出比例", "其中進系統前就少了", "其中進系統後才改", "拆單接回原單的金額"]
    if minutes:
        hdr.append(f"估計工時（每次 {minutes} 分鐘）")
    ws.append(hdr)
    for m in d["months"] + [d["total"]]:
        row = [m["month"], m["pos"], m["items"], m["skus"], m["changed_pos"], f"{m['changed_pct']}%", m["rounds"],
               m["stages"]["1"], m["stages"]["2"], m["stages"]["3"], m["kinds"]["qty"], m["kinds"]["date"], m["kinds"]["split"],
               m["events"], m["pending"], m["orig_amount"], m["cur_amount"], m["short_amount"], f"{m['short_pct']}%",
               m["pre_amount"], m["post_amount"], m["moved_amount"]]
        if minutes:
            row.append(round(m["rounds"] * minutes / 60, 1))
        ws.append(row)
    _style(ws, hdr, bold, head, tot, total_row=True)
    for r in ws.iter_rows(min_row=2, min_col=14, max_col=14):
        r[0].font = grey

    ws2 = wb.create_sheet("主因")
    hdr2 = ["類別", "主因", "次數"]
    ws2.append(hdr2)
    for r in d["reasons"]:
        ws2.append([r["cat_label"], r["reason"], r["n"]])
    ws2.append(["合計", "", sum(r["n"] for r in d["reasons"])])
    _style(ws2, hdr2, bold, head, tot, total_row=True)
    ws2.column_dimensions["B"].width = 28

    ws3 = wb.create_sheet("各階段做了什麼")
    hdr3 = ["階段", "修改次數（OK）", "有下修", "有改交期／倉別", "有拆單"]
    ws3.append(hdr3)
    for st in d["stages"]:
        ws3.append([st["label"], st["rounds"], st["kinds"]["qty"], st["kinds"]["date"], st["kinds"]["split"]])
    t = d["total"]
    ws3.append(["合計", t["rounds"], t["kinds"]["qty"], t["kinds"]["date"], t["kinds"]["split"]])
    _style(ws3, hdr3, bold, head, tot, total_row=True)

    ws4 = wb.create_sheet("修改次數最多的 PO")
    hdr4 = ["PO 單號", "月份", "線別", "品項數", "修改次數（OK）", "各輪", "主因"]
    ws4.append(hdr4)
    for p in d["top_pos"]:
        ws4.append([p["po_number"], p["month"], p["lines"], p["items"], p["rounds"],
                    "；".join(f"{r['ok_at']} {r['stage_short']} {'、'.join(r['kind_labels'])}" for r in p["list"]), p["reasons"]])
    _style(ws4, hdr4, bold, head, tot, total_row=False)
    ws4.column_dimensions["F"].width = 60; ws4.column_dimensions["G"].width = 40

    ws5 = wb.create_sheet("修改紀錄明細")
    hdr5 = ["OK 時間", "月份", "PO 單號", "線別", "階段", "做了什麼", "主因", "說明", "記的人"]
    ws5.append(hdr5)
    for r in rrows:
        ws5.append([r["ok_at"], r["month"], r["po_number"], r["lines"], r["stage_label"], "、".join(r["kind_labels"]), r["reason"], r["note"], r["operator"]])
    _style(ws5, hdr5, bold, head, tot, total_row=False)
    ws5.column_dimensions["A"].width = 20; ws5.column_dimensions["G"].width = 36; ws5.column_dimensions["H"].width = 40

    ws6 = wb.create_sheet("還沒記 OK 的 PO")
    hdr6 = ["PO 單號", "月份", "線別", "到貨日", "上一個 OK", "之後的改動"]
    ws6.append(hdr6)
    for r in d["pending"]:
        ws6.append([r["po_number"], r["month"], r["lines"], r["delivery_date"], r["last_ok"] or "—",
                    "；".join(r["changes"]) + (f"…另 {r['more']} 項" if r["more"] else "")])
    _style(ws6, hdr6, bold, head, tot, total_row=False)
    ws6.column_dimensions["F"].width = 80

    ws7 = wb.create_sheet("系統抓到的變動")
    hdr7 = ["時間", "月份", "PO 單號", "線別", "操作人員", "來源", "記在哪個系統", "改了什麼", "修改內容", "金額差", "說明"]
    ws7.append(hdr7)
    for r in _event_rows(keep, events):
        ws7.append([r["when"], r["month"], r["po_number"], r["lines"], r["operator"], r["source_label"], r["system"],
                    r["kind_labels"], r["summary"], r["delta"], r["note"]])
    _style(ws7, hdr7, bold, head, tot, total_row=False)
    ws7.column_dimensions["I"].width = 60; ws7.column_dimensions["A"].width = 20

    ws8 = wb.create_sheet("怎麼算的")
    for line_text in [
        ("範圍：全部期間" if month_from is None else f"範圍：{month_from} ～ {month_to}") + (f"，線別 {line}" if line else "，全部線別"),
        "PO 歸屬月份：依目前的到貨日（多個日期取最早那天）。未排到貨日的另列「未排日期」。",
        f"修改次數（OK）：以酷澎訂單管理系統 PO 視窗的「修改紀錄」為準，一輪處理完按一次 OK 算一次。修改紀錄自 {rounds.SINCE} 起才有，之前的月份修改次數為 0。",
        "曾改單的 PO：至少有一筆修改紀錄的 PO。平均每張改幾次＝修改次數 ÷ 曾改單的 PO。",
        "第一／二／三階段：記 OK 時選的「這一輪在哪個階段結束」（收單確認／已排單確認／出貨確認）。",
        "有下修／有改交期／倉別／有拆單：記 OK 時勾的「這一輪做了什麼」，一輪可以勾好幾項，所以三項加起來會大於修改次數。",
        "主因：記 OK 時勾的主因（照小真的清單，依類別分組），一輪可以勾好幾個，所以主因加總會大於修改次數；沒用到的主因不列。",
        "還沒記 OK 的 PO：上一個 OK 之後酷澎或我們有改過（下單或出貨數量變少、到貨日或倉別變了、拆出新單），還沒按「這一輪處理完了（OK）」。",
        "系統抓到的變動次數（灰字）：系統從兩個系統的匯入紀錄和人工修改抓出來的每一次變動，同一張 PO、同一來源、同一時間算一次；好幾次變動可能其實是同一輪處理，"
        "所以不當修改次數，留著對照用，每個月都有。改單時勾了「這是測試」的不算" + ("" if _include_test() else f"（這段期間有 {n_test} 次）") + "。",
        "金額：數量 × 下單單價（一張單有箱、盒、包，數量直接相加沒有意思）。原始下單＝兩個系統最早看到的下單數量；目前出貨＝最近更新那邊的出貨數量，品項移除算 0。",
        "少出拆兩段：進系統前就少了＝第一次匯入時出貨數量就比下單數量少（在 Kate 系統先改的，沒有修改紀錄）；進系統後才改＝之後的變動。",
        "酷澎拆單：原單拿掉的品項在拆出來的新單裡（原訂單欄），新單那項接回原單算，不算少出、也不算新的下單。",
        "估計工時：修改次數（OK）× 每次幾分鐘 ÷ 60。每次幾分鐘是人填的，不是系統算的。",
        "資料從系統開始用（2026 年 9 月中）才有，之前 Excel 時代的紀錄不在裡面。",
    ]:
        ws8.append([line_text])
    ws8.column_dimensions["A"].width = 110
    for r in ws8.iter_rows():
        r[0].alignment = Alignment(wrap_text=True, vertical="top")

    label = "全部期間" if month_from is None else month_from if month_from == month_to else f"{month_from}～{month_to}"
    return _xlsx_response(wb, f"改單統計_{label}.xlsx")


def _style(ws, hdr, bold, head, tot, total_row):
    from openpyxl.utils import get_column_letter
    for c in ws[1]:
        c.font = bold; c.fill = head
    if total_row and ws.max_row > 1:
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = tot
    for i, h in enumerate(hdr, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(h)) * 2 + 4))
    ws.freeze_panes = "A2"


__all__ = ["api_stats", "api_stats_rounds", "api_stats_events", "api_stats_export", "_build_stats", "_collect", "_event_rows"]
