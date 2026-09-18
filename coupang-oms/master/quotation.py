"""匯出專案報價檔（一個交貨日一個分頁，A～R 欄）。"""
from .common import *  # noqa: F401,F403 — 共用工具、Flask、db、openpyxl 都從這裡來


@master_bp.route("/api/master/export/daily")
def api_export_daily():
    """專案報價檔：一個交貨日一個分頁（0904交貨、0903交貨…新的在前），A～R 每欄照
    原本公式的來源填。匯出的就是畫面上篩出來的（線別／日期／PO／倉別／品牌／關鍵字），
    什麼都沒篩才是整月；帶 month_to 可以一次匯好幾個月。"""
    month = norm_text(request.args.get("month")) or _this_month()
    month_to = norm_text(request.args.get("month_to")) or month
    if not _valid_month(month) or not _valid_month(month_to):
        return jsonify({"error": "月份格式要像 2026-09。"}), 400
    months = _month_span(month, month_to)       # 可以跨月：7 月～9 月就是三個月的交貨日全部進來
    filters = _read_filters(request.args)
    cfg = _line_groups()
    conn = get_conn()
    try:
        orders = [o for o in _month_orders(conn, months, cfg) if _keep(o, filters)]
    finally:
        conn.close()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    head = ["SKU ID", "永豐料號", "國條", "品類", "品牌", "品名", "下單數量(酷澎單位)", "出貨數量", "箱入數",
            "出貨數量(箱)", "單價(含稅)", "酷澎下單價(含稅)", "箱單價(含稅)", "總計(含稅)", "備註",
            "驗收完成請打勾", "簽單完成請打勾", "交貨日"]
    batch = filters.get("batch")
    if batch:
        head.append("本次變動")   # 只匯某一次匯入時多一欄：新增／出貨數量 20→0／交貨日 9/4→9/8
    yellow = PatternFill("solid", fgColor="FFFF00"); head_fill = PatternFill("solid", fgColor="F8CBAD"); bold = Font(bold=True)
    changed_fill = PatternFill("solid", fgColor="FFF3C4")
    by_date = collections.OrderedDict()
    for o in orders:
        by_date.setdefault(o["delivery_date"] or "", collections.OrderedDict()).setdefault(o["po_number"], []).append(o)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for date in sorted(by_date.keys(), key=lambda d: d or "0000", reverse=True):
        pos = by_date[date]
        if date:
            d = _dt.date.fromisoformat(date); title = f"{d.month:02d}{d.day:02d}交貨"; label = f"{d.month}/{d.day}交貨"
        else:
            title, label = "未排日期", ""
        ws = wb.create_sheet(title[:31]); ws.append(head)
        for c in ws[1]:
            c.font = bold; c.fill = head_fill
        first = True
        for po, rows in pos.items():
            if not first:
                ws.append([None] * len(head))
                for c in ws[ws.max_row]:
                    c.fill = yellow
            first = False
            po_first_row = ws.max_row + 1
            for i, o in enumerate(rows):
                box = o["box_size"]; cp = o["unit_price"]
                r = ws.max_row + 1
                # J／M／N 三欄寫成活的 Excel 公式，不是算好的數字（Chloe 2026-09-17：業務會在檔裡改
                # 數量、單價，公式才會跟著動）。J＝出貨數量÷箱入數、M＝酷澎下單價×箱入數、N＝箱單價×箱數。
                # 箱入數或單價空白時留空，不要跳 #DIV/0! 或算出 0 元誤導。
                cases_f = f'=IF(OR(I{r}="",I{r}=0),"",H{r}/I{r})'
                box_price_f = f'=IF(OR(L{r}="",I{r}=""),"",L{r}*I{r})'
                total_f = f'=IF(OR(M{r}="",J{r}=""),"",M{r}*J{r})'
                po_cell = (f"{po}_{label}({o['warehouse']})" if o["warehouse"] else f"{po}_{label}") if i == 0 else None
                cells = [o["sku_id"], o.get("yf_sku_master") or o["yf_sku"], o["barcode"], o.get("category") or "",
                         o.get("brand_master") or o["brand"], o["product_name"], o["qty_coupang"], o["qty_ship"],
                         box, cases_f, o.get("cost_price"), cp, box_price_f, total_f, o["export_note"], "", "", po_cell]
                if batch:
                    is_new = o.get("first_batch_id") == batch
                    cells.append("新增" if is_new else (o.get("last_batch_changes") or "有變"))
                ws.append(cells)
                if batch and not is_new:
                    for c in ws[ws.max_row]:
                        c.fill = changed_fill
            # 同一張 PO 的「交貨日」欄合併成一格（Chloe 2026-09-17：同一張 PO 請跨欄位），
            # 值在第一列、垂直置中，看起來就是一張單一塊。
            if ws.max_row > po_first_row:
                ws.merge_cells(start_row=po_first_row, start_column=18, end_row=ws.max_row, end_column=18)
            ws.cell(row=po_first_row, column=18).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for i, w in enumerate([16, 15, 15, 8, 14, 44, 10, 9, 8, 11, 10, 12, 11, 12, 18, 8, 8, 30, 36], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=3):
            for c in row:
                c.number_format = "@"
        # 數字欄不設格式（用 Excel 的「通用」）：之前設過「#,##0.##」，整數會被顯示成「100.」還帶逗號，
        # Chloe 2026-09-18 抓到。通用格式整數就是整數、小數照實際位數顯示。
    if not wb.sheetnames:
        ws = wb.create_sheet("無資料"); ws.append(["目前的篩選條件下沒有任何訂單"])
    # 檔名照同事原本的檔就叫「專案報價檔」，分頁是 0904交貨、0903交貨…新的在前
    return _xlsx_response(wb, "專案報價檔.xlsx")

