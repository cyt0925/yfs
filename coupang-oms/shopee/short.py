"""蝦皮特選缺貨維護與通知（需求第八章）。

流程（Jerry 2026-10-12 問過主管）：
- 寶僑：採購單打完、核單後，採購寄信通知缺貨 → 人工進系統填缺貨 → 第一次下修通知蝦皮；
  之後拋訂單時又缺 → 人工再填 → 第二次通知。
- 紙潔：Teams 通知缺貨 → 人工填 → 通知蝦皮（一次）。
- 瑪氏：勇信配送明細表（跟酷澎瑪氏的 PDF 長一樣）系統自動核對（yx.py）→ 寫成第一次 → 通知蝦皮。只有一次、不能人工填。
- 竹運出貨：竹運出貨頁產拋單檔時判斷出來的缺貨，直接帶進來當第一次（不用再填）。
- 通知蝦皮用 LINE：照 Alice 的 LINE 機器人訊息（Jerry 2026-10-12 貼的範例）產文字，一張 PR、一種處理（刪單／下修）一則，
  另外附 Excel 明細（下修明細、刪單明細、最終可出貨量）。

規則：
- 下修數量＝原始訂購量－最終可出貨量，算累計的（需求 8.4）。原始訂購量＝我們第一次下修前的蝦皮數量
  （蝦皮自己改過數量、差異也確認過的話，以確認後的為準，不把蝦皮自己的變動算成我們的下修）；還沒下修就是目前的蝦皮數量。
- 一次一筆：填的是「這次缺多少」，系統算這次之後的可出貨量；整筆缺＝0＝刪單。
- 還沒通知的那一次可以改、可以刪；通知過就鎖住，之後再缺要開下一次（第一次的紀錄不能蓋掉，需求 8.2）。
"""
import collections

import openpyxl
from openpyxl.styles import Font, PatternFill

from .common import *  # noqa: F401,F403
from .orders import order_no

MANUAL_FULFILS = ("竹運採購進貨", "供應商直送")
SHIP_FULFIL = "竹運出貨"
MAX_ROUNDS = {"寶僑": 2, "紙潔": 1, "瑪氏": 1}
REASONS = ("庫存不足", "效期不足")
LINES_HERE = ("寶僑", "紙潔")                     # 可以人工填缺貨的線別
LIST_LINES = LINES_HERE + ("瑪氏",)              # 清單上看得到的（瑪氏只有供應商直送，缺貨由勇信核對寫進來）
MARKET = {"寶僑": "P&G", "紙潔": "紙潔", "瑪氏": "瑪氏"}          # 訊息裡的「賣場」（Alice 範例寶僑寫 P&G）
ROUND_TXT = {1: "首次通知", 2: "第二次通知"}


def _rounds(conn, ids):
    if not ids:
        return {}
    marks = ",".join("?" * len(ids))
    out = collections.defaultdict(list)
    for r in _rows(conn.execute(f"SELECT s.*, n.created_at AS notify_at, n.created_by AS notify_by FROM shp_short s "
                                f"LEFT JOIN shp_notify n ON n.id = s.notify_id WHERE s.order_id IN ({marks}) ORDER BY s.round", list(ids))):
        out[r["order_id"]].append(r)
    return out


def short_item(it, rounds, done):
    """一個品項目前的缺貨狀況。rounds＝shp_short 那幾筆；done＝竹運出貨拋單檔快照（沒有就 None）。"""
    cur = it["qty"] or 0
    rs = [{"round": r["round"], "source": r["source"], "before": r["qty_before"], "after": r["qty_after"], "short": r["qty_before"] - r["qty_after"],
           "reasons": [x for x in (r["reasons"] or "").split("、") if x], "note": r["note"] or "", "operator": r["operator"], "updated_at": r["updated_at"],
           "notified": bool(r["notify_id"]), "notify_at": r.get("notify_at") or "", "saved": True} for r in rounds]
    if it["fulfil"] == SHIP_FULFIL and done and not rs:
        short = max((done["qty_order"] or 0) - (done["qty_ship"] or 0), 0)
        if short:                                   # 竹運出貨判斷出的缺貨：還沒通知前先列著，通知時才存成第一次
            rs.append({"round": 1, "source": SHIP_FULFIL, "before": done["qty_order"] or 0, "after": done["qty_ship"] or 0, "short": short,
                       "reasons": [x for x in (done["short_reasons"] or "").split("、") if x], "note": f"竹運出貨拋單檔 {done['filename']}",
                       "operator": done["created_by"], "updated_at": done["created_at"], "notified": False, "notify_at": "", "saved": False})
    orig = rs[0]["before"] if rs else cur
    final = rs[-1]["after"] if rs else cur
    last = rs[-1] if rs else None
    max_r = MAX_ROUNDS.get(it["line"], 1)
    manual = it["fulfil"] in MANUAL_FULFILS and it["line"] in LINES_HERE
    if not manual:
        edit = None
    elif last and not last["notified"]:
        edit = last["round"]                        # 還沒通知：改這一次
    elif len(rs) < max_r:
        edit = len(rs) + 1                          # 通知過了：開下一次
    else:
        edit = None
    return {"order_id": it["id"], "supplier_sku_id": it["supplier_sku_id"], "shopee_sku_id": it["shopee_sku_id"], "name": it["sku_name"],
            "pr_id": it["pr_id"] or "", "po_id": it["po_id"] or "",
            "unit": it["unit_name"] or UNIT_CODES.get(it["unit_code"] or "", ""), "inbound_ids": it["inbound_ids"] or "", "orig": orig, "qty": cur,
            "final": final, "cut": max(orig - final, 0), "rounds": rs, "pending": any(not r["notified"] for r in rs),
            "notified_round": max((r["round"] for r in rs if r["notified"]), default=0), "edit_round": edit, "max_rounds": max_r,
            "prev_final": (rs[edit - 2]["after"] if edit >= 2 else cur) if edit else None}


def build_short(conn, d1="", d2="", line="", fulfil="", order_nos=None):
    where, params = ["removed = 0", "line IN (?, ?, ?)", "(line != ? OR fulfil = ?)"], list(LIST_LINES) + ["瑪氏", "供應商直送"]
    fulfils = (fulfil,) if fulfil in MANUAL_FULFILS + (SHIP_FULFIL,) else MANUAL_FULFILS + (SHIP_FULFIL,)
    where.append(f"fulfil IN ({','.join('?' * len(fulfils))})"); params += list(fulfils)
    if d1:
        where.append("expected_date >= ?"); params.append(d1)
    if d2:
        where.append("expected_date <= ?"); params.append(d2)
    if line in LIST_LINES:
        where.append("line = ?"); params.append(line)
    items = _rows(conn.execute(f"SELECT * FROM shp_orders WHERE {' AND '.join(where)} ORDER BY expected_date, warehouse, id", params))
    if order_nos is not None:
        want = set(order_nos)
        items = [i for i in items if order_no(i) in want]
    ids = [i["id"] for i in items]
    rounds = _rounds(conn, ids)
    from .ship import ship_status_map
    done = ship_status_map(conn, [i["id"] for i in items if i["fulfil"] == SHIP_FULFIL])
    groups = collections.OrderedDict()
    for it in items:
        no = order_no(it)
        g = groups.setdefault(no, {"order_no": no, "pr_id": it["pr_id"], "po_id": it["po_id"], "line": it["line"], "fulfil": it["fulfil"],
                                   "warehouse": it["warehouse"], "warehouse_name": wh_name(it["warehouse"]), "expected_date": it["expected_date"], "items": []})
        g["items"].append(short_item(it, rounds.get(it["id"], []), done.get(it["id"])))
    out = []
    for g in groups.values():
        its = g["items"]
        rs = [r for p in its for r in p["rounds"]]
        g.update(orig=sum(p["orig"] for p in its), final=sum(p["final"] for p in its), cut=sum(p["cut"] for p in its),
                 deleted=sum(1 for p in its if p["orig"] and not p["final"]), short_items=sum(1 for p in its if p["rounds"]),
                 pending=sum(1 for p in its if p["pending"]), notified_round=max((p["notified_round"] for p in its), default=0),
                 status=notify_status(rs))
        out.append(g)
    return out


def notify_status(rs):
    if any(not r["notified"] for r in rs):
        return "待通知"
    if rs:
        return f"已通知（第 {max(r['round'] for r in rs)} 次）"
    return "未通知"


def notify_status_map(conn, ids):
    """訂單總覽用：品項 → 待通知／已通知第幾次。只看有存的下修紀錄（竹運出貨還沒通知的在總覽算「待通知」交給呼叫端判斷）。"""
    out = {}
    for oid, rs in _rounds(conn, ids).items():
        out[oid] = {"pending": any(not r["notify_id"] for r in rs), "round": max((r["round"] for r in rs if r["notify_id"]), default=0),
                    "final": rs[-1]["qty_after"], "orig": rs[0]["qty_before"]}
    return out


@shopee_bp.route("/api/shopee/short/plan")
def api_short_plan():
    a = request.args
    conn = get_conn()
    try:
        orders = build_short(conn, a.get("from", ""), a.get("to", ""), a.get("line", ""), a.get("fulfil", ""))
        notes = _rows(conn.execute("SELECT id, orders, items, created_by, created_at FROM shp_notify ORDER BY id DESC LIMIT 20"))
    finally:
        conn.close()
    return jsonify({"orders": orders, "notifies": notes, "reasons": list(REASONS)})


@shopee_bp.route("/api/shopee/short/save", methods=["POST"])
def api_short_save():
    """填這次缺多少（需求 8.2、8.3）。short＝0 而且這次還沒通知＝取消這次。"""
    p = request.get_json(silent=True) or {}
    try:
        oid, short = int(p.get("order_id")), int(str(p.get("short", "")).replace(",", "") or 0)
    except (TypeError, ValueError):
        return jsonify({"error": "缺貨數量請填整數。"}), 400
    reasons = [r for r in (p.get("reasons") or []) if r in REASONS]
    note = str(p.get("note") or "").strip()
    operator = _operator()
    conn = get_conn()
    try:
        it = _row(conn.execute("SELECT * FROM shp_orders WHERE id = ?", (oid,)))
        if it is None:
            return jsonify({"error": "找不到這個品項，請重新整理頁面。"}), 404
        if it["fulfil"] not in MANUAL_FULFILS or it["line"] not in LINES_HERE:
            return jsonify({"error": "竹運出貨的缺貨請在竹運出貨頁修改出貨量；瑪氏的缺貨由勇信效期表核對。"}), 400
        cur = short_item(it, _rounds(conn, [oid]).get(oid, []), None)
        rnd = cur["edit_round"]
        if not rnd:
            return jsonify({"error": f"{it['line']}最多下修 {cur['max_rounds']} 次，已經都通知過了。"}), 400
        prev = cur["prev_final"]
        if short < 0 or short > prev:
            return jsonify({"error": f"缺貨數量要在 0 到 {prev} 之間（目前可出貨量 {prev}）。"}), 400
        old = _row(conn.execute("SELECT * FROM shp_short WHERE order_id = ? AND round = ?", (oid, rnd)))
        no = order_no(it)
        if short == 0:
            if old:
                conn.execute("DELETE FROM shp_short WHERE id = ?", (old["id"],))
                _log(conn, it["line"], no, it["shopee_sku_id"], "", "shopee_short", f"第 {rnd} 次下修", f"可出貨 {old['qty_after']}", "取消（不缺）", operator,
                     "manual", it["supplier_sku_id"])
                conn.commit()
            return jsonify({"ok": True, "cleared": True})
        if not reasons:
            return jsonify({"error": "請勾選缺貨原因（庫存不足、效期不足，可複選）。"}), 400
        after, stamp = prev - short, now()
        if old:
            conn.execute("UPDATE shp_short SET qty_before = ?, qty_after = ?, reasons = ?, note = ?, operator = ?, updated_at = ? WHERE id = ?",
                         (prev, after, "、".join(reasons), note, operator, stamp, old["id"]))
        else:
            conn.execute("INSERT INTO shp_short (order_id, round, source, qty_before, qty_after, reasons, note, operator, updated_at) VALUES (?,?,?,?,?,?,?,?,?)",
                         (oid, rnd, "人工", prev, after, "、".join(reasons), note, operator, stamp))
        _log(conn, it["line"], no, it["shopee_sku_id"], "", "shopee_short", f"第 {rnd} 次下修", f"可出貨 {prev}", f"可出貨 {after}（缺 {short}）", operator,
             "manual", f"{it['supplier_sku_id']}；{'、'.join(reasons)}{('；' + note) if note else ''}", "、".join(reasons))
        conn.commit()
        return jsonify({"ok": True, "round": rnd, "after": after})
    finally:
        conn.close()


# ── 通知蝦皮（LINE 文字＋Excel 明細） ───────────────────────────────────────
def line_messages(groups):
    """照 Alice 的 LINE 機器人訊息：一張 PR、一種處理（刪單／下修）、同一次通知一則。刪單在前、下修在後。"""
    buckets = collections.OrderedDict()
    for g in groups:
        for p in g["items"]:
            if not p["pending"]:
                continue
            kind = "刪單" if not p["final"] else "下修"
            key = (p["pr_id"] or g["order_no"], 0 if kind == "刪單" else 1, p["rounds"][-1]["round"])
            buckets.setdefault(key, {"g": g, "kind": kind, "items": []})["items"].append(p)
    msgs = []
    for (pr, _k, rnd), b in sorted(buckets.items(), key=lambda kv: kv[0]):
        g, kind, its = b["g"], b["kind"], b["items"]
        out = [f"【蝦皮缺貨{kind}通知｜{ROUND_TXT.get(rnd, f'第{rnd}次通知')}】", f"賣場：{MARKET.get(g['line'], g['line'])}", f"PR ID：{pr}",
               f"預計到貨日：{g['expected_date'] or ''}", f"到貨倉別：{g['warehouse'] or ''}", "", f"{kind}SKU（{len(its)}項）："]
        for i, p in enumerate(its, 1):
            out += [f"{i}. PO ID：{p['po_id'] or '尚未產生'}", f"蝦皮SKU ID：{p['shopee_sku_id']}", f"{p['supplier_sku_id']} {p['name']}"]
            if kind == "下修":
                out += [f"訂購{p['orig']}／可出貨{p['final']}／缺貨{p['cut']}", f"處理方式：下修至{p['final']}"]
            out.append("")
        out.append("請協助處理刪單，完成後請回覆，謝謝。" if kind == "刪單" else "請協助下修訂單數量，完成後請回覆，謝謝。")
        msgs.append("\n".join(out))
    return msgs


def notify_rows(groups):
    out = []
    for g in groups:
        for p in g["items"]:
            if not p["pending"]:
                continue
            r = p["rounds"][-1]
            out.append({"order_no": g["order_no"], "date": g["expected_date"], "wh": g["warehouse_name"], "line": g["line"], "fulfil": g["fulfil"],
                        "inbound_id": p["inbound_ids"], "sku": p["supplier_sku_id"], "shopee_sku": p["shopee_sku_id"], "name": p["name"], "unit": p["unit"],
                        "orig": p["orig"], "final": p["final"], "cut": p["cut"], "reasons": "、".join(r["reasons"]), "round": r["round"],
                        "operator": r["operator"], "updated_at": r["updated_at"], "note": r["note"]})
    return out


def notify_xlsx(rows):
    """需求 8.4／8.5：下修明細、刪單明細、最終可出貨量清單。"""
    head = ["訂單編號", "到貨日", "倉別", "線別", "履約方式", "入庫單號", "廠商料號", "蝦皮商品編號", "品名", "單位", "原始訂購量", "最終可出貨量", "下修數量",
            "缺貨原因", "下修次數", "更新人員", "更新時間", "說明"]
    keys = ["order_no", "date", "wh", "line", "fulfil", "inbound_id", "sku", "shopee_sku", "name", "unit", "orig", "final", "cut", "reasons", "round",
            "operator", "updated_at", "note"]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, pick in (("下修明細", lambda r: r["final"] > 0), ("刪單明細", lambda r: r["final"] == 0), ("最終可出貨量清單", lambda r: True)):
        ws = wb.create_sheet(title)
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FBE4DC")
        for r in rows:
            if pick(r):
                ws.append([str(r[k])[:16].replace("T", " ") if k == "updated_at" else r[k] for k in keys])
        for col, w in zip("ABCDEFGHIJKLMNOPQR", (22, 11, 8, 6, 12, 20, 18, 26, 46, 6, 10, 11, 9, 16, 8, 9, 16, 24)):
            ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@shopee_bp.route("/api/shopee/short/notify", methods=["POST"])
def api_short_notify():
    """勾選的單裡「待通知」的品項 → 一則 LINE 通知文字＋Excel 明細；那幾次下修標成已通知（鎖住）。"""
    p = request.get_json(silent=True) or {}
    nos = [str(x) for x in (p.get("order_nos") or []) if str(x).strip()]
    if not nos:
        return jsonify({"error": "請先勾選要通知的訂單。"}), 400
    operator, stamp = _operator(), now()
    conn = get_conn()
    try:
        groups = [g for g in build_short(conn, order_nos=nos) if g["pending"]]
        if not groups:
            return jsonify({"error": "勾選的訂單沒有待通知的缺貨。"}), 400
        msgs = line_messages(groups)
        text = "\n\n".join(msgs)
        rows = notify_rows(groups)
        cur = conn.execute("INSERT INTO shp_notify (text, messages_json, rows_json, orders, items, created_by, created_at) VALUES (?,?,?,?,?,?,?)",
                           (text, json.dumps(msgs, ensure_ascii=False), json.dumps(rows, ensure_ascii=False), len(groups), len(rows), operator, stamp))
        nid = cur.lastrowid
        for g in groups:
            for x in g["items"]:
                for r in x["rounds"]:
                    if r["notified"]:
                        continue
                    if r["saved"]:
                        conn.execute("UPDATE shp_short SET notify_id = ? WHERE order_id = ? AND round = ?", (nid, x["order_id"], r["round"]))
                    else:                                   # 竹運出貨帶進來的：通知時存成第一次
                        conn.execute("INSERT INTO shp_short (order_id, round, source, qty_before, qty_after, reasons, note, operator, updated_at, notify_id) "
                                     "VALUES (?,?,?,?,?,?,?,?,?,?)", (x["order_id"], r["round"], r["source"], r["before"], r["after"], "、".join(r["reasons"]),
                                                                    r["note"], r["operator"], r["updated_at"], nid))
            _log(conn, g["line"], g["order_no"], "", "", "shopee_notify", "缺貨通知蝦皮", "待通知", f"已通知（{g['pending']} 品）", operator, "manual", f"通知 #{nid}")
        conn.commit()
        return jsonify({"ok": True, "id": nid, "messages": msgs, "items": len(rows)})
    finally:
        conn.close()


@shopee_bp.route("/api/shopee/short/notify/<int:nid>")
def api_short_notify_get(nid):
    conn = get_conn()
    try:
        n = _row(conn.execute("SELECT * FROM shp_notify WHERE id = ?", (nid,)))
    finally:
        conn.close()
    if n is None:
        return jsonify({"error": "找不到這次通知。"}), 404
    return jsonify({**{k: n[k] for k in ("id", "text", "orders", "items", "created_by", "created_at")}, "messages": json.loads(n["messages_json"] or "[]")})


@shopee_bp.route("/api/shopee/short/notify/<int:nid>/xlsx")
def api_short_notify_xlsx(nid):
    conn = get_conn()
    try:
        n = _row(conn.execute("SELECT * FROM shp_notify WHERE id = ?", (nid,)))
    finally:
        conn.close()
    if n is None:
        return jsonify({"error": "找不到這次通知。"}), 404
    d = (n["created_at"] or "")[:10].replace("-", "")
    return send_file(io.BytesIO(notify_xlsx(json.loads(n["rows_json"] or "[]"))), as_attachment=True,
                     download_name=f"蝦皮特選缺貨下修通知_{d}_{nid}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


__all__ = ["short_item", "build_short", "notify_status", "notify_status_map", "line_messages", "notify_rows", "notify_xlsx", "api_short_plan",
           "api_short_save", "api_short_notify", "api_short_notify_get", "api_short_notify_xlsx"]
