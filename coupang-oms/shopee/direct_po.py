"""蝦皮特選直送採購單：給採購轉給原廠的檔（Alice 2026-10-21，範本「蝦皮特選直送採購單」在 \\\\10.37.90.10\\電子商務\\營運管理\\系統優化\\蝦皮特選\\給採購的檔案）。

竹運採購進貨本來就有「採購下採明細」給採購；供應商直送之前只有 EIP 採購單，採購還要自己做一份給原廠，這裡照 Alice 的範本直接產。
一張 EIP 對一份（寶僑箱單位、小單位 EIP 分兩個檔，這個也跟著分兩份；紙潔一份）。回填 EIP 採購單號後才能產（D1 要填單號），跟瑪氏採購單一樣。
格子照範本（Alice 的欄位填寫規則）：
  A1 進貨通路名「蝦皮特選」    C1「採購單號」 D1 EIP 採購單號    G1「下單」 H1 下載當天    G2「到庫」 H2 到貨日
  第 3 列表頭，第 4 列起一品一列：採購單ID（有入庫單號填入庫單號，沒有填 PR 單號）、供應商料號（蝦皮 Supplier SKU ID）、SKU名稱、
  數量（蝦皮數量）、單位、箱入數（價格本）、轉換箱數（＝數量÷箱入數，公式）、備註；最後一列 Total。範本留 15 列，品項多就往下加。
  B21「請安排 M/D 到貨」  A22「蝦皮特選店-倉名」  C22 那個倉的進倉說明（WH_NOTES，可在資料與設定改）  C23 固定的進倉注意事項。
採購單ID 用產檔當下的入庫單號（蝦皮入庫單常常是 EIP 產出之後才上傳），其他照 EIP 存下來的內容。
"""
import datetime as _dt
import zipfile

import openpyxl

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .common import *  # noqa: F401,F403

DIRECT = "供應商直送"
MARS_LINE = "瑪氏"
CHANNEL = "蝦皮特選"
# 各倉進倉說明（Alice 2026-10-21 給的，照蝦皮進倉規範）；C22 用。可在資料與設定「蝦皮倉資料」改，存在 mst_meta shopee_wh 的 notes
WH_NOTES = {
    "TWK": "進倉時間: 星期二至星期四 07:30–12:00；星期日不收貨。最晚進倉時間 11:00，最晚離倉時間 12:00。\n物流中心地址: 桃園市楊梅區梅獅路二段 888 號（5F）\n連絡電話: 0975213830（威獅進貨組）\n注意事項:活動日進倉時間需注意",
    "TWW": "進倉時間: 星期一至星期六 07:30–12:00；星期日不收貨。最晚進倉時間 11:00，最晚離倉時間 12:00。\n物流中心地址: 桃園市楊梅區梅獅路二段 888 號（3F／5F／7F；請依指定樓層進倉）\n連絡電話: 0975213830（威獅進貨組）\n注意事項:活動日進倉時間需注意",
    "TWH": "進倉時間: 星期一至星期五 14:00–18:00；星期六、日不收貨。最晚進倉時間 17:00，最晚離倉時間 18:00。\n物流中心地址: 桃園市楊梅區和平路 576 號（物流共和國・楊梅 2A，08–13碼頭）\n連絡電話: 03-262-6560，分機 3151",
    "TWA": "進倉時間: 星期一至星期六 07:00–12:00、13:00–16:00；星期日不收貨。最晚進倉時間 11:00、15:00，最晚離倉時間 12:00、16:00。\n物流中心地址: 桃園市觀音區寶倉街 108 號 C 棟 1 樓\n連絡電話: 03-2800-500-5924（進貨組）",
    "TWR": "進倉時間: 星期一至星期六 07:00–12:00；星期日不收貨，國定假日是否收貨不固定（需提前確認）。最晚進倉時間 11:00，最晚離倉時間 12:00。\n物流中心地址: 桃園市楊梅區梅獅路 395 巷 168 號\n連絡電話: 0917-179560（張先生）",
    "TWG": "進倉時間: 星期一至星期五 14:00–18:00；星期日不收貨。最晚進倉時間 17:00，最晚離倉時間 18:00。\n物流中心地址: 桃園市中壢區高鐵南路四段 888 號\n連絡電話: 0968233806（高鐵南進貨組）",
    "TWT": "進倉時間: 星期一至星期六 08:00–11:30、13:00–14:30；星期日不收貨，國定假日是否收貨不固定（需提前確認）。最晚進倉時間 11:30、14:30，須於該時段前完成下貨和驗收。\n物流中心地址: 台南市安南區新吉一路 118 號\n連絡電話: 06-600-1450，分機 5871（安南一進貨組）",
    "TWC": "進倉時間: 星期一至星期六 08:00–11:00、13:00–14:00；星期日不收貨，國定假日是否收貨不固定（需提前確認）。最晚進倉時間 11:00、14:00，須於該時段前完成下貨和驗收。\n物流中心地址: 台南市安定區新吉里安新二路 568 號\n連絡電話: 0983-655-699（曾小姐）",
    "TWX": "進倉時間: 星期一至星期六 08:00–11:30、13:00–14:30；星期日不收貨，國定假日是否收貨不固定（需提前確認）。最晚進倉時間 11:30、14:30，須於該時段前完成下貨和驗收。\n物流中心地址: 台南市安定區安新二路 688 號\n連絡電話: 06-600-1450，分機 5883（安南三進貨組）",
}
NOTICE = ("●同批進貨商品效期須一致\n●不同商品須分開包裝，並將「入庫通知單」與商品捆綁在一起\n●在同一箱品分多箱/板入庫時，請依上列方法，將「入庫通知單」分別與商品放置一起\n"
          "●同批進貨商品效期須一致，若有不同效期，請以不同入庫單號進倉\n●同批進貨商品若有不同包裝，不可同一時間進貨\n●現場收貨，以現場實際到貨點收數量為主\n●請依照預計入庫日期及預約時間準時送達倉庫")
HEAD = ["", "採購單ID", "供應商料號", "SKU名稱", "數量", "單位", "箱入數", "轉換箱數", "備註"]
MIN_ROWS = 15


def wh_settings(conn):
    from .mars_po import wh_settings as _ws      # 放在函式裡，避免跟 mars_po 互相 import
    return _ws(conn)


def wh_notes(conn):
    """{倉代碼: 進倉說明}：資料與設定「蝦皮倉資料」改過的蓋預設。"""
    return {code: w.get("notes") or WH_NOTES.get(code, "") for code, w in wh_settings(conn).items()}


def direct_rows(conn, e):
    """一張 EIP（shp_eip 那列）→ 一品一列：採購單ID 用現在的入庫單號（沒有就 PR）、數量照 EIP 存的 parts（沒有 parts 的舊 EIP 照合併列）。"""
    rows = json.loads(e["rows_json"] or "[]")
    ids = [r["order_id"] for r in _rows(conn.execute("SELECT order_id FROM shp_eip_items WHERE eip_id = ?", (e["id"],)))]
    orders = {}
    if ids:
        marks = ",".join("?" * len(ids))
        orders = {o["id"]: o for o in _rows(conn.execute(f"SELECT * FROM shp_orders WHERE id IN ({marks})", ids))}
    out = []
    for r in rows:
        parts = r.get("parts") or []
        if not parts:                                   # 2026-10-21 以前產的 EIP 沒存每張單的量：整列當一品
            o = next((orders[i] for i in ids if orders.get(i) and orders[i]["base_code"] == r["code"] and (orders[i]["unit_name"] or "") == (r.get("in_unit") or "")), None)
            parts = [{"order_id": o["id"] if o else None, "use_qty": r.get("in_qty", r["qty"])}]
        for p in parts:
            o = orders.get(p.get("order_id")) or {}
            out.append({"po_id": (o.get("inbound_ids") or "").split(",")[0] or o.get("pr_id") or o.get("po_id") or "",
                        "sku": o.get("supplier_sku_id") or (r["code"] + "_" + (o.get("unit_code") or "")),
                        "name": o.get("sku_name") or r["name"], "qty": p["use_qty"], "unit": r.get("in_unit") or o.get("unit_name") or "",
                        "box": r.get("conv") or 1})
    return out


def direct_name(e):
    d = e["expected_date"] or ""
    return f"{e['eip_po'] or ''}蝦皮特選直送採購單_{d[5:7]}{d[8:10]}_{wh_name(e['warehouse'])}" + (f"_{e['unit_group']}" if e["unit_group"] else "") + ".xlsx"


def direct_xlsx(e, rows, note, today=None):
    """照 Alice 的範本畫。rows 來自 direct_rows。"""
    today = today or _dt.date.today()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "直送採購單"
    hei = Font(name="微軟正黑體", size=20, bold=True); mid = Font(name="微軟正黑體", size=16); body = Font(name="微軟正黑體", size=11)
    thin = Side(style="thin"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center"); left = Alignment(horizontal="left", vertical="center")
    grey = PatternFill("solid", fgColor="D9D9D9"); yellow = PatternFill("solid", fgColor="FFFF00")
    for col, w in zip("ABCDEFGHI", (5.6, 31.5, 24.5, 55.5, 13.9, 9.8, 10.7, 22.4, 8.6)):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:B2"); ws.merge_cells("C1:C2"); ws.merge_cells("D1:D2")
    ws["A1"] = CHANNEL; ws["A1"].font = hei; ws["A1"].alignment = center; ws["A1"].border = box
    ws["C1"] = "採購單號"; ws["C1"].font = hei; ws["C1"].alignment = center
    ws["D1"] = e["eip_po"] or ""; ws["D1"].font = Font(name="微軟正黑體", size=24, bold=True); ws["D1"].alignment = center
    ws["G1"] = "下單"; ws["G1"].font = mid; ws["G1"].alignment = center
    ws["H1"] = today; ws["H1"].number_format = "mm-dd-yy"; ws["H1"].font = mid; ws["H1"].alignment = center
    ws["G2"] = "到庫"; ws["G2"].font = mid; ws["G2"].alignment = center
    ws["H2"] = _dt.date.fromisoformat(e["expected_date"]) if e["expected_date"] else ""; ws["H2"].number_format = "mm-dd-yy"; ws["H2"].font = hei; ws["H2"].alignment = center
    ws.row_dimensions[1].height = 20.5; ws.row_dimensions[2].height = 26; ws.row_dimensions[3].height = 32
    for i, h in enumerate(HEAD, start=1):
        c = ws.cell(row=3, column=i, value=h); c.font = Font(name="新細明體", size=12, bold=True); c.fill = grey; c.border = box; c.alignment = center
    n = max(len(rows), MIN_ROWS)
    for i in range(n):
        r = 4 + i; row = rows[i] if i < len(rows) else None
        vals = [i + 1] + ([row["po_id"], row["sku"], row["name"], row["qty"], row["unit"], row["box"], f"=E{r}/G{r}", ""] if row else [""] * 8)
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v); c.font = body; c.border = box
            c.alignment = center if j in (5, 8) else left
            if j == 5 and row:
                c.fill = yellow
            if j == 8:
                c.number_format = "0.0"; c.fill = grey
        ws.row_dimensions[r].height = 16
    tr = 4 + n
    ws.cell(row=tr, column=4, value="Total").font = Font(name="微軟正黑體", size=12)
    ws.cell(row=tr, column=5, value=f"=SUM(E4:E{tr - 1})"); ws.cell(row=tr, column=8, value=f"=SUM(H4:H{tr - 1})")
    for j in range(1, 10):
        c = ws.cell(row=tr, column=j); c.border = box; c.alignment = center
        if j in (5, 8):
            c.font = Font(name="微軟正黑體", size=14, bold=True, color="FF0000")
    ws.row_dimensions[tr].height = 18.5
    d = e["expected_date"] or ""
    r1, r2, r3 = tr + 2, tr + 3, tr + 4
    ws.cell(row=r1, column=2, value=f"請安排{int(d[5:7])}/{int(d[8:10])} 到貨" if len(d) >= 10 else "請安排到貨").font = Font(name="Arial", size=12, bold=True)
    ws.row_dimensions[r1].height = 43
    ws.merge_cells(start_row=r2, start_column=1, end_row=r3, end_column=2); ws.merge_cells(start_row=r2, start_column=3, end_row=r2, end_column=9); ws.merge_cells(start_row=r3, start_column=3, end_row=r3, end_column=9)
    a = ws.cell(row=r2, column=1, value=f"{CHANNEL}店-{wh_name(e['warehouse'])}"); a.font = Font(name="細明體", size=14, bold=True); a.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c22 = ws.cell(row=r2, column=3, value=note); c22.font = Font(name="Arial", size=14, bold=True); c22.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c23 = ws.cell(row=r3, column=3, value=NOTICE); c23.font = Font(name="微軟正黑體", size=12); c23.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for rr in (r2, r3):
        for j in range(1, 10):
            ws.cell(row=rr, column=j).border = box
    ws.row_dimensions[r2].height = 24 + 18 * (note.count("\n") + 1); ws.row_dimensions[r3].height = 115
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def direct_bytes(conn, e, notes=None):
    notes = notes or wh_notes(conn)
    return direct_name(e), direct_xlsx(e, direct_rows(conn, e), notes.get(e["warehouse"], ""))


def _eligible(e):
    return e["fulfil"] == DIRECT and e["line"] != MARS_LINE


@shopee_bp.route("/api/shopee/eips/<int:eid>/direct_po")
def api_shopee_direct_po(eid):
    conn = get_conn()
    try:
        e = _row(conn.execute("SELECT * FROM shp_eip WHERE id = ?", (eid,)))
        if e is None or not _eligible(e):
            return jsonify({"error": "找不到這張供應商直送的 EIP 採購單，請重新整理頁面。"}), 404
        if not e["eip_po"]:
            return jsonify({"error": "請先回填 EIP 採購單號，才能產出直送採購單。"}), 400
        name, data = direct_bytes(conn, e)
    finally:
        conn.close()
    return send_file(io.BytesIO(data), as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@shopee_bp.route("/api/shopee/direct_po.zip")
def api_shopee_direct_po_zip():
    """到貨日區間內、已回填 EIP 採購單號的直送採購單（寶僑、紙潔）一次下載。"""
    where, params = ["fulfil = ?", "line != ?", "eip_po != ''"], [DIRECT, MARS_LINE]
    if request.args.get("from"):
        where.append("expected_date >= ?"); params.append(request.args["from"])
    if request.args.get("to"):
        where.append("expected_date <= ?"); params.append(request.args["to"])
    usql, up = upload_cond(request.args.get("up_from", ""), request.args.get("up_to", ""), "o.")
    if usql:
        where.append(f"id IN (SELECT i.eip_id FROM shp_eip_items i JOIN shp_orders o ON o.id = i.order_id WHERE {usql})"); params += up
    conn = get_conn()
    try:
        eips = _rows(conn.execute(f"SELECT * FROM shp_eip WHERE {' AND '.join(where)} ORDER BY expected_date, warehouse, id", params))
        if not eips:
            return jsonify({"error": "這段到貨日沒有已回填 EIP 採購單號的直送採購單。"}), 400
        notes = wh_notes(conn)
        files = [direct_bytes(conn, e, notes) for e in eips]
    finally:
        conn.close()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files:
            zf.writestr(name, data)
    buf.seek(0)
    ds = sorted({e["expected_date"] for e in eips})
    md = lambda d: f"{d[5:7]}{d[8:10]}"  # noqa: E731
    zname = f"蝦皮特選直送採購單_{md(ds[0])}" + (f"-{md(ds[-1])}" if ds[-1] != ds[0] else "") + ".zip"
    resp = send_file(buf, as_attachment=True, download_name=zname, mimetype="application/zip")
    resp.headers["X-Shopee-Files"] = str(len(files))
    return resp


__all__ = ["WH_NOTES", "wh_notes", "direct_rows", "direct_name", "direct_xlsx", "direct_bytes", "api_shopee_direct_po", "api_shopee_direct_po_zip"]
