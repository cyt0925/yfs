"""匯出訂單明細（Alice 2026-10-30）：清單勾一張或多張單，匯出一個 Excel。

Jerry 定的：一張工作表、一品一列，每列前面帶訂單資料（方便篩選），後面欄位照訂單視窗的品項表；有缺貨的那列整列黃底。
數字跟訂單視窗用同一份資料（side_data），竹運出貨的出貨量也照畫面上的到貨日／線別篩選算（庫存是照到貨日分的）。
"""
import datetime as _dt

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .common import *  # noqa: F401,F403
from .ship import build_ship
from .side import side_data

HEAD = ("訂單編號", "PR 單號", "線別", "倉別", "到貨日", "履約方式", "永豐料號", "品名", "入庫單號", "蝦皮數量", "單位", "原始訂購量",
        "轉換率", "下採數量", "EIP 數量", "EIP 單位", "出貨量", "缺貨", "缺貨原因", "效期", "調整原因", "缺貨通知", "狀態")
WIDTH = (20, 22, 7, 13, 11, 12, 22, 40, 22, 10, 6, 10, 8, 10, 10, 8, 9, 8, 14, 24, 16, 13, 12)
SHORT_FILL = PatternFill("solid", fgColor="FFF2A8")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")


def item_rows(d):
    """一張單 → 每品一列（dict：欄名 → 值），另外回這列有沒有缺貨。"""
    g = d["order"]
    by = lambda grp: {x["order_id"]: x for x in (grp or {}).get("items", [])}
    buy, ship, short = by(d["buy"]), by(d["ship"]), by(d["short"])
    out = []
    for i in g["items"]:
        p, z, k = buy.get(i["id"]), ship.get(i["id"]), short.get(i["id"])
        final = z["final"] if z else k["final"] if k else None
        cut = (z["short"] or 0) if z else (k["cut"] or 0) if k else None
        reasons = z["reasons"] if z and z["short"] else sorted({x for r in (k["rounds"] if k else []) for x in r["reasons"]})
        notice = ("待通知" if k["pending"] else f"已通知第 {k['notified_round']} 次" if k["notified_round"] else "") if k and k["rounds"] else ""
        state = "已刪除" if g.get("trash") else "已移除" if i["removed"] else "差異待確認" if i["diff_pending"] else ""
        row = {"訂單編號": g["order_no"], "PR 單號": g["pr_id"] if g["po_id"] else "", "線別": i["line"], "倉別": wh_name(i["warehouse"]),
               "到貨日": i["expected_date"], "履約方式": i["fulfil"], "永豐料號": i["supplier_sku_id"], "品名": i["sku_name"],
               "入庫單號": i["inbound_ids"] or "", "蝦皮數量": i["qty"], "單位": i["unit_name"] or i["unit_code"],
               "原始訂購量": i["qty_original"], "轉換率": p["conv"] if p else None, "下採數量": p["use_qty"] if p else None,
               "EIP 數量": p["out_qty"] if p else None, "EIP 單位": (p["out_unit"] or "") if p else "",
               "出貨量": final, "缺貨": cut, "缺貨原因": "、".join(reasons),
               "效期": "、".join(f"{a['exp']}（{a['qty']:,}）" for a in z["alloc"]) if z else "",
               "調整原因": (z["adj"]["reason"] or "") if z and z["adj"] else (p["adj"]["reason"] or "") if p and p.get("adj") else "",
               "缺貨通知": notice, "狀態": state}
        out.append((row, bool(cut) and not i["removed"]))
    return out


@shopee_bp.route("/api/shopee/orders/export", methods=["POST"])
def api_orders_export():
    p = request.get_json(silent=True) or {}
    nos = [str(x).strip() for x in (p.get("order_nos") or []) if str(x).strip()]
    if not nos:
        return jsonify({"error": "請先勾選要匯出的訂單。"}), 400
    conn = get_conn()
    try:
        ship_groups = None
        datas = []
        for no in nos:
            if ship_groups is None:                    # 竹運出貨的計畫照畫面的篩選算一次就好
                ship_groups = build_ship(conn, p.get("from", ""), p.get("to", ""), p.get("line", ""))[0]
            d = side_data(conn, no, p, ship_groups)
            if d is None:
                return jsonify({"error": f"找不到訂單 {no}，請重新整理頁面。"}), 404
            datas.append(d)
    finally:
        conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "訂單明細"
    ws.append(HEAD)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True), HEAD_FILL, Alignment(horizontal="center", vertical="center")
    n_items = n_short = 0
    for d in datas:
        for row, short in item_rows(d):
            ws.append([row[h] for h in HEAD])
            n_items += 1
            if short:
                n_short += 1
                for c in ws[ws.max_row]:
                    c.fill = SHORT_FILL
    for k, w in enumerate(WIDTH, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(k)].width = w
    for col in ("J", "L", "M", "N", "O", "Q", "R"):
        for c in ws[col][1:]:
            c.number_format = "#,##0.##"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = send_file(buf, as_attachment=True, download_name=f"蝦皮特選訂單明細_{_dt.date.today():%m%d}_{len(datas)}張.xlsx",
                  mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r.headers["X-Shopee-Items"], r.headers["X-Shopee-Short"] = str(n_items), str(n_short)
    return r


__all__ = ["api_orders_export", "item_rows"]
