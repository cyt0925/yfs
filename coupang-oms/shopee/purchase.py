"""蝦皮特選第二段：竹運採購進貨、供應商直送 → EIP 採購單＋採購下採明細（需求第六、七、十章）。

寶僑、紙潔 2026-10-08 做；瑪氏直送 2026-10-12 做（EIP 在這裡，瑪氏制式採購單在 mars_po.py）。
- 轉換率（箱入數）從各線別的價格本來：PG價格本整合.xlsx「價格本整合」、紙潔價格本整合.xlsx「紙品價格本」「潔品價格本」，
  都是「料號3＝料號＋單位代碼」配一個箱入數，跟蝦皮 Supplier SKU（1022248_CNN → 1022248C）對得起來。上傳時同一線別整份覆蓋。
- 寶僑（Jerry 2026-10-07 看範例檔定的，跟需求 10.1／10.2 對得起來）：轉換率 1 放「箱單位」檔、不是 1 放「小單位」檔；
  料號放國條、品名用蝦皮品名、備註「蝦皮特選」、第 10 欄（標題空白）放轉換率。
    竹運採購進貨：單位、數量照蝦皮原本的（PG竹運進貨範例）。
    供應商直送：小單位那檔單位寫「箱」、數量＝蝦皮數量÷轉換率（PG直出範例 216÷4＝54）。
- 紙潔（Jerry 2026-10-08：照需求 10.3）：一律換成箱，單位「箱」、數量＝蝦皮數量÷轉換率，一個檔；備註空白；
  「寄銷倉庫存」「前 30 天實銷」先留空（Jerry：不知道要填什麼）。
- 分檔（寶僑再分箱單位／小單位）：供應商直送一天一倉一個檔；竹運採購進貨整批一個檔（照 PG竹運進貨範例，跨天跨倉合在一起）。
  同一個料號同一個單位合成一列。
- 數量除不盡轉換率（需求 10.5）：不自動進位或捨去，列為異常，要人填修改後數量和原因才能產出；改的紀錄留原始數量、轉換率、人員、時間。
- 差異還沒確認的單、已回填 EIP 單號的單不能產（要重產先清掉單號）。
- 瑪氏（需求 10.4，Jerry 2026-10-12「照建議」）：只做供應商直送。單位、數量照蝦皮（不換箱），備註
  「MARS入倉20261013台灣瑪氏送蝦皮特選-觀音」，套瑪氏的 EIP 範本（mars_template.xls，跟酷澎瑪氏同一份）。
  商品資料用瑪氏出貨的商品總表（mst_mars_products，永豐料號＋單位），轉換率＝箱入數，除不盡一樣要改數量。
  一份瑪氏採購單對一張 EIP：依倉、到貨日、品類、單位、中標拆（需求 11.3），不出採購下採明細。
"""
import datetime as _dt
import os
import zipfile

import openpyxl
import xlrd
from openpyxl.styles import Font, PatternFill
from xlutils.copy import copy as xl_copy

from mars.common import CAT_FULL, cat_code

from .common import *  # noqa: F401,F403
from .orders import order_no

TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "purchase_templates", "pg_template.xls")
WIPE_ROWS = 249                       # 範本「產品採購單」原本有資料到第 249 列，全部清掉再寫
MARS_TEMPLATE = os.path.join(os.path.dirname(TEMPLATE), "mars_template.xls")   # 瑪氏 EIP 範本（工作表「產品領用單」）
MARS_WIPE_ROWS = 60
EIP_PO_RE = re.compile(r"^PO\d{9}$")  # EIP 採購單號：PO202610003
BUY_FULFILS = ("竹運採購進貨", "供應商直送")
FULFIL_SHORT = {"竹運採購進貨": "竹運進貨", "供應商直送": "直送"}
SUPPORTED = ("寶僑", "紙潔", "瑪氏")
MARS = "瑪氏"


# ── 價格本 ───────────────────────────────────────────────────────────────────
def parse_conv(data, filename=""):
    """價格本 → {料號3: {...}}。找表頭有「料號3」「箱入數」的工作表；名字有「價格本」的優先（紙潔檔另外有一頁
    「工作表4」是兩頁合起來的，不重複讀）。回傳 (rows, warnings)。"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"{filename}：無法開啟此檔案，請確認是 .xlsx。") from exc
    sheets = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        for i, row in enumerate(it):
            heads = [str(v).strip() if v is not None else "" for v in (row or ())]
            if "料號3" in heads and "箱入數" in heads:
                sheets.append((ws.title, heads, it))
                break
            if i >= 5:
                break
    if any("價格本" in t for t, _, _ in sheets):
        sheets = [x for x in sheets if "價格本" in x[0]]
    if not sheets:
        raise ValueError(f"{filename}：找不到有「料號3」「箱入數」欄位的工作表，請上傳寶僑或紙潔的價格本整合檔。")
    out, warns = {}, []
    for title, heads, it in sheets:
        idx = {h: i for i, h in enumerate(heads) if h}
        for row in it:
            row = list(row or ())
            get = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None  # noqa: E731
            code3 = str(get("料號3") or "").strip()
            if not code3 or code3 == "None":
                continue
            base, ucode = code3[:-1], code3[-1:].upper()
            box = get("箱入數")
            try:
                box = int(float(box)) if box not in (None, "") else None
            except (TypeError, ValueError):
                box = None
            rec = {"code3": base + ucode, "line": line_of(base), "base_code": base, "unit_code": ucode,
                   "unit_name": str(get("單位") or "").strip() or UNIT_CODES.get(ucode, ""), "box_qty": box if box and box > 0 else None,
                   "name": str(get("品名") or "").strip()}
            old = out.get(rec["code3"])
            if old and old["box_qty"] and rec["box_qty"] and old["box_qty"] != rec["box_qty"]:
                warns.append(f"{rec['code3']} 在不同頁的箱入數不同（{old['box_qty']}／{rec['box_qty']}），以先讀到的為準")
            if not old or (not old["box_qty"] and rec["box_qty"]):
                out[rec["code3"]] = rec
    return out, warns


def conv_status(conn):
    """各線別轉換率的來源狀態。瑪氏用瑪氏商品總表（跟瑪氏出貨頁共用；蝦皮特選的上傳入口也收）。"""
    st = {}
    for r in _rows(conn.execute("SELECT line, COUNT(*) AS n, SUM(CASE WHEN box_qty IS NULL THEN 1 ELSE 0 END) AS blank, "
                                "MAX(updated_at) AS at FROM shp_conv GROUP BY line")):
        st[r["line"] or "未知"] = {"count": r["n"], "blank": r["blank"] or 0, "updated_at": r["at"] or ""}
    for r in _rows(conn.execute("SELECT line, source_file, updated_by FROM shp_conv WHERE updated_at IN (SELECT MAX(updated_at) FROM shp_conv GROUP BY line)")):
        if r["line"] in st:
            st[r["line"]].update(file=r["source_file"], by=r["updated_by"])
    n = conn.execute("SELECT COUNT(*) AS n FROM mst_mars_products").fetchone()["n"]
    if n:
        last = _row(conn.execute("SELECT * FROM mst_mars_uploads ORDER BY id DESC LIMIT 1"))
        st[MARS] = {"count": n, "blank": 0, "updated_at": (last or {}).get("uploaded_at", ""), "file": (last or {}).get("filename", ""),
                    "by": (last or {}).get("operator", ""), "external": True}
    return st


@shopee_bp.route("/api/shopee/conv")
def api_conv():
    conn = get_conn()
    try:
        return jsonify({"lines": conv_status(conn)})
    finally:
        conn.close()


def _waiting_items(conn, line):
    """還沒產採購單、要用這個線別價格本的品項（竹運採購進貨、供應商直送；還沒掛到任何 EIP）→ {料號3: 品名}。"""
    marks = ",".join("?" * len(BUY_FULFILS))
    out = {}
    for r in _rows(conn.execute(f"SELECT base_code, unit_code, sku_name FROM shp_orders o WHERE line = ? AND removed = 0 AND fulfil IN ({marks}) "
                                "AND NOT EXISTS (SELECT 1 FROM shp_eip_items i WHERE i.order_id = o.id) ORDER BY id", [line, *BUY_FULFILS])):
        if r["base_code"]:
            out.setdefault((r["base_code"] or "") + (r["unit_code"] or ""), r["sku_name"] or "")
    return out


def _conv_gaps(conn, codes):
    """codes 裡價格本找不到的、有但沒有箱入數的（跟產採購單時的查法一樣，不分線別）。"""
    have = _conv_map(conn, codes)
    return {c for c in codes if c not in have}, {c for c in codes if c in have and not box_of(have[c])}


@shopee_bp.route("/api/shopee/conv/import", methods=["POST"])
def api_conv_import():
    """上傳價格本：檔案裡出現的線別整份換掉（別的線別不動）。"""
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "沒有收到檔案。"}), 400
    try:
        rows, warns = parse_conv(f.read(), f.filename)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    # 一份價格本就是一個線別：用多數料號的長相決定（紙潔價格本裡有 AAA… 組合品料號、M 開頭的料號，不能一列一列認）
    votes = collections.Counter(r["line"] for r in rows.values() if r["line"])
    by_line = collections.defaultdict(list)
    if votes:
        main = votes.most_common(1)[0][0]
        for r in rows.values():
            r["line"] = main
            by_line[main].append(r)
    if not by_line:
        return jsonify({"error": f"{f.filename}：認不出是哪個線別的價格本（料號應為國條或 7 碼永豐料號）。"}), 400
    operator, stamp = _operator(), now()
    conn = get_conn()
    try:
        result, pending = {}, {}
        for line, rs in by_line.items():
            # Becky 2026-10-16：「更新價格本了 是不是得重拋採購單？」——上傳完直接講還在等價格的品項找到了沒有
            waiting = _waiting_items(conn, line)
            miss0, nobox0 = _conv_gaps(conn, set(waiting))
            before = conn.execute("SELECT COUNT(*) AS n FROM shp_conv WHERE line = ?", (line,)).fetchone()["n"]
            conn.execute("DELETE FROM shp_conv WHERE line = ?", (line,))
            for r in rs:
                conn.execute("INSERT INTO shp_conv (code3, line, base_code, unit_code, unit_name, box_qty, name, source_file, updated_by, updated_at) "
                             "VALUES (?,?,?,?,?,?,?,?,?,?)", (r["code3"], line, r["base_code"], r["unit_code"], r["unit_name"], r["box_qty"],
                                                              r["name"], f.filename, operator, stamp))
            blank = sum(1 for r in rs if not r["box_qty"])
            result[line] = {"count": len(rs), "before": before, "blank": blank}
            miss1, nobox1 = _conv_gaps(conn, set(waiting))
            named = lambda cs: [{"code3": c, "name": waiting[c]} for c in sorted(cs)]   # noqa: E731
            pending[line] = {"waiting": len(waiting), "was_blocked": len(miss0 | nobox0), "fixed": named((miss0 | nobox0) - miss1 - nobox1),
                             "missing": named(miss1), "no_box": named(nobox1)}
            _log(conn, line, "", "", "", "shopee_conv", "價格本", f"{before} 筆", f"{len(rs)} 筆（箱入數空白 {blank} 筆）", operator, "import", f.filename)
        conn.commit()
        return jsonify({"ok": True, "lines": result, "pending": pending, "warnings": warns, "status": conv_status(conn)})
    finally:
        conn.close()


# ── 產出計畫：每個品項怎麼寫進 EIP、有沒有問題 ──────────────────────────────
def _conv_map(conn, codes):
    if not codes:
        return {}
    marks = ",".join("?" * len(codes))
    return {r["code3"]: r for r in _rows(conn.execute(f"SELECT * FROM shp_conv WHERE code3 IN ({marks})", list(codes)))}


def _mars_map(conn, items):
    """瑪氏品項 → 商品總表那列，用（永豐料號, 單位）對。"""
    codes = sorted({i["base_code"] for i in items if i["line"] == MARS and i["base_code"]})
    if not codes:
        return {}
    marks = ",".join("?" * len(codes))
    return {(r["yf_sku"], r["unit"]): r for r in _rows(conn.execute(f"SELECT * FROM mst_mars_products WHERE yf_sku IN ({marks})", codes))}


def conv_for_item(conn, it):
    """改數量時用：這個品項的轉換率來源（瑪氏是商品總表，其他是價格本）。"""
    if it["line"] == MARS:
        return _row(conn.execute("SELECT * FROM mst_mars_products WHERE yf_sku = ? AND unit = ?", (it["base_code"], it["unit_name"])))
    return _row(conn.execute("SELECT * FROM shp_conv WHERE code3 = ?", ((it["base_code"] or "") + (it["unit_code"] or ""),)))


def box_of(conv):
    """箱入數轉整數（瑪氏商品總表存成小數 24.0）；沒有或不是正數回 None。"""
    try:
        v = float(conv["box_qty"]) if conv and conv["box_qty"] not in (None, "") else 0
    except (TypeError, ValueError):
        return None
    return int(v) if v > 0 and v.is_integer() else (v if v > 0 else None)


def _adj_map(conn, ids):
    if not ids:
        return {}
    marks = ",".join("?" * len(ids))
    return {r["order_id"]: r for r in _rows(conn.execute(f"SELECT * FROM shp_qty_adj WHERE order_id IN ({marks})", list(ids)))}


def _eip_map(conn, ids):
    """品項 → 掛在哪一張已產出的 EIP（最新那張）。"""
    if not ids:
        return {}
    marks = ",".join("?" * len(ids))
    out = {}
    for r in _rows(conn.execute(f"SELECT i.order_id, e.id, e.eip_po, e.filename FROM shp_eip_items i JOIN shp_eip e ON e.id = i.eip_id "
                                f"WHERE i.order_id IN ({marks}) ORDER BY e.id", list(ids))):
        out[r["order_id"]] = {"eip_id": r["id"], "eip_po": r["eip_po"] or "", "filename": r["filename"]}
    return out


def item_plan(it, conv, adj, eip, other_units=()):
    """一個品項要怎麼下：換算、寫進哪個檔、有什麼問題。other_units：瑪氏商品總表有這個料號、但單位不同的那幾列（找不到時講清楚）。"""
    line, fulfil = it["line"], it["fulfil"]
    code3 = (it["base_code"] or "") + (it["unit_code"] or "")
    box = box_of(conv)
    qty = it["qty"] or 0
    use = adj["qty_after"] if adj and adj["qty_before"] == qty else qty      # 蝦皮數量又變了，之前改的就不算
    p = {"order_id": it["id"], "supplier_sku_id": it["supplier_sku_id"], "code": it["base_code"], "code3": code3,
         "name": it["sku_name"], "unit": it["unit_name"] or UNIT_CODES.get(it["unit_code"] or "", ""), "qty": qty, "use_qty": use,
         "conv": box, "adj": adj if adj and adj["qty_before"] == qty else None, "problems": [], "blocking": False,
         "eip": eip, "unit_group": "", "out_unit": "", "out_qty": None, "remark": "", "rem": 0}
    if line not in SUPPORTED:
        p["problems"].append("認不出線別")
        p["blocking"] = True
        return p
    if line == MARS and fulfil != "供應商直送":
        p["problems"].append("瑪氏目前只開放供應商直送")
        p["blocking"] = True
        return p
    if it["diff_pending"]:
        p["problems"].append("這張單有差異待確認，請先到訂單總覽確認")
        p["blocking"] = True
    if eip and eip["eip_po"]:
        p["problems"].append(f"已回填 EIP 採購單號 {eip['eip_po']}；要重新產出請先清除單號")
        p["blocking"] = True
    if conv is None:
        p["problems"].append((f"瑪氏商品總表有 {it['base_code']} 的{'、'.join(other_units)}，沒有{p['unit']}，請按「補商品資料」填寫" if other_units
                              else f"瑪氏商品總表找不到 {it['base_code']}（{p['unit']}），請按「補商品資料」填寫") if line == MARS
                             else f"價格本找不到 {code3}，請更新{line}價格本")
        p["blocking"] = True
        return p
    if not box:
        p["problems"].append(f"瑪氏商品總表 {it['base_code']}（{p['unit']}）沒有箱入數" if line == MARS
                             else f"價格本 {code3} 沒有箱入數，請補上後重新上傳")
        p["blocking"] = True
        return p
    p["rem"] = use % box
    if p["rem"]:
        p["problems"].append(f"{fmt_n(use)} 除以轉換率 {box} 除不盡（餘 {p['rem']}），請修改數量")
        p["blocking"] = True
    if line == "寶僑":
        p["unit_group"] = "箱單位" if box == 1 else "小單位"
        if fulfil == "供應商直送" and box != 1:
            p["out_unit"], p["out_qty"] = "箱", use // box if not p["rem"] else round(use / box, 2)
        else:
            p["out_unit"], p["out_qty"] = p["unit"], use
        p["remark"] = "蝦皮特選"
    elif line == MARS:                               # 瑪氏：單位、數量照蝦皮，不換箱（需求 10.4）；箱數給瑪氏採購單用
        cat = cat_code(conv["category"])
        lab = "V" if (conv["inner_label"] or "").strip().upper() == "V" else ""
        p["unit_group"] = f"{p['unit']}_{cat or '品類不明'}_{'需貼中標' if lab else '不貼中標'}"
        p["out_unit"], p["out_qty"] = p["unit"], use
        p["remark"] = f"MARS入倉{(it['expected_date'] or '').replace('-', '')}台灣瑪氏送蝦皮特選-{wh_name(it['warehouse'])}"
        p["cases"] = use // box if not p["rem"] else round(use / box, 2)
        p["mars"] = {"category": cat, "label": lab, "mars_code": conv["mars_code"] or "", "mars_name": conv["name"] or "",
                     "price": conv["price"], "pcs_per_case": conv["pcs_per_case"], "inner_per_case": conv["inner_per_case"],
                     "note": conv["note"] or "", "po_case_note": conv["po_case_note"] or ""}
    else:                                            # 紙潔：一律換成箱（需求 10.3）
        p["out_unit"], p["out_qty"] = "箱", use // box if not p["rem"] else round(use / box, 2)
    return p


def fmt_n(n):
    return f"{n:,}" if isinstance(n, int) else str(n)


def build_plan(conn, fulfil, d1="", d2="", line="", order_nos=None, up1="", up2=""):
    where, params = ["fulfil = ?", "removed = 0"], [fulfil]
    if d1:
        where.append("expected_date >= ?"); params.append(d1)
    if d2:
        where.append("expected_date <= ?"); params.append(d2)
    usql, up = upload_cond(up1, up2)
    if usql:
        where.append(usql); params += up
    if line:
        where.append("line = ?"); params.append(line)
    items = _rows(conn.execute(f"SELECT * FROM shp_orders WHERE {' AND '.join(where)} ORDER BY expected_date, warehouse, id", params))
    if order_nos is not None:
        want = set(order_nos)
        items = [i for i in items if order_no(i) in want]
    ids = [i["id"] for i in items]
    convs = _conv_map(conn, {(i["base_code"] or "") + (i["unit_code"] or "") for i in items if i["line"] != MARS})
    mars = _mars_map(conn, items)
    adjs, eips = _adj_map(conn, ids), _eip_map(conn, ids)
    groups = collections.OrderedDict()
    for it in items:
        no = order_no(it)
        g = groups.setdefault(no, {"order_no": no, "pr_id": it["pr_id"], "po_id": it["po_id"], "line": it["line"], "warehouse": it["warehouse"],
                                   "warehouse_name": wh_name(it["warehouse"]), "expected_date": it["expected_date"], "items": [], "_rows": []})
        g["_rows"].append(it)
        conv = mars.get((it["base_code"], it["unit_name"])) if it["line"] == MARS else convs.get((it["base_code"] or "") + (it["unit_code"] or ""))
        others = sorted(u for (c, u) in mars if c == it["base_code"] and u != it["unit_name"]) if it["line"] == MARS and conv is None else ()
        g["items"].append(item_plan(it, conv, adjs.get(it["id"]), eips.get(it["id"]), others))
    out = []
    for g in groups.values():
        its = g["items"]
        filled = [p for p in its if p["eip"] and p["eip"]["eip_po"]]
        made = [p for p in its if p["eip"]]
        g.update(group_meta(g.pop("_rows")))
        g.update(blocking=sum(1 for p in its if p["blocking"]), supported=g["line"] in SUPPORTED,
                 po_status="已回填" if its and len(filled) == len(its) else "已產出" if made else "未產出",
                 qty=sum(p["qty"] for p in its))
        out.append(g)
    return out


@shopee_bp.route("/api/shopee/purchase/plan")
def api_purchase_plan():
    fulfil = request.args.get("fulfil", "")
    if fulfil not in BUY_FULFILS:
        return jsonify({"error": "履約方式只能是竹運採購進貨或供應商直送。"}), 400
    conn = get_conn()
    try:
        orders = build_plan(conn, fulfil, request.args.get("from", ""), request.args.get("to", ""), request.args.get("line", ""),
                            up1=request.args.get("up_from", ""), up2=request.args.get("up_to", ""))
        st = conv_status(conn)
    finally:
        conn.close()
    return jsonify({"orders": orders, "conv": st})


# ── 補登瑪氏商品（Jerry 2026-10-14）──────────────────────────────────────────
# 瑪氏新品不在商品總表就產不了 EIP／瑪氏採購單。以前要改 Alice 的 Excel、再到瑪氏出貨頁整份重傳；
# 需求書只說商品總表是對照來源、要可維護（3.2、第十三章），沒說只能改 Excel，所以在面板直接補。
# 存進同一份 mst_mars_products（酷澎瑪氏也用得到），manual_by 記是誰補的；之後上傳 Excel 不會蓋掉，Excel 有了才以 Excel 為準。
MARS_UNITS = ("箱", "盒", "包")


def _mars_item(conn, order_id):
    it = _row(conn.execute("SELECT * FROM shp_orders WHERE id = ?", (order_id,)))
    if it is None:
        return None, None, (jsonify({"error": "找不到這個品項，請重新整理頁面。"}), 404)
    if it["line"] != MARS:
        return None, None, (jsonify({"error": "只有瑪氏的品項要補瑪氏商品總表。"}), 400)
    unit = it["unit_name"] or UNIT_CODES.get(it["unit_code"] or "", "")
    if unit not in MARS_UNITS:
        return None, None, (jsonify({"error": f"單位「{unit or '不明'}」不是箱、盒、包，無法補登。"}), 400)
    return it, unit, None


@shopee_bp.route("/api/shopee/mars_product")
def api_mars_product_get():
    """補商品資料視窗：這一品的料號、單位、蝦皮品名，和同料號其他單位那幾列（可帶入品類、瑪氏貨號、價格）。"""
    conn = get_conn()
    try:
        it, unit, err = _mars_item(conn, request.args.get("order_id"))
        if err:
            return err
        same = _rows(conn.execute("SELECT * FROM mst_mars_products WHERE yf_sku = ? ORDER BY unit", (it["base_code"],)))
    finally:
        conn.close()
    return jsonify({"code": it["base_code"], "unit": unit, "name": it["sku_name"], "exists": any(r["unit"] == unit for r in same),
                    "siblings": [r for r in same if r["unit"] != unit],
                    "categories": [{"value": v, "label": CAT_FULL[c]} for v, c in (("Chocolate", "CHO"), ("Gum", "GUM"), ("Petcare", "PET"))]})


def _num(v, kind, label, required=False):
    s_ = str(v if v is not None else "").strip().replace(",", "")
    if s_ == "":
        if required:
            raise ValueError(f"請填寫{label}。")
        return None
    try:
        n = float(s_)
    except ValueError:
        raise ValueError(f"{label}要填數字。") from None
    if n <= 0 if kind == "int" else n < 0:
        raise ValueError(f"{label}要大於 0。" if kind == "int" else f"{label}不能是負數。")
    if kind == "int":
        if not n.is_integer():
            raise ValueError(f"{label}要填整數。")
        return int(n)
    return n


@shopee_bp.route("/api/shopee/mars_product", methods=["POST"])
def api_mars_product_add():
    p = request.get_json(silent=True) or {}
    cat = str(p.get("category") or "").strip()
    if not cat_code(cat):
        return jsonify({"error": "請選擇品類。"}), 400
    try:
        box = _num(p.get("box_qty"), "int", "箱入數", required=True)
        price = _num(p.get("price"), "float", "系統價格")
        days = _num(p.get("shelf_days"), "int", "效期天")
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    label = "V" if str(p.get("inner_label") or "").strip().upper() == "V" else ""
    mars_code = str(p.get("mars_code") or "").strip()
    operator = _operator()
    conn = get_conn()
    try:
        it, unit, err = _mars_item(conn, p.get("order_id"))
        if err:
            return err
        code = it["base_code"]
        if _row(conn.execute("SELECT id FROM mst_mars_products WHERE yf_sku = ? AND unit = ?", (code, unit))):
            return jsonify({"error": f"商品總表已經有 {code}（{unit}），請重新整理頁面。"}), 409
        name = re.sub(r"^【[^】]*】", "", it["sku_name"] or "").strip()[:120]
        conn.execute("""INSERT INTO mst_mars_products (yf_sku, unit, mars_code, category, name, price, unit_code, box_qty, inner_label,
                        shelf_days, manual_by, manual_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                     (code, unit, mars_code, cat, name, price, it["unit_code"] or "", box, label, days, operator, now()))
        _log(conn, MARS, order_no(it), it["supplier_sku_id"], "", "shopee_mars_product", "補登瑪氏商品總表", "",
             f"{code}（{unit}）{cat}、箱入數 {box}、{'需貼中標' if label else '不貼中標'}"
             + (f"、瑪氏貨號 {mars_code}" if mars_code else "") + (f"、價格 {price:g}" if price is not None else "")
             + (f"、效期天 {days}" if days else ""), operator, "manual", "在蝦皮特選補登；總表 Excel 裡還沒有")
        conn.commit()
    finally:
        conn.close()
    return jsonify({"ok": True, "code": code, "unit": unit})


@shopee_bp.route("/api/shopee/purchase/adjust", methods=["POST"])
def api_purchase_adjust():
    """除不盡轉換率時改下採數量（需求 10.5）。改後要整除、要寫原因；清空＝取消修改。"""
    payload = request.get_json(silent=True) or {}
    try:
        oid = int(payload.get("order_id"))
    except (TypeError, ValueError):
        return jsonify({"error": "缺少品項。"}), 400
    reason = str(payload.get("reason") or "").strip()
    raw = payload.get("qty")
    operator = _operator()
    conn = get_conn()
    try:
        it = _row(conn.execute("SELECT * FROM shp_orders WHERE id = ?", (oid,)))
        if it is None:
            return jsonify({"error": "找不到這個品項，請重新整理頁面。"}), 404
        no = order_no(it)
        old = _row(conn.execute("SELECT * FROM shp_qty_adj WHERE order_id = ?", (oid,)))
        if raw in (None, ""):
            if old:
                conn.execute("DELETE FROM shp_qty_adj WHERE order_id = ?", (oid,))
                _log(conn, it["line"], no, it["shopee_sku_id"], "", "shopee_qty_adj", "下採數量", old["qty_after"], f"取消修改（回到 {it['qty']}）", operator, "manual", it["supplier_sku_id"])
                conn.commit()
            return jsonify({"ok": True, "cleared": True})
        try:
            qty = int(str(raw).replace(",", ""))
        except ValueError:
            return jsonify({"error": "修改後數量請填整數。"}), 400
        box = box_of(conv_for_item(conn, it))
        if not box:
            return jsonify({"error": "瑪氏商品總表沒有這個品項的箱入數，無法判斷，請先更新商品總表。" if it["line"] == MARS
                            else "價格本沒有這個品項的箱入數，無法判斷，請先更新價格本。"}), 400
        if qty < 0 or qty % box:
            return jsonify({"error": f"修改後數量 {qty} 仍除不盡轉換率 {box}，請改成 {box} 的倍數（例如 {it['qty'] // box * box} 或 {(it['qty'] // box + 1) * box}）。"}), 400
        if not reason:
            return jsonify({"error": "請填寫修改原因。"}), 400
        stamp = now()
        if old:
            conn.execute("UPDATE shp_qty_adj SET qty_before = ?, qty_after = ?, conv = ?, reason = ?, operator = ?, updated_at = ? WHERE order_id = ?",
                         (it["qty"], qty, box, reason, operator, stamp, oid))
        else:
            conn.execute("INSERT INTO shp_qty_adj (order_id, qty_before, qty_after, conv, reason, operator, updated_at) VALUES (?,?,?,?,?,?,?)",
                         (oid, it["qty"], qty, box, reason, operator, stamp))
        _log(conn, it["line"], no, it["shopee_sku_id"], "", "shopee_qty_adj", "下採數量", it["qty"], qty, operator, "manual",
             f"{it['supplier_sku_id']}；轉換率 {box}；{reason}", reason)
        conn.commit()
        return jsonify({"ok": True, "qty": qty})
    finally:
        conn.close()


# ── 產檔 ────────────────────────────────────────────────────────────────────
def _merge_rows(plans):
    """同一個料號同一個單位合成一列（不同 PO 同一天同倉會撞在一起）。"""
    rows = collections.OrderedDict()
    for p in plans:
        k = (p["code"], p["out_unit"])
        r = rows.get(k)
        if r is None:
            r = rows[k] = {"code": p["code"], "name": p["name"], "unit": p["out_unit"], "qty": 0, "in_unit": p["unit"], "in_qty": 0,
                           "conv": p["conv"], "remark": p["remark"], "parts": []}       # parts：每張單各下多少（直送採購單一品一列、寫採購單ID）
            if p.get("mars"):                        # 瑪氏採購單要的商品總表欄位一起存，之後照這份產
                r.update(p["mars"], cases=0)
        r["qty"] += p["out_qty"]
        r["in_qty"] += p["use_qty"]
        r["parts"].append({"order_id": p["order_id"], "use_qty": p["use_qty"]})
        if "cases" in r:
            r["cases"] += p.get("cases") or 0
    return list(rows.values())


def eip_xls(rows, line):
    """套公司 EIP 範本（purchase_templates/pg_template.xls，跟蝦皮範例「產品採購單」同一張表；瑪氏用 mars_template.xls）。
    寶僑照蝦皮範例：第 10 欄標題清空、放轉換率。"""
    with_conv = line == "寶僑"
    template, wipe = (MARS_TEMPLATE, MARS_WIPE_ROWS) if line == MARS else (TEMPLATE, WIPE_ROWS)
    rb = xlrd.open_workbook(template, formatting_info=True)
    wb = xl_copy(rb)
    ws = wb.get_sheet(0)
    for r in range(1, max(wipe, len(rows) + 1)):
        for c in range(10):
            ws.write(r, c, "")
    if with_conv:
        ws.write(0, 9, "")
    for i, row in enumerate(rows, start=1):
        ws.write(i, 0, i)
        ws.write(i, 1, str(row["code"]))
        ws.write(i, 2, row["name"])
        ws.write(i, 3, row["unit"])
        ws.write(i, 6, row["qty"])
        ws.write(i, 8, row["remark"])
        if with_conv:
            ws.write(i, 9, row["conv"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def detail_xlsx(rows, line="寶僑"):
    """採購下採明細：照「蝦皮特選-商品下採.xlsx」，寶僑兩頁「箱」（轉換率 1）「小單位」；紙潔 EIP 一律換箱、不分，明細也只出一頁
    （Jerry 2026-10-28）。數量/箱＝進貨÷轉換率（公式）。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head = ["永豐料號", "品名", "單位", "進貨", "轉換率", "數量/箱"]
    pages = (("箱", lambda r: r["conv"] == 1), ("小單位", lambda r: r["conv"] != 1)) if line == "寶僑" else (("下採", lambda r: True),)
    for title, pick in pages:
        ws = wb.create_sheet(title)
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FBE4DC")
        for i, r in enumerate([x for x in rows if pick(x)], start=2):
            ws.append([r["code"], r["name"], r["in_unit"], r["in_qty"], r["conv"], f"=D{i}/E{i}"])
        for col, w in zip("ABCDEF", (19, 90, 6, 8, 9, 9)):
            ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _md(d):
    return f"{d[5:7]}{d[8:10]}" if d and len(d) >= 10 else "日期不明"


def _when_where(dates, whs):
    """檔名裡的到貨日與倉：0916到觀音；跨天跨倉的寫 1019-1021到觀音、安南三。"""
    d = _md(dates[0]) + (f"-{_md(dates[-1])}" if dates[-1] != dates[0] else "")
    return f"{d}到{'、'.join(wh_name(w) for w in whs)}"


def eip_filename(line, fulfil, dates, whs, unit_group):
    return f"產品採購表上傳-{line}{FULFIL_SHORT.get(fulfil, fulfil)}_{_when_where(dates, whs)}" + (f"-{unit_group}" if unit_group else "") + ".xls"


def detail_filename(line, fulfil, dates, whs):
    return f"蝦皮特選-商品下採_{line}{FULFIL_SHORT.get(fulfil, fulfil)}_{_when_where(dates, whs)}.xlsx"


@shopee_bp.route("/api/shopee/purchase/generate", methods=["POST"])
def api_purchase_generate():
    """勾選的單 → zip（EIP 採購單＋採購下採明細）。有任何擋下來的品項就整批不產，列出原因。"""
    payload = request.get_json(silent=True) or {}
    fulfil = payload.get("fulfil", "")
    nos = [str(x) for x in (payload.get("order_nos") or []) if str(x).strip()]
    if fulfil not in BUY_FULFILS:
        return jsonify({"error": "履約方式只能是竹運採購進貨或供應商直送。"}), 400
    if not nos:
        return jsonify({"error": "請先勾選要產出採購單的訂單。"}), 400
    operator, stamp = _operator(), now()
    conn = get_conn()
    try:
        orders = build_plan(conn, fulfil, order_nos=nos)
        if not orders:
            return jsonify({"error": "勾選的訂單已不是這個履約方式，請重新整理頁面。"}), 400
        bad = [f"{g['order_no']} {p['supplier_sku_id']}：{'；'.join(p['problems'])}" for g in orders for p in g["items"] if p["blocking"]]
        if bad:
            return jsonify({"error": f"有 {len(bad)} 個品項需要先處理，無法產出。", "details": bad[:30]}), 400
        # 分檔：供應商直送一天一倉一個檔（貨直接送蝦皮倉，PG直出範例一張 PO 一組）；
        # 竹運採購進貨是進我們的竹運倉，整批只分箱單位／小單位（PG竹運進貨範例：10/19～10/21、觀音和安南三合在同一個檔）
        files, details = collections.OrderedDict(), collections.OrderedDict()
        for g in orders:
            key = (g["line"], g["expected_date"], g["warehouse"]) if fulfil == "供應商直送" else (g["line"], "", "")
            for p in g["items"]:
                p["_date"], p["_wh"] = g["expected_date"], g["warehouse"]
                files.setdefault(key + (p["unit_group"],), []).append(p)
                if g["line"] != MARS:                # 瑪氏出瑪氏制式採購單，不出下採明細（需求第七章產出表）
                    details.setdefault(key, []).append(p)
        ids = [p["order_id"] for g in orders for p in g["items"]]
        marks = ",".join("?" * len(ids))
        # 之前產過、還沒回填單號的：從舊的那張拿掉（舊的那張空了就刪），這次重產為準
        old_links = _rows(conn.execute(f"SELECT i.id, i.eip_id FROM shp_eip_items i JOIN shp_eip e ON e.id = i.eip_id "
                                       f"WHERE i.order_id IN ({marks}) AND (e.eip_po = '' OR e.eip_po IS NULL)", ids))
        for x in old_links:
            conn.execute("DELETE FROM shp_eip_items WHERE id = ?", (x["id"],))
        for eid in {x["eip_id"] for x in old_links}:
            if conn.execute("SELECT COUNT(*) AS n FROM shp_eip_items WHERE eip_id = ?", (eid,)).fetchone()["n"] == 0:
                conn.execute("DELETE FROM shp_eip WHERE id = ?", (eid,))
        batch = stamp
        buf = io.BytesIO()
        names = []
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for (line, date, whc, ug), ps in files.items():
                rows = [r for r in _merge_rows(ps) if r["qty"]]       # 數量改成 0＝這品不下採，不寫進 EIP
                if not rows:
                    continue
                dates, whs = sorted({p["_date"] for p in ps}), sorted({p["_wh"] for p in ps})
                date, whc = dates[0], "、".join(whs)
                name = eip_filename(line, fulfil, dates, whs, ug)
                zf.writestr(name, eip_xls(rows, line))
                names.append(name)
                cur = conn.execute("INSERT INTO shp_eip (batch, line, fulfil, warehouse, expected_date, unit_group, filename, rows_json, eip_po, created_by, created_at, updated_by, updated_at) "
                                   "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", (batch, line, fulfil, whc, date, ug, name, json.dumps(rows, ensure_ascii=False), "", operator, stamp, operator, stamp))
                eid = cur.lastrowid
                for p in ps:
                    conn.execute("INSERT INTO shp_eip_items (eip_id, order_id) VALUES (?,?)", (eid, p["order_id"]))
            for (line, _d, _w), ps in details.items():
                name = detail_filename(line, fulfil, sorted({p["_date"] for p in ps}), sorted({p["_wh"] for p in ps}))
                zf.writestr(name, detail_xlsx([r for r in _merge_rows(ps) if r["in_qty"]], line))
                names.append(name)
        for g in orders:
            _log(conn, g["line"], g["order_no"], "", "", "shopee_eip", "產出 EIP 採購單", "", f"{fulfil}；{len(g['items'])} 品項", operator, "manual",
                 "、".join(n for n in names if n.endswith(".xls")))
        conn.commit()
    finally:
        conn.close()
    dates = sorted({p["_date"] for ps in files.values() for p in ps})
    zname = f"蝦皮特選{FULFIL_SHORT[fulfil]}採購單_{_md(dates[0])}" + (f"-{_md(dates[-1])}" if dates[-1] != dates[0] else "") + ".zip"
    buf.seek(0)
    resp = send_file(buf, as_attachment=True, download_name=zname, mimetype="application/zip")
    resp.headers["X-Shopee-Eip-Files"] = str(sum(1 for n in names if n.endswith(".xls")))
    resp.headers["X-Shopee-Files"] = str(len(names))
    return resp


# ── 採購單產出：已產出的 EIP 檔、回填 EIP 單號 ───────────────────────────────
@shopee_bp.route("/api/shopee/eips")
def api_eips():
    where, params = ["1=1"], []
    if request.args.get("from"):
        where.append("expected_date >= ?"); params.append(request.args["from"])
    if request.args.get("to"):
        where.append("expected_date <= ?"); params.append(request.args["to"])
    usql, up = upload_cond(request.args.get("up_from", ""), request.args.get("up_to", ""), "o.")
    if usql:                                     # 上傳日：看這個 EIP 檔裡的訂單
        where.append(f"id IN (SELECT i.eip_id FROM shp_eip_items i JOIN shp_orders o ON o.id = i.order_id WHERE {usql})"); params += up
    conn = get_conn()
    try:
        eips = _rows(conn.execute(f"SELECT * FROM shp_eip WHERE {' AND '.join(where)} ORDER BY expected_date DESC, warehouse, line, fulfil, unit_group, id", params))
        links = collections.defaultdict(list)
        if eips:
            marks = ",".join("?" * len(eips))
            for r in _rows(conn.execute(f"SELECT i.eip_id, o.po_id, o.pr_id FROM shp_eip_items i JOIN shp_orders o ON o.id = i.order_id WHERE i.eip_id IN ({marks})",
                                        [e["id"] for e in eips])):
                links[r["eip_id"]].append(r["po_id"] or r["pr_id"])
    finally:
        conn.close()
    out = []
    for e in eips:
        rows = json.loads(e["rows_json"] or "[]")
        out.append({k: e[k] for k in ("id", "line", "fulfil", "warehouse", "expected_date", "unit_group", "filename", "eip_po", "created_by", "created_at", "updated_by", "updated_at")}
                   | {"warehouse_name": "、".join(wh_name(w) for w in (e["warehouse"] or "").split("、") if w), "rows": len(rows), "qty": sum(r["qty"] for r in rows),
                      "mars": e["line"] == MARS, "direct": e["fulfil"] == "供應商直送" and e["line"] != MARS,
                      "orders": sorted(set(links.get(e["id"], []))), "note": e.get("note") or ""})
    return jsonify({"eips": out})


@shopee_bp.route("/api/shopee/eips/<int:eid>", methods=["PUT"])
def api_eip_update(eid):
    payload = request.get_json(silent=True) or {}
    eip = str(payload.get("eip_po") or "").strip().upper().replace(" ", "")
    if eip and not EIP_PO_RE.match(eip):
        return jsonify({"error": f"EIP 採購單號格式為 PO 加 9 碼數字（例如 PO202610003），目前填的是「{eip}」。"}), 400
    operator = _operator()
    conn = get_conn()
    try:
        e = _row(conn.execute("SELECT * FROM shp_eip WHERE id = ?", (eid,)))
        if e is None:
            return jsonify({"error": "找不到這張 EIP 採購單，可能已重新產出，請重新整理頁面。"}), 404
        if eip:
            dup = _row(conn.execute("SELECT filename FROM shp_eip WHERE eip_po = ? AND id != ?", (eip, eid)))
            if dup:
                return jsonify({"error": f"{eip} 已用於「{dup['filename']}」，請確認號碼是否正確。"}), 400
        if eip != (e["eip_po"] or ""):
            conn.execute("UPDATE shp_eip SET eip_po = ?, updated_by = ?, updated_at = ? WHERE id = ?", (eip, operator, now(), eid))
            _log(conn, e["line"], "", "", "", "shopee_eip_po", "EIP 採購單號", e["eip_po"], eip, operator, "manual", e["filename"])
            conn.commit()
        return jsonify({"ok": True, "eip_po": eip})
    finally:
        conn.close()


@shopee_bp.route("/api/shopee/eips/<int:eid>", methods=["DELETE"])
def api_eip_cancel(eid):
    """取消產出（Jerry 2026-10-21「產出 EIP 之後反悔，回到待產採購單」）：這份檔刪掉，裡面的訂單退回待產採購單。
    已回填單號的擋下來，要先清掉單號，免得手滑把回填過的拿掉。蝦皮那邊已經上傳的 EIP 要自己撤。"""
    from .orders import order_no as _no
    operator = _operator()
    conn = get_conn()
    try:
        e = _row(conn.execute("SELECT * FROM shp_eip WHERE id = ?", (eid,)))
        if e is None:
            return jsonify({"error": "找不到這張 EIP 採購單，可能已重新產出，請重新整理頁面。"}), 404
        if e["eip_po"]:
            return jsonify({"error": f"這份已回填 EIP 採購單號 {e['eip_po']}，請先清掉單號再取消產出。"}), 400
        ids = [r["order_id"] for r in _rows(conn.execute("SELECT order_id FROM shp_eip_items WHERE eip_id = ?", (eid,)))]
        nos = []
        if ids:
            marks = ",".join("?" * len(ids))
            nos = sorted({_no(o) for o in _rows(conn.execute(f"SELECT * FROM shp_orders WHERE id IN ({marks})", ids))})
        conn.execute("DELETE FROM shp_eip_items WHERE eip_id = ?", (eid,))
        conn.execute("DELETE FROM shp_eip WHERE id = ?", (eid,))
        for no in nos:
            _log(conn, e["line"], no, "", "", "shopee_eip", "取消產出 EIP 採購單", e["filename"], "未產出", operator, "manual")
        conn.commit()
        return jsonify({"ok": True, "orders": nos})
    finally:
        conn.close()


@shopee_bp.route("/api/shopee/eips/<int:eid>/file")
def api_eip_file(eid):
    """照當時存下來的內容重新下載（檔名前面加上 EIP 單號，有填的話）。"""
    conn = get_conn()
    try:
        e = _row(conn.execute("SELECT * FROM shp_eip WHERE id = ?", (eid,)))
    finally:
        conn.close()
    if e is None:
        return jsonify({"error": "找不到這張 EIP 採購單。"}), 404
    name = e["filename"] if not e["eip_po"] else e["filename"].replace(".xls", f"({e['eip_po']}).xls")
    return send_file(io.BytesIO(eip_xls(json.loads(e["rows_json"] or "[]"), e["line"])), as_attachment=True, download_name=name,
                     mimetype="application/vnd.ms-excel")


def po_status_map(conn, ids):
    """訂單總覽「採購單產出狀態」用：品項 id → 已產出／已回填。"""
    return {k: ("已回填" if v["eip_po"] else "已產出") for k, v in _eip_map(conn, ids).items()}


__all__ = ["parse_conv", "conv_status", "conv_for_item", "box_of", "item_plan", "build_plan", "eip_xls", "detail_xlsx", "eip_filename", "detail_filename", "po_status_map",
           "api_conv", "api_conv_import", "api_purchase_plan", "api_purchase_adjust", "api_purchase_generate", "api_eips", "api_eip_update", "api_eip_file"]
