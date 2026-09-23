"""瑪氏出貨（mars/）的端到端驗證：假的瑪氏商品總表＋假訂單彙總表裡線別是瑪氏的單，
走「商品總表上傳 → ② 匯訂單 → 拆單 → 產出 zip（拆單表＋EIP 採購表）→ 回填 EIP 採購單號／約倉時間 → 訂單改了之後」。

執行：python test_mars.py（設了 DATABASE_URL 就跑 PostgreSQL）
"""
import datetime
import io
import json
import os
import sys
import tempfile
import zipfile
from urllib.parse import unquote

import openpyxl
import xlrd

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FAKE_M = os.path.join(BASE_DIR, "samples", "master", "fake")
FAKE = os.path.join(BASE_DIR, "samples", "mars", "fake")
MASTER = os.path.join(FAKE, "假_瑪氏商品總表.xlsx")

_tmp = tempfile.mkdtemp(prefix="oms_mars_test_")
sys.path.insert(0, BASE_DIR)
import db  # noqa: E402
db.DATA_DIR = os.path.join(_tmp, "資料與設定")
db.DB_PATH = os.path.join(db.DATA_DIR, "test.db")
db.BACKUP_DIR = os.path.join(db.DATA_DIR, "backups")
import app as app_module  # noqa: E402
import mars.common as mc  # noqa: E402

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"  {'✓' if cond else '✗'} {name}" + (f"  — {detail}" if detail else ""))


def up(client, url, name_or_bytes, filename=None, **form):
    data = open(name_or_bytes, "rb").read() if isinstance(name_or_bytes, str) else name_or_bytes
    d = dict(form); d["file"] = (io.BytesIO(data), filename or os.path.basename(name_or_bytes))
    return client.post(url, data=d, content_type="multipart/form-data")


def import_orders(client, data, filename):
    pv = up(client, "/api/master/import/preview", data, filename).get_json()
    return client.post("/api/master/import/commit", json={"batch_id": pv["batch_id"], "add_missing_products": True})


def splits(client, d1, d2=None):
    return client.get(f"/api/mars/splits?from={d1}&to={d2 or d1}").get_json()


def special_sheet(rows):
    """照整合表的表頭做一份只有瑪氏幾列的訂單彙總表。"""
    wb = openpyxl.load_workbook(os.path.join(FAKE_M, "假_訂單彙總表_9月_第一次.xlsx")); ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        ws.append([r.get(h) for h in hdr])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def base_row(**kw):
    r = {"訂單類型": "一般", "線別": "瑪氏", "PO單號": "13000000699901", "到貨倉別": "TAO3", "地址": "桃園市大園區中山南路472號",
         "交付日期": datetime.datetime(2026, 9, 30), "NO": 1, "品牌": "測試", "酷澎下單單價(含稅)": 100}
    r.update(kw)
    return r


def main():
    db.init_db(); app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()

    print("\n【0】登入、頁面、選單")
    check("沒登入打 API → 401", c.get("/api/mars/status").status_code == 401)
    c.post("/login", data={"username": "小真", "password": "changeme123"})
    # CI 把幾套測試跑在同一個 PostgreSQL 裡，前一套會留下訂單、匯入批次；這裡先清空，不依賴資料庫是空的
    app_module._write_users(app_module.get_users(), {"Jerry", "小真"})
    r = c.post("/api/master/reset", json={"confirm": "清空資料", "orders": True, "products": True})
    check("開始前先清空訂單與主檔（含瑪氏的表）", r.status_code == 200, r.get_data(as_text=True)[:120])
    html = c.get("/mars").get_data(as_text=True)
    check("/mars 頁面打得開、有拆單與 JS 版本號", "瑪氏出貨" in html and "btn-gen" in html and "mars.js?v=" in html)
    check("線別工具的瑪氏欄有「瑪氏出貨」，不再是準備中", "/mars" in c.get("/master").get_data(as_text=True))

    print("\n【1】瑪氏商品總表")
    check("還沒上傳時狀態是空的", c.get("/api/mars/status").get_json()["last_upload"] is None)
    r = up(c, "/api/mars/products/import", MASTER); d = r.get_json()
    check("上傳：7 個料號、21 列（箱盒包各一）、品類分得出來", r.status_code == 200 and d["codes"] == 7 and d["rows"] == 21
          and d["by_category"] == {"CHO": 2, "GUM": 2, "PET": 3} and not d["warnings"], str(d)[:200])
    st = c.get("/api/mars/status").get_json()
    check("狀態：上次上傳、料號數、② 裡有瑪氏訂單的日子", st["last_upload"]["filename"] == "假_瑪氏商品總表.xlsx" and st["codes"] == 7 and st["last_order_import"] is None)
    check("不是商品總表的檔 → 400 講清楚要哪幾欄", up(c, "/api/mars/products/import", os.path.join(FAKE_M, "假_酷澎主檔_瑪氏.xlsx")).status_code == 400)
    check("壞掉的檔 → 400", up(c, "/api/mars/products/import", b"not excel", "x.xlsx").status_code == 400)
    wb = openpyxl.load_workbook(MASTER); ws = wb["商品表"]
    ws.append([1, "Candy", "x", "M00000001", "怪品類", 1, 1, 1, 1, "", "", "箱", "R", 1, "", 1, "", ""])
    ws.append([2, "Gum", "x", "M00000002", "怪單位", 1, 1, 1, 1, "", "", "打", "", 1, "", 1, "", ""])
    buf = io.BytesIO(); wb.save(buf)
    d = up(c, "/api/mars/products/import", buf.getvalue(), "改過的.xlsx").get_json()
    check("單位不是箱盒包的列跳過、Category 認不得的列標出來", d["codes"] == 8 and any("打" in w for w in d["warnings"]) and any("Candy" in w for w in d["warnings"]), str(d["warnings"]))
    d = up(c, "/api/mars/products/import", MASTER).get_json()
    check("再傳原本那份＝整份覆蓋：跟上一版比拿掉 1 個", d["codes"] == 7 and d["removed"] == 1 and d["removed_examples"] == ["M00000001"], str(d)[:160])

    print("\n【2】② 匯訂單：地址、報價備註也存下來")
    for n in ("假_訂單彙總表_9月_第一次.xlsx", "假_訂單彙總表_9月_第二次.xlsx"):
        check(f"匯 {n}", import_orders(c, os.path.join(FAKE_M, n), None).status_code == 200)
    conn = db.get_conn()
    o = dict(conn.execute("SELECT * FROM mst_orders WHERE po_number = '13000000600028' AND yf_sku = 'M36752197'").fetchone()); conn.close()
    check("mst_orders 有存地址", o["address"] != "", o["address"])
    st = c.get("/api/mars/status").get_json()
    check("狀態：最近一次有瑪氏單的匯入是第二次那份", st["last_order_import"]["filename"] == "假_訂單彙總表_9月_第二次.xlsx" and "2026-09-18" in st["dates"], str(st["last_order_import"]))

    print("\n【3】拆單：9/18 那張 PO 拆成 5 份")
    v = splits(c, "2026-09-18")
    names = [s["filename"] for s in v["splits"]]
    check("一張 PO、5 份（品類 × 單位 × 中標各不同）", v["summary"]["pos"] == 1 and v["summary"]["files"] == 5, str(names))
    check("檔名照 Alice 的規則：訂單系統拆單表_PO_到貨日_倉_品類_單位_中標_序號", "訂單系統拆單表_13000000600028_20260918_TAO1_CHO_盒_需貼中標_01.xlsx" in names
          and "訂單系統拆單表_13000000600028_20260918_TAO1_PET_箱_不貼中標_01.xlsx" in names, str(names))
    s_cho = next(s for s in v["splits"] if s["category"] == "CHO" and s["label"] == "")
    it = s_cho["items"][0]
    check("箱數 = 出貨數量 ÷ 整合表箱入數（36 ÷ 12 = 3）", it["cases"] == 3 and s_cho["cases_total"] == 3, str(it["cases"]))
    check("整合表箱入數跟商品總表不一樣 → 提醒、箱數照整合表", any("整合表 12、商品總表 6" in x for x in it["issues"]), str(it["issues"]))
    check("帶進商品總表的欄（瑪氏貨號、品名、價格、每箱中盒數…）", it["mars_code"] == "36752197" and it["mars_name"].startswith("【巧克棒】") and it["price"] == 2410.0 and it["inner_per_case"] == 12)
    check("全部都還沒產出", all(s["status"] == "new" and s["id"] is None for s in v["splits"]))
    total_file = sum(r["cases_total"] for r in v["splits"])
    check("拆完的箱數加總 = 這張 PO 的箱數加總（3+1+7+1+3）", total_file == 15, str(total_file))
    v2 = splits(c, "2026-09-11")
    pv = [s for s in v2["splits"] if s["po_number"] == "13000000600020" and s["category"] == "PET"]
    check("9/11 TAO5：同品類、同單位、同中標的兩個品項放同一份", len(pv) == 1 and pv[0]["item_count"] == 2, str([(s["filename"], s["item_count"]) for s in pv]))
    mc.SPLIT_BY_UNIT = False
    try:
        import mars.split as ms
        ms.SPLIT_BY_UNIT = False
        vn = splits(c, "2026-09-05")
    finally:
        mc.SPLIT_BY_UNIT = True; ms.SPLIT_BY_UNIT = True
    vy = splits(c, "2026-09-05")
    check("關掉「依單位拆」：檔名單位那段固定寫箱", all("_箱_" in s["filename"] for s in vn["splits"]) and vn["split_by_unit"] is False)

    print("\n【4】特殊狀況：組出商品、對不到、箱數不是整數")
    sheet = special_sheet([
        base_row(**{"SKU ID": "900000000000001", "條碼(國條)": "4700000000001", "永豐料號": "M60055599", "報價備註": "M55500001", "品名": "[組出] 口香糖王 超值包",
                    "下單數量(酷澎單位)": 60, "出貨數量": 60, "單位": "盒", "箱入數": 6}),
        base_row(**{"SKU ID": "900000000000002", "條碼(國條)": "4700000000002", "永豐料號": "M99999999", "報價備註": "M99999999", "品名": "商品總表沒有的",
                    "下單數量(酷澎單位)": 10, "出貨數量": 10, "單位": "盒", "箱入數": 10}),
        base_row(**{"SKU ID": "900000000000003", "條碼(國條)": "4712227918056", "永豐料號": "M81232885", "品名": "口香糖王 薄荷",
                    "下單數量(酷澎單位)": 25, "出貨數量": 25, "單位": "盒", "箱入數": 20}),
        base_row(**{"SKU ID": "900000000000004", "條碼(國條)": "4717986828288", "永豐料號": "M69072565", "品名": "喵愛餡 盒",
                    "下單數量(酷澎單位)": 24, "出貨數量": 24, "單位": "盒", "箱入數": 12}),
        base_row(**{"SKU ID": "900000000000005", "條碼(國條)": "4717986828289", "永豐料號": "M69072565", "品名": "喵愛餡 包",
                    "下單數量(酷澎單位)": 144, "出貨數量": 144, "單位": "包", "箱入數": 144}),
    ])
    check("匯特殊狀況那份", import_orders(c, sheet, "特殊狀況.xlsx").status_code == 200)
    conn = db.get_conn()
    q = dict(conn.execute("SELECT quote_note FROM mst_orders WHERE sku_id = '900000000000001'").fetchone()); conn.close()
    check("報價備註存進 mst_orders", q["quote_note"] == "M55500001", str(q))
    v = splits(c, "2026-09-30")
    combo = next(i for s in v["splits"] for i in s["items"] if i["sku_id"] == "900000000000001")
    check("組出商品：料號對不到 → 用報價備註對到，下採料號用報價備註那個", combo["via"] == "報價備註" and combo["purchase_code"] == "M55500001" and combo["cases"] == 10, str(combo)[:200])
    check("對不到商品總表的列出來、不進任何一份", [u["yf_sku"] for u in v["unmatched"]] == ["M99999999"]
          and all(i["yf_sku"] != "M99999999" for s in v["splits"] for i in s["items"]))
    frac = next(s for s in v["splits"] if any(i["sku_id"] == "900000000000003" for i in s["items"]))
    check("箱數 25÷20 不是整數 → 那份標成擋下來", frac["blocking"] and "不是整數" in frac["blocking"][0], str(frac["blocking"]))
    pet = [s for s in v["splits"] if s["category"] == "PET"]
    check("同料號、盒跟包：依單位拆成兩份", len(pet) == 2 and {s["unit"] for s in pet} == {"盒", "包"}, str([s["filename"] for s in pet]))
    check("箱那列才有的採購單箱備註，盒的品項不會亂帶", all(not i["po_case_note"] for s in pet for i in s["items"]))

    print("\n【5】產出拆單表")
    r = c.post("/api/mars/splits/generate", json={"from": "2026-09-30", "to": "2026-09-30"})
    check("有箱數不是整數的 → 400、點名是哪個、什麼都沒存", r.status_code == 400 and "M81232885" in json.dumps(r.get_json(), ensure_ascii=False)
          and splits(c, "2026-09-30")["splits"][0]["id"] is None, str(r.get_json())[:200])
    orow = c.get("/api/master/orders?month=2026-09&q=900000000000003").get_json()["rows"][0]
    r = c.put(f"/api/master/orders/{orow['id']}", json={"version": orow["version"], "qty_ship": 40, "reason": "缺貨"})
    check("到 ② 把出貨數量改成 40（2 箱）", r.status_code == 200, str(r.get_json())[:120])
    r = c.post("/api/mars/splits/generate", json={"from": "2026-09-30", "to": "2026-09-30"})
    check("還有對不到的 → 409 要人確認，列出是哪個", r.status_code == 409 and r.get_json()["needs_ack"] and "M99999999" in r.get_json()["details"][0])
    r = c.post("/api/mars/splits/generate", json={"from": "2026-09-30", "to": "2026-09-30", "ack_unmatched": True})
    cd = unquote(r.headers.get("Content-Disposition", ""))
    check("確認後下載 zip，檔名帶到貨日", r.status_code == 200 and "瑪氏拆單_20260930.zip" in cd, cd)
    z = zipfile.ZipFile(io.BytesIO(r.data)); names = z.namelist()
    v = splits(c, "2026-09-30")
    check("zip 裡每份兩個檔：拆單表 .xlsx ＋ _EIP上傳.xls", len(names) == 2 * len(v["splits"]) and all(f"{s['filename'][:-5]}_EIP上傳.xls" in names and s["filename"] in names for s in v["splits"]), str(names))
    check("產出後狀態：已產出、等 EIP 單號", all(s["status"] == "generated" and s["id"] for s in v["splits"]))
    gum_combo = next(s for s in v["splits"] if any(i["sku_id"] == "900000000000001" for i in s["items"]))
    wsx = openpyxl.load_workbook(io.BytesIO(z.read(gum_combo["filename"]))).active
    hdr = [x.value for x in wsx[1]]
    check("拆單表帶了 Alice 指定的商品總表欄（瑪氏貨號、Category、品名、價格、每箱中盒數、每箱產品數、單位類別、中盒貼標、備註、採購單箱備註）",
          all(h in hdr for h in ("瑪氏貨號", "Category", "品名(永豐建檔)", "系統價格(未稅)", "每箱中盒數", "每箱產品數(最小單位)", "單位類別", "中盒貼標", "備註", "採購單箱備註")), str(hdr))
    check("拆單表最後兩欄是 EIP 採購單號、約倉時間（還沒填 → 空）", hdr[-2:] == ["EIP採購單號", "約倉時間"] and wsx.cell(2, len(hdr) - 1).value in (None, ""))
    tot = [x.value for x in wsx[wsx.max_row]]
    check("最後一列合計：箱數加總", tot[0] == "合計" and tot[hdr.index("箱數")] == gum_combo["cases_total"], str(tot[:15]))
    sh = xlrd.open_workbook(file_contents=z.read(gum_combo["filename"][:-5] + "_EIP上傳.xls")).sheet_by_index(0)
    eip = [sh.row_values(i) for i in range(1, sh.nrows) if sh.cell_value(i, 1)]
    check("EIP 採購表：同一個下採料號加總、單位箱、數量是整數、備註照採購表轉換的瑪氏格式",
          {e[1]: e[6] for e in eip} == {"M55500001": 10.0, "M81232885": 2.0} and all(e[3] == "箱" for e in eip) and eip[0][8] == "MARS入倉9/30瑪氏送酷澎-TAO3倉", str(eip))
    check("對不到的 M99999999 不在任何一個檔", all("M99999999" not in json.dumps([openpyxl.load_workbook(io.BytesIO(z.read(n))).active.cell(2, 7).value]) for n in names if n.endswith(".xlsx")))

    print("\n【6】回填 EIP 採購單號、約倉時間")
    s1, s2 = v["splits"][0], v["splits"][1]
    put = lambda sid, body: c.put(f"/api/mars/splits/{sid}", json=body)  # noqa: E731
    r = put(s1["id"], {"eip_po": "PO12345"})
    check("格式不對 → 400，講要長怎樣", r.status_code == 400 and "PO202609004" in r.get_json()["error"])
    r = put(s1["id"], {"eip_po": " po202609300 "})
    check("前後空白、小寫都接受，存成 PO202609300", r.status_code == 200 and r.get_json()["split"]["eip_po"] == "PO202609300")
    r = put(s2["id"], {"eip_po": "PO202609300"})
    check("同一個單號填到第二份 → 400，說已經填在哪一份", r.status_code == 400 and s1["filename"] in r.get_json()["error"])
    r = put(s1["id"], {"slot_time": "12:30~15:30（1台車）"})
    check("約倉時間存得進去", r.status_code == 200 and r.get_json()["split"]["slot_time"] == "12:30~15:30（1台車）")
    v = splits(c, "2026-09-30"); f1 = next(s for s in v["splits"] if s["id"] == s1["id"])
    check("狀態變成已回填", f1["status"] == "filled" and f1["eip_po"] == "PO202609300")
    wsx = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/mars/splits/{s1['id']}/file?kind=split").data)).active
    check("單份重新下載：拆單表裡有 EIP 採購單號與約倉時間", wsx.cell(2, wsx.max_column - 1).value == "PO202609300" and wsx.cell(2, wsx.max_column).value == "12:30~15:30（1台車）")
    r = c.get(f"/api/mars/splits/{s1['id']}/file?kind=eip")
    check("單份 EIP 採購表也能重新下載", r.status_code == 200 and "_EIP上傳.xls" in unquote(r.headers.get("Content-Disposition", "")))
    logs = c.get("/api/master/logs?q=PO202609300&limit=1000").get_json()["logs"]
    check("回填有記修改歷程", any(l["field"] == "mars_eip_po" and l["new_value"] == "PO202609300" for l in logs))

    print("\n【7】回填之後訂單又改了：已送 EIP 的那份不偷偷改")
    target = next(i for i in f1["items"])
    orow = next(o for o in c.get("/api/master/orders?month=2026-09&q=" + target["sku_id"]).get_json()["rows"])
    r = c.put(f"/api/master/orders/{orow['id']}", json={"version": orow["version"], "qty_ship": orow["qty_ship"] + orow["box_size_file"], "reason": "酷澎要求"})
    check("到 ② 多加一箱", r.status_code == 200, str(r.get_json())[:100])
    v = splits(c, "2026-09-30"); f1 = next(s for s in v["splits"] if s["id"] == s1["id"])
    check("狀態：EIP 送出後訂單有變、寫出差在哪", f1["status"] == "changed_after_eip" and "出貨" in f1["diff"], f1.get("diff"))
    check("畫面上顯示的還是送出去的那份（舊數量）", next(i for i in f1["items"] if i["sku_id"] == target["sku_id"])["qty_ship"] == target["qty_ship"])
    c.post("/api/mars/splits/generate", json={"from": "2026-09-30", "to": "2026-09-30", "ack_unmatched": True})
    wsx = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/mars/splits/{s1['id']}/file?kind=split").data)).active
    hdr = [x.value for x in wsx[1]]
    qcol = hdr.index("出貨數量") + 1
    scol = hdr.index("SKU ID") + 1
    got = [wsx.cell(rr, qcol).value for rr in range(2, wsx.max_row + 1) if wsx.cell(rr, scol).value == target["sku_id"]]
    check("再按產出：已回填的那份內容不動（還是舊數量）", got == [target["qty_ship"]], f"{got} vs {target['qty_ship']}")
    put(s1["id"], {"eip_po": ""})
    c.post("/api/mars/splits/generate", json={"from": "2026-09-30", "to": "2026-09-30", "ack_unmatched": True})
    v = splits(c, "2026-09-30"); f1 = next(s for s in v["splits"] if s["id"] == s1["id"])
    check("把單號清掉再產出：更新成新數量、狀態回到已產出", f1["status"] == "generated" and next(i for i in f1["items"] if i["sku_id"] == target["sku_id"])["qty_ship"] == target["qty_ship"] + orow["box_size_file"])

    print("\n【8】訂單改期：拆不出來的舊拆單")
    put(s2["id"], {"eip_po": "PO202609301"})
    moved = next(i for i in next(s for s in v["splits"] if s["id"] == s2["id"])["items"])
    r = c.put("/api/master/pos/date", json={"po_number": "13000000699901", "delivery_date": "2026-09-29", "reason": "沒車"})
    check("整張 PO 改期到 9/29", r.status_code == 200, str(r.get_json())[:120])
    v = splits(c, "2026-09-30")
    gone = [s for s in v["splits"] if s["status"] == "gone"]
    check("9/30 存過的三份都標「訂單已沒有這份」（有人可能已經拿著檔），不會拆出新的", len(gone) == 3 and s2["id"] in [s["id"] for s in gone]
          and not [s for s in v["splits"] if s["status"] == "new"], str([(s['status'], s['filename']) for s in v['splits']]))
    v29 = splits(c, "2026-09-29")
    check("9/29 拆出新的（還沒產出）", v29["splits"] and all(s["status"] == "new" for s in v29["splits"]), str([s["filename"] for s in v29["splits"]]))
    r = c.post("/api/mars/splits/generate", json={"from": "2026-09-29", "to": "2026-09-30", "ack_unmatched": True})
    v = splits(c, "2026-09-29", "2026-09-30")
    check("兩天一起產出：9/30 沒回填的舊拆單刪掉，只剩已回填的那份", r.status_code == 200 and [s["id"] for s in v["splits"] if s["delivery_date"] == "2026-09-30"] == [s2["id"]]
          and all(s["status"] == "generated" for s in v["splits"] if s["delivery_date"] == "2026-09-29"), str([(s["delivery_date"], s["status"]) for s in v["splits"]]))

    print("\n【9】清除資料")
    app_module._write_users(app_module.get_users(), {"Jerry", "小真"})
    r = c.post("/api/master/reset", json={"confirm": "清空資料", "orders": True, "products": True})
    conn = db.get_conn()
    left = {t: conn.execute(f"SELECT COUNT(*) AS n FROM {t}").fetchone()["n"] for t in ("mst_mars_splits", "mst_mars_products", "mst_mars_uploads")}; conn.close()
    check("清訂單＋主檔：瑪氏拆單、商品總表一起清", r.status_code == 200 and not any(left.values()), str(left))
    r = c.post("/api/mars/splits/generate", json={"from": "2026-09-29", "to": "2026-09-30"})
    check("沒有商品總表時產出 → 400 講清楚", r.status_code == 400 and "商品總表" in r.get_json()["error"])

    print(f"\n通過 {len(PASS)} 項，失敗 {len(FAIL)} 項")
    if FAIL:
        print("失敗項目："); [print("  -", f) for f in FAIL]; sys.exit(1)


if __name__ == "__main__":
    main()
