"""蝦皮特選寄倉訂單系統（/shopee）端到端測試。
第一段：上傳採購單／入庫單、重傳差異、版本、確認、改履約方式、總覽。
第二段：價格本、竹運採購進貨／供應商直送產 EIP 採購單＋採購下採明細、除不盡改數量、回填 EIP 單號。
瑪氏直送：瑪氏 EIP 採購單（單位數量照蝦皮、依品類／單位／中標分檔）、回填後產瑪氏制式採購單、蝦皮倉資料。
竹運出貨：EC 庫存＋竹運效期表＋商品主檔 → 效期判斷、逐列扣庫存、批次與人工出貨量 → 訂單系統拋單檔＋下修刪單清單。
缺貨維護與通知：寶僑兩次、紙潔一次人工填缺貨；竹運出貨缺貨自動帶入；LINE 通知文字＋Excel 明細；通知過鎖住。
瑪氏缺貨：勇信配送明細表 PDF 對回瑪氏 EIP，逐品比箱數，寫成缺貨紀錄後一樣產 LINE 通知。

測試檔照蝦皮後台下載檔的表頭用程式造（真的業務檔不放進 repo）。
執行：python test_shopee.py（設了 DATABASE_URL 就跑 PostgreSQL）
"""
import io
import os
import sys
import tempfile
import zipfile

import openpyxl
import xlrd

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

_tmp = tempfile.mkdtemp(prefix="oms_shopee_test_")
import db  # noqa: E402
db.DATA_DIR = os.path.join(_tmp, "資料與設定")
db.DB_PATH = os.path.join(db.DATA_DIR, "test.db")
db.BACKUP_DIR = os.path.join(db.DATA_DIR, "backups")

import app as app_module  # noqa: E402
from shopee.common import line_of, split_supplier_sku  # noqa: E402
from shopee.parse import ParseError, parse_file  # noqa: E402

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("  ✓ " if cond else "  ✗ ") + name + ("" if cond or not detail else f"  — {detail}"))


PO_HDR = ["PR ID", "PO ID", "Is Jit", "Warehouse", "PR Creation Date", "PO Creation Date", "Expected Delivery Time", "Shopee SKU ID",
          "Supplier SKU ID", "Shopee SKU Name", "Shopee Requested Qty (unit)", "Shopee Requested Qty (pcs)", "Unit Name",
          "Shopee Request Price (before-tax)", "Shopee Request Price (after-tax)", "Tax %", "Shopee Requested Value (before-tax)",
          "Shopee Requested Value (after-tax)", "Supplier Confirmed Qty (unit)", "Supplier Confirmed Qty (pcs)",
          "Supplier Confirmed Price (before-tax)", "Supplier Confirmed Price (after-tax)", "Supplier Confirmed Value (after-tax)",
          "Brand", "EAN / UPC", "EAN / UPC2", "MOQ", "Sourcing Status", "Next Available Date", "Selling Type", "Inbound ID"]
IN_HDR = ["Inbound ID", "is_jit", "Shopee SKU ID", "Supplier SKU ID", "SKU Name", "PO ID", "Estimated purchase date", "Actual purchase date",
          "Estimated purchase quantity", "Warehouse acceptance quantity", "ASN Status", "Is the Payment Requested", "Warehouse",
          "Selling Type", "Expected Qty (Unit)", "Expected Qty (Pcs)", "Shortage (Pcs)", "Surplus (Pcs)", "Damaged (Pcs)",
          "Received (Pcs)", "Arrived (Pcs)", "Unit Name", "Purchase Type", "Batch ID", "Production Date", "Expiration Date", "Lifecycle"]


def yx_pdf(pages):
    """假的勇信配送明細表，文字排法照真檔（跟 test_mars.yx_pdf 一樣）：pages = [(EIP單號, 指送日期 yyyy/mm/dd, [(瑪氏貨號, 箱數, 效期), ...])]。
    同一品兩批效期寫 (瑪氏貨號, [(箱數, 效期), (箱數, 效期)])，照真檔一批印「箱數、批號、效期」（TAO5 17C→2027/11/21、2C→2027/11/22）。"""
    import pymupdf
    doc = pymupdf.open()
    for eip, date, items in pages:
        lines = ["桃園市觀音區", "客戶單號：7816797526", "BMS100230-永豐商店-蝦皮特選", "出貨廠商名稱：A5888-台灣瑪氏股份有限公司",
                 f"指送日期：{date}", "配 送 明 細 表(預排鮮度)", "產品編號", "產品名稱", "入數", "配送數量", f"收貨單號：{eip}-蝦皮-T"]
        total = 0
        for n, it in enumerate(items, start=1):
            batches = it[1] if isinstance(it[1], list) else [(it[1], it[2])]
            lines += [f"58880{it[0]}", str(n), f"假品名 {it[0]} 1:6:10"]
            for q, exp in batches:
                lines += [f"{q}C  ", "625D1TAP01", exp]; total += q
        lines += ["備註:蝦皮特選嘜頭(入庫通知單)+採購單", "單據金額:", f"合計箱數:{total}", "客戶簽章:"]
        page = doc.new_page(); y = 40
        for ln in lines:
            page.insert_text((40, y), ln, fontname="china-t", fontsize=9); y += 14
    return doc.tobytes()


def po_pdf(po, rows, text=True):
    """假的蝦皮後台採購單 PDF（照 678050_POTWCTWA2610121L1_PO.pdf：橫式、表格線、表頭每頁重複）。rows = [(Supplier SKU ID, 數量)]，一頁 8 列。
    text=False 模擬「Microsoft Print to PDF」印出來的：只有線、沒有字。"""
    import pymupdf
    doc = pymupdf.open()
    heads = ["No.", "Shopee SKU ID", "Supplier SKU ID", "Barcode", "Item Name", "Purchase QTY (unit)", "Unit Name", "Purchase QTY (pcs)", "Comment"]
    xs = [20, 45, 120, 200, 260, 400, 470, 530, 600, 700]
    for k in range(0, max(len(rows), 1), 8):
        page = doc.new_page(width=842, height=595)
        t = (lambda pt, s_, **kw: page.insert_text(pt, s_, fontname="helv", fontsize=7, **kw)) if text else (lambda *a, **kw: None)
        y0 = 40
        if k == 0:
            t((300, 30), "PURCHASE ORDER"); t((40, 50), "Purchase Order ID"); t((160, 50), po); y0 = 80
        ys = [y0, y0 + 30] + [y0 + 30 + 40 * (j + 1) for j in range(len(rows[k:k + 8]))]
        for y in ys:
            page.draw_line((xs[0], y), (xs[-1], y), width=0.8)
        for x in xs:
            page.draw_line((x, ys[0]), (x, ys[-1]), width=0.8)
        for x, h in zip(xs, heads):
            t((x + 2, y0 + 18), h)
        for j, (sku, q) in enumerate(rows[k:k + 8]):
            yy = ys[1] + 40 * j + 22
            for x, v in zip(xs, [str(k + j + 1), "123_456", sku, "471", "item", str(q), "Pieces", str(q), ""]):
                t((x + 2, yy), v)
    return doc.tobytes()


def purchase_xlsx(rows, pr="PRTWCOTWA202610010001", po=None, wh="TWA", date="2026-10-19"):
    """rows = [(Shopee SKU ID, Supplier SKU ID, 品名, 數量, Selling Type)]"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"; ws.append(PO_HDR)
    for sku, sup, name, qty, st in rows:
        r = dict.fromkeys(PO_HDR)
        r.update({"PR ID": pr, "PO ID": po, "Is Jit": "No", "Warehouse": wh, "PR Creation Date": "2026-10-01 13:16",
                  "Expected Delivery Time": date, "Shopee SKU ID": sku, "Supplier SKU ID": sup, "Shopee SKU Name": name,
                  "Shopee Requested Qty (unit)": str(qty), "Shopee Requested Qty (pcs)": str(qty), "Unit Name": "Pieces",
                  "Tax %": "5", "Brand": "NON-BRAND", "EAN / UPC": sup.split("_")[0], "MOQ": "1", "Sourcing Status": "Available",
                  "Selling Type": st})
        ws.append([r[h] for h in PO_HDR])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def inbound_xlsx(rows, po, wh="TWA", date="2026-10-19"):
    """rows = [(Inbound ID, Shopee SKU ID, Supplier SKU ID, 品名, 數量)]"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "ASN Detail"; ws.append(IN_HDR)
    for inb, sku, sup, name, qty in rows:
        r = dict.fromkeys(IN_HDR, "/")
        r.update({"Inbound ID": inb, "is_jit": "No", "Shopee SKU ID": sku, "Supplier SKU ID": sup, "SKU Name": name, "PO ID": po,
                  "Estimated purchase date": date, "Actual purchase date": "", "Estimated purchase quantity": qty,
                  "Warehouse acceptance quantity": 0, "ASN Status": "Pending Delivery", "Is the Payment Requested": "No",
                  "Warehouse": wh, "Selling Type": "Pcs", "Expected Qty (Unit)": qty, "Expected Qty (Pcs)": qty,
                  "Unit Name": "Pieces", "Purchase Type": "Consignment"})
        ws.append([r[h] for h in IN_HDR])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def price_xlsx(sheets):
    """價格本：{工作表名: [(料號3, 單位, 品名, 箱入數)]}，表頭照 PG價格本整合／紙潔價格本整合。"""
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title); ws.append(["料號", "單位", "品名", "料號3", "箱入數", "進價"])
        for code3, unit, name, box in rows:
            ws.append([code3[:-1], unit, name, code3, box, 100])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def xls_rows(data):
    sh = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
    head = [sh.cell_value(0, c) for c in range(10)]
    rows = [[sh.cell_value(r, c) for c in range(10)] for r in range(1, sh.nrows) if sh.cell_value(r, 1) != ""]
    return head, rows


PG_ROWS = [("41511463369_285962283475", "4987176232878_CNN", "幫寶適 清新幫 拉拉褲(XXL)26片", 240, "Pcs"),
           ("50403302885_340318306491", "4987176232878_RNN", "幫寶適 清新幫 拉拉褲 箱購(XXL)26片X4包", 26, "Carton"),
           ("43811463251_265962855996", "6903148358214_CNN", "幫寶適 極上守護 拉拉褲(L+)30片", 216, "Pcs")]


def main():
    db.init_db()
    app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()
    with c.session_transaction() as s:
        s["user"] = "Jerry"

    def preview(data, name):
        return c.post("/api/shopee/import/preview", data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")

    def orders(**q):
        q.setdefault("from", "2026-01-01")
        return c.get("/api/shopee/orders", query_string=q).get_json()

    print("\n【0】頁面、首頁卡片、選單")
    page = c.get("/shopee").get_data(as_text=True)
    check("蝦皮特選頁打得開、蝦皮橘、一個上傳入口、作業流程格子、訂單清單，沒有舊的六個頁籤（Jerry 2026-10-13 改版）",
          "蝦皮特選訂單管理系統" in page and "#ee4d2d" in page and 'id="dz"' in page and 'id="stages"' in page and 'id="tbl"' in page
          and 'id="tabs"' not in page and 'data-tab="zhu"' not in page and page.count('class="dzbox"') == 1 and 'id="btn-set"' in page and 'id="view-set"' in page)
    home = c.get("/").get_data(as_text=True)
    check("首頁左側選單的蝦皮特選連到 /shopee", 'href="/shopee"' in home)
    ch = next(x for x in c.get("/api/portal/summary").get_json()["channels"] if x["key"] == "shopee")
    check("首頁蝦皮卡片：上線、連 /shopee、還沒資料時數字是 0、沒有箱數", ch["live"] and ch["url"] == "/shopee" and ch["new_today"] == 0 and ch["month_boxes"] is None, str(ch))

    print("\n【1】讀檔：認檔案種類、線別、單位代碼")
    check("料號拆底線：國條、瑪氏帶 -1、沒有底線", split_supplier_sku("4987176232878_CNN ") == ("4987176232878", "C")
          and split_supplier_sku("M60002536-1_CMN") == ("M60002536-1", "C") and split_supplier_sku("ABC") == ("ABC", ""))
    check("線別：M 開頭瑪氏、7 碼紙潔、13 碼與 14 碼寶僑、其他認不出", line_of("M10273409") == "瑪氏" and line_of("1022257") == "紙潔"
          and line_of("4987176232878") == "寶僑" and line_of("14987176176172") == "寶僑" and line_of("DH56M") == "")
    d = parse_file(purchase_xlsx(PG_ROWS), "PurchaseOrder_x.xlsx")
    check("採購單：看表頭認得、3 列、單位代碼 C＝包 R＝箱、到貨日、數量轉數字", d["kind"] == "purchase" and len(d["rows"]) == 3
          and [r["unit_name"] for r in d["rows"]] == ["包", "箱", "包"] and d["rows"][0]["expected_date"] == "2026-10-19"
          and d["rows"][0]["qty"] == 240 and d["lines"] == {"寶僑": 3}, str(d["lines"]))
    mars = parse_file(purchase_xlsx([("1_1", "M10273409_QMN", "M&M'S", 48, "Pcs"), ("1_2", "DH56M_QMN", "士力架 5入", 45, "Pcs")]), "x.xlsx")
    check("認不出線別的（DH56M）用同檔多數補、還提醒", [r["line"] for r in mars["rows"]] == ["瑪氏", "瑪氏"] and any("DH56M" in e for e in mars["errors"]), str(mars["errors"]))
    ind = parse_file(inbound_xlsx([("IN1", "s1", "1022257_PNN", "橘子工坊", 30), ("IN2", "s1", "1022257_PNN", "橘子工坊", 18)], "POX"), "InboundOrder_x.xlsx")
    check("入庫單：同一張 PO 同一品項兩列（雙效期）合成一列，數量加總、入庫單號串起來", ind["kind"] == "inbound" and len(ind["rows"]) == 1
          and ind["rows"][0]["qty"] == 48 and ind["rows"][0]["inbound_id"] == "IN1,IN2" and ind["rows"][0]["line"] == "紙潔", str(ind["rows"]))
    wb = openpyxl.Workbook(); wb.active.append(["隨便", "一個", "表"]); b = io.BytesIO(); wb.save(b)
    for name, data in (("不是蝦皮的表", b.getvalue()), ("壞掉的檔", b"not excel")):
        try:
            parse_file(data, "x.xlsx"); ok = False
        except ParseError as e:
            ok = "Excel" in str(e) or "表頭" in str(e)
        check(f"{name} → 講清楚、不出現英文錯誤", ok)
    r = preview(b"not excel", "x.xlsx")
    check("上傳壞檔 → 400", r.status_code == 400 and "Traceback" not in r.get_data(as_text=True))

    print("\n【2】第一次上傳採購單（PR 階段，還沒有 PO）")
    raw = purchase_xlsx(PG_ROWS)
    pv = preview(raw, "PurchaseOrder_20261001.xlsx").get_json()
    check("預覽：採購單、新增 3、要選履約方式、還沒寫進訂單", pv["kind"] == "purchase" and pv["counts"]["new"] == 3 and pv["needs_fulfil"]
          and orders()["summary"]["orders"] == 0, str(pv["counts"]))
    check("還沒匯入：全部訂單 0 張（畫面據此寫「尚未匯入蝦皮訂單」）", orders()["all_orders"] == {"count": 0, "first": "", "last": ""}, str(orders()["all_orders"]))
    pend = c.get("/api/shopee/import/pending").get_json()["pending"]
    check("還沒按確認匯入：重新整理後拿得回這張預覽（同一個上傳、新增 3、要選履約方式、線別寶僑 3 項）",
          len(pend) == 1 and pend[0]["upload_id"] == pv["upload_id"] and pend[0]["counts"]["new"] == 3 and pend[0]["needs_fulfil"]
          and pend[0]["lines"] == pv["lines"] and pend[0]["orders"] == pv["orders"], str(pend)[:300])
    with c.session_transaction() as s_:
        s_["user"] = "Alice"
    check("別人的預覽不會跑到自己畫面上", c.get("/api/shopee/import/pending").get_json()["pending"] == [])
    with c.session_transaction() as s_:
        s_["user"] = "Jerry"
    r = c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"]})
    check("沒選履約方式 → 400", r.status_code == 400 and "履約方式" in r.get_json()["error"])
    r = c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"], "fulfil": "竹運採購進貨"})
    check("選了竹運採購進貨 → 匯入 3 項", r.status_code == 200 and r.get_json()["counts"]["new"] == 3, str(r.get_json()))
    check("確認匯入後就不再列為待確認", c.get("/api/shopee/import/pending").get_json()["pending"] == [])
    check("同一個預覽再按一次 → 404（不會重複匯入）", c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"], "fulfil": "竹運出貨"}).status_code == 404)
    o = orders()
    g = o["orders"][0]
    out = orders(**{"from": "2027-01-01"})
    check("到貨日篩掉全部：清單 0 張，但告訴畫面系統共有 1 張、到貨日 10/19（畫面據此給「清除篩選」）",
          out["orders"] == [] and out["all_orders"] == {"count": 1, "first": "2026-10-19", "last": "2026-10-19"}, str(out["all_orders"]))
    check("總覽：一張單、訂單編號先用 PR、線別寶僑、倉 TWA、到貨日、履約方式、原始訂購量 482", len(o["orders"]) == 1 and g["order_no"] == "PRTWCOTWA202610010001"
          and g["line"] == "寶僑" and g["warehouse"] == "TWA" and g["expected_date"] == "2026-10-19" and g["fulfil"] == "竹運採購進貨"
          and g["qty_original"] == 482 and g["status"] == "已匯入", str(g)[:300])
    ups = c.get("/api/shopee/uploads").get_json()["uploads"]
    check("上傳紀錄：檔名、上傳人員、匯入結果", ups[0]["filename"] == "PurchaseOrder_20261001.xlsx" and ups[0]["uploaded_by"] == "Jerry" and ups[0]["new_count"] == 3)
    r = c.get(f"/api/shopee/uploads/{ups[0]['id']}/file")
    check("原始檔下載回來跟上傳的一模一樣", r.status_code == 200 and r.data == raw)
    logs = c.get("/api/master/logs?scope=shopee&limit=50").get_json()["logs"]
    check("歷程視窗（本頁）看得到這次匯入，系統欄是蝦皮特選訂單管理系統", any(l["field"] == "shopee_import" for l in logs)
          and all(l.get("system") in (None, "蝦皮特選訂單管理系統") for l in logs), str([l["field"] for l in logs]))

    print("\n【3】同一份再傳一次：內容完全一樣")
    pv = preview(raw, "PurchaseOrder_20261001.xlsx").get_json()
    check("預覽：3 項無異動、沒有差異、不用選履約方式", pv["counts"]["same"] == 3 and not pv["diffs"] and not pv["needs_fulfil"], str(pv["counts"]))
    r = c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"]})
    check("確認後只記上傳紀錄，訂單還是已匯入", r.status_code == 200 and orders()["orders"][0]["status"] == "已匯入"
          and len(c.get("/api/shopee/uploads").get_json()["uploads"]) == 2)

    print("\n【4】重傳：變成 PO、數量改、到貨日改、少一個品項")
    rows2 = [PG_ROWS[0][:3] + (200, "Pcs"), PG_ROWS[1]]           # 第一項 240→200、第三項不見了
    pv = preview(purchase_xlsx(rows2, po="POTWCTWA26100514L", date="2026-10-20"), "PurchaseOrder_20261005.xlsx").get_json()
    kinds = sorted((x["label"], x["old"], x["new"]) for x in pv["diffs"])
    check("預覽列出差異：數量 240→200（-40）、到貨日 ×2、品項移除；PR 變 PO 不算差異", ("數量", 240, 200) in kinds
          and sum(1 for k in kinds if k[0] == "到貨日") == 2 and any(k[0] == "品項" for k in kinds) and not any(k[0] == "PO 號碼" for k in kinds)
          and any(x["qty_delta"] == -40 for x in pv["diffs"]) and pv["counts"]["remove"] == 1, str(kinds))
    c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"]})
    o = orders(); g = o["orders"][0]
    check("總覽：訂單編號換成 PO、PR 還留著、匯入差異待確認、目前數量 226、原始訂購量還是 482、移除 1 項",
          g["order_no"] == "POTWCTWA26100514L" and g["pr_id"] == "PRTWCOTWA202610010001" and g["status"] == "匯入差異待確認"
          and g["qty"] == 226 and g["qty_original"] == 482 and g["removed_count"] == 1 and o["pending_orders"] == 1, str(g)[:300])
    st = {x["key"]: x["count"] for x in o["stages"]}
    check("作業流程：差異待確認的單排第一步，「下一步」寫確認差異，格子數字對得上", g["todo"][:1] == ["diff"] and g["next"] == "確認差異"
          and st["diff"] == 1 and [x["key"] for x in o["stages"]] == ["diff", "buy", "ship", "eip", "notify", "done"], str((g["todo"], st)))
    it = next(i for i in g["items"] if i["supplier_sku_id"] == "4987176232878_CNN")
    vs = c.get(f"/api/shopee/versions/{it['id']}").get_json()["versions"]
    check("版本紀錄：舊版不刪，第 1、2 版都在（內容一樣的那次重傳不另記一版），第 1 版數量 240、最新版 200", [v["version"] for v in vs] == [2, 1]
          and vs[-1]["data"]["qty"] == 240 and vs[0]["data"]["qty"] == 200, str([(v["version"], v["data"]["qty"]) for v in vs]))
    dd = c.get("/api/shopee/diffs", query_string={"order_no": "POTWCTWA26100514L"}).get_json()["diffs"]
    check("差異視窗：訂單編號、商品料號、欄位、上一版、新版、數量增減、上傳人員、時間、待確認", dd and all(
        x["order_no"] == "POTWCTWA26100514L" and x["supplier_sku_id"] and x["field_label"] and x["uploaded_by"] == "Jerry" and x["uploaded_at"]
        and not x["confirmed"] for x in dd), str(dd[0]) if dd else "")
    t = c.get("/api/portal/summary").get_json()["tasks"]
    check("首頁今日事項出現「蝦皮特選重傳後有差異 1 張單」", any(x["key"] == "shopee_diff" and x["count"] == 1 and x["url"] == "/shopee#diff" for x in t))
    check("篩「只看差異待確認」只剩這張", [x["order_no"] for x in orders(pending="1")["orders"]] == ["POTWCTWA26100514L"])

    print("\n【5】確認差異")
    r = c.post("/api/shopee/diffs/confirm", json={"order_nos": ["POTWCTWA26100514L"]})
    g = orders()["orders"][0]
    dd = c.get("/api/shopee/diffs", query_string={"order_no": "POTWCTWA26100514L"}).get_json()["diffs"]
    check("確認後回到已匯入、差異記下確認人和時間", r.status_code == 200 and g["status"] == "已匯入" and all(x["confirmed"] and x["confirmed_by"] == "Jerry" for x in dd))
    check("首頁今日事項的蝦皮差異消失", not any(x["key"] == "shopee_diff" for x in c.get("/api/portal/summary").get_json()["tasks"]))
    check("沒選訂單就按確認 → 400", c.post("/api/shopee/diffs/confirm", json={"order_nos": []}).status_code == 400)

    print("\n【6】入庫單：補入庫單號、數量不一樣算差異；只有入庫單的 PO 直接建")
    inb = inbound_xlsx([("INTWA0002610050005", PG_ROWS[0][0], PG_ROWS[0][1], PG_ROWS[0][2], 190),
                        ("INTWA0002610050006", PG_ROWS[1][0], PG_ROWS[1][1], PG_ROWS[1][2], 26)], "POTWCTWA26100514L", date="2026-10-20")
    pv = preview(inb, "InboundOrder_20261006.xlsx").get_json()
    check("入庫單預覽：找到採購單那兩筆、數量 200→190 算差異、不用選履約方式", pv["kind"] == "inbound" and pv["counts"]["new"] == 0
          and [(x["label"], x["old"], x["new"]) for x in pv["diffs"]] == [("數量", 200, 190)] and not pv["needs_fulfil"], str(pv["diffs"]))
    c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"]})
    g = orders()["orders"][0]
    it = next(i for i in g["items"] if i["supplier_sku_id"] == "4987176232878_CNN")
    check("補上入庫單號、入庫單數量 190，又變差異待確認", it["inbound_ids"] == "INTWA0002610050005" and it["inbound_qty"] == 190 and g["status"] == "匯入差異待確認", str(it)[:200])
    c.post("/api/shopee/diffs/confirm", json={"order_nos": ["POTWCTWA26100514L"]})
    only = inbound_xlsx([("INTWA0002610050266", "56402432052_307219345020", "M60003196_QMN", "希寶 誘惑泥 55入", 48),
                         ("INTWA0002610050265", "53902436305_360235791151", "M10281171_RMN", "偉嘉 貓罐頭 箱購", 18)], "POTWCTWA26100514N", date="2026-10-13")
    pv = preview(only, "InboundOrder_瑪氏.xlsx").get_json()
    check("只有入庫單、找不到採購單 → 新增 2 項、要選履約方式", pv["counts"]["new"] == 2 and pv["needs_fulfil"])
    c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"], "fulfil": "供應商直送"})
    g2 = next(x for x in orders()["orders"] if x["order_no"] == "POTWCTWA26100514N")
    check("建成一張 PO：瑪氏、供應商直送、2 項、單位 盒／箱", g2["line"] == "瑪氏" and g2["fulfil"] == "供應商直送" and g2["item_count"] == 2
          and sorted(i["unit_name"] for i in g2["items"]) == ["盒", "箱"])
    pr_only = purchase_xlsx([("70000000001_1", "1022257_PNN", "橘子工坊浴廁清潔劑", 360, "Pcs")], pr="PRTWCOTWX202610010777", wh="TWX", date="2026-10-21")
    c.post("/api/shopee/import/commit", json={"upload_id": preview(pr_only, "PurchaseOrder_紙潔.xlsx").get_json()["upload_id"], "fulfil": "竹運出貨"})
    pv = preview(inbound_xlsx([("INTWX1", "70000000001_1", "1022257_PNN", "橘子工坊浴廁清潔劑", 360)], "POTWCTWX26100599A", wh="TWX", date="2026-10-21"), "InboundOrder_紙潔.xlsx").get_json()
    check("採購單還在 PR 階段時，入庫單用「同商品、同倉、同到貨日」對回去，不另外建", pv["counts"]["new"] == 0 and not pv["diffs"], str(pv["counts"]))
    c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"]})
    g3 = next(x for x in orders()["orders"] if x["pr_id"] == "PRTWCOTWX202610010777")
    check("那張 PR 補上 PO 號碼、入庫單號，線別紙潔", g3["order_no"] == "POTWCTWX26100599A" and g3["items"][0]["inbound_ids"] == "INTWX1" and g3["line"] == "紙潔")

    print("\n【7】改履約方式：原因選填、可以好幾張一起改（Alice 2026-10-21）、留紀錄")
    r0 = c.post("/api/shopee/fulfil", json={"order_no": "POTWCTWA26100514N", "fulfil": "竹運出貨", "reason": ""})
    check("沒寫原因也能改（原因選填）", r0.status_code == 200 and r0.get_json()["changed"] and r0.get_json()["count"] == 1, str(r0.get_json()))
    check("沒選訂單 → 400", c.post("/api/shopee/fulfil", json={"order_nos": [], "fulfil": "竹運出貨"}).status_code == 400)
    rb = c.post("/api/shopee/fulfil", json={"order_nos": ["POTWCTWA26100514N", "NOPE"], "fulfil": "供應商直送", "reason": "改回去"})
    check("好幾張一起改：找得到的改、找不到的列在 missing；全部找不到才 404", rb.status_code == 200 and rb.get_json()["count"] == 1 and rb.get_json()["missing"] == ["NOPE"]
          and c.post("/api/shopee/fulfil", json={"order_nos": ["NOPE", "NOPE2"], "fulfil": "供應商直送"}).status_code == 404, str(rb.get_json()))
    check("亂填履約方式 → 400", c.post("/api/shopee/fulfil", json={"order_no": "POTWCTWA26100514N", "fulfil": "宅配", "reason": "x"}).status_code == 400)
    r = c.post("/api/shopee/fulfil", json={"order_no": "POTWCTWA26100514N", "fulfil": "竹運出貨", "reason": "竹運有現貨"})
    g2 = next(x for x in orders()["orders"] if x["order_no"] == "POTWCTWA26100514N")
    logs = [l for l in c.get("/api/master/logs?scope=shopee&limit=100").get_json()["logs"] if l["field"] == "shopee_fulfil"]
    check("改成竹運出貨，歷程記下原本、修改後、修改人員、時間、原因", r.get_json()["changed"] and g2["fulfil"] == "竹運出貨" and logs
          and logs[0]["old_value"] == "供應商直送" and logs[0]["new_value"] == "竹運出貨" and logs[0]["reason"] == "竹運有現貨"
          and logs[0]["operator"] == "Jerry" and logs[0]["changed_at"], str(logs[:1]))
    check("找不到的單 → 404", c.post("/api/shopee/fulfil", json={"order_no": "NOPE", "fulfil": "竹運出貨", "reason": "x"}).status_code == 404)

    print("\n【8】總覽篩選、首頁數字")
    check("篩線別：瑪氏只剩一張", [x["order_no"] for x in orders(line="瑪氏")["orders"]] == ["POTWCTWA26100514N"])
    check("篩履約方式：竹運採購進貨只剩寶僑那張", [x["order_no"] for x in orders(fulfil="竹運採購進貨")["orders"]] == ["POTWCTWA26100514L"])
    check("搜尋入庫單號找得到那張", [x["order_no"] for x in orders(q="INTWX1")["orders"]] == ["POTWCTWX26100599A"])
    check("到貨日區間：10/13 只剩瑪氏那張", [x["order_no"] for x in orders(**{"from": "2026-10-13", "to": "2026-10-13"})["orders"]] == ["POTWCTWA26100514N"])
    # 上傳日篩選、每張單的上傳日和永豐料號（OP 2026-10-21，Jerry 看過示意圖）
    today_ = db.now()[:10]
    all_ = orders()["orders"]
    check("上傳日＝那張單第一次進系統那天：今天上傳的每張都是今天；永豐料號＝Supplier SKU ID 清單（沒移除的品項）",
          all_ and all(x["upload_date"] == today_ and x["codes"] and all("_" in cd for cd in x["codes"]) for x in all_), str([(x["upload_date"], x["codes"][:2]) for x in all_[:2]]))
    check("上傳日篩選：今天～今天全部都在；2030 以後一張都沒有；只填起日也行",
          [x["order_no"] for x in orders(up_from=today_, up_to=today_)["orders"]] == [x["order_no"] for x in all_]
          and orders(up_from="2030-01-01")["orders"] == [] and orders(up_to="2000-01-01")["orders"] == []
          and len(orders(up_from="2000-01-01")["orders"]) == len(all_))
    check("其他步驟的 API 也吃上傳日：採購計畫、缺貨計畫、EIP 檔清單用 2030 以後都空的",
          c.get("/api/shopee/purchase/plan", query_string={"fulfil": "供應商直送", "from": "2026-01-01", "up_from": "2030-01-01"}).get_json()["orders"] == []
          and c.get("/api/shopee/short/plan", query_string={"from": "2026-01-01", "up_from": "2030-01-01"}).get_json()["orders"] == []
          and c.get("/api/shopee/eips", query_string={"from": "2026-01-01", "up_from": "2030-01-01"}).get_json()["eips"] == [])
    pl_ = c.get("/api/shopee/purchase/plan", query_string={"fulfil": "供應商直送", "from": "2026-01-01", "up_to": today_}).get_json()["orders"]
    check("採購計畫那邊每張單也有上傳日、永豐料號", all(g["upload_date"] == today_ and g["codes"] for g in pl_), str(pl_[:1])[:200])
    ch = next(x for x in c.get("/api/portal/summary").get_json()["channels"] if x["key"] == "shopee")
    check("首頁蝦皮卡片：今日新進 3 張、有最後匯入時間", ch["new_today"] == 3 and ch["last_import"], str(ch))
    recent = c.get("/api/portal/summary").get_json()["recent"]
    check("首頁最近操作紀錄標出蝦皮特選訂單管理系統", any(x.get("system") == "蝦皮特選訂單管理系統" for x in recent), str([x.get("system") for x in recent]))
    r = c.post("/api/shopee/import/cancel", json={"upload_id": preview(raw, "x.xlsx").get_json()["upload_id"]})
    check("預覽後按取消 → 不寫訂單、不進上傳紀錄", r.status_code == 200 and len(c.get("/api/shopee/uploads").get_json()["uploads"]) == 7,
          str(len(c.get("/api/shopee/uploads").get_json()["uploads"])))

    stage2(c, preview, orders)

    print("\n" + "=" * 62)
    print(f"通過 {len(PASS)} 項，失敗 {len(FAIL)} 項")
    if FAIL:
        print("失敗項目："); [print("  -", n) for n in FAIL]
        sys.exit(1)


def stage2(c, preview, orders):
    def up(url, data, name):
        return c.post(url, data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")

    def imp(data, fulfil):
        return c.post("/api/shopee/import/commit", json={"upload_id": preview(data, "PurchaseOrder_t.xlsx").get_json()["upload_id"], "fulfil": fulfil}).get_json()

    def plan(fulfil):
        return {g["order_no"]: g for g in c.get("/api/shopee/purchase/plan", query_string={"fulfil": fulfil}).get_json()["orders"]}

    def gen(fulfil, nos):
        return c.post("/api/shopee/purchase/generate", json={"fulfil": fulfil, "order_nos": nos})

    def item(g, sup):
        return next(p for p in g["items"] if p["supplier_sku_id"] == sup)

    print("\n【9】價格本（轉換率）")
    page = c.get("/shopee").get_data(as_text=True)
    check("待產採購單、待回填 EIP 採購單號兩步的清單；價格本狀態在資料與設定；修改數量視窗", all(x in page for x in (
        'id="st-buy"', 'id="b-tbl"', 'id="st-eip"', 'id="e-tbl"', 'id="conv-st"', 'id="dlg-adj"')))
    pg = price_xlsx({"價格本整合": [("4987176232878C", "包", "幫寶適 清新幫 拉拉褲", 4), ("4987176232878R", "箱", "幫寶適 清新幫 拉拉褲 箱購", 1),
                                   ("6903148358214C", "包", "幫寶適 極上守護 拉拉褲", 3)]})
    r = up("/api/shopee/conv/import", pg, "PG價格本整合.xlsx").get_json()
    check("上傳寶僑價格本：寶僑 3 筆", r["ok"] and r["lines"] == {"寶僑": {"count": 3, "before": 0, "blank": 0}}, str(r))
    paper = price_xlsx({"紙品價格本": [("1099999P", "包", "紙品 沒填箱入數", None)],
                        "潔品價格本": [("1022257P", "瓶", "橘子工坊浴廁清潔劑", 12), ("1022248C", "包", "橘子工坊洗碗精補充包", 12),
                                      ("1021315R", "箱", "橘子工坊洗衣精補充包 箱購", 1), ("AAA0001C", "包", "組合品", 6)],
                        "工作表4": [("1022257P", "瓶", "兩頁合起來的，不另外讀", 99)]})
    r = up("/api/shopee/conv/import", paper, "紙潔價格本整合.xlsx").get_json()
    check("上傳紙潔價格本：組合品料號（AAA…）也算紙潔、「工作表4」不重複讀、箱入數空白 1 筆", r["lines"] == {"紙潔": {"count": 5, "before": 0, "blank": 1}}
          and not r["warnings"], str(r))
    st = c.get("/api/shopee/conv").get_json()["lines"]
    check("價格本狀態：兩個線別、筆數、檔名、上傳人員", st["寶僑"]["count"] == 3 and st["紙潔"]["count"] == 5 and st["紙潔"]["file"] == "紙潔價格本整合.xlsx"
          and st["寶僑"]["by"] == "Jerry", str(st))
    r = up("/api/shopee/conv/import", pg, "PG價格本整合.xlsx").get_json()
    check("同一線別再傳一次：整份換掉，不會變兩倍", r["lines"]["寶僑"] == {"count": 3, "before": 3, "blank": 0}, str(r["lines"]))
    wb = openpyxl.Workbook(); wb.active.append(["隨便", "一個", "表"]); b = io.BytesIO(); wb.save(b)
    r = up("/api/shopee/conv/import", b.getvalue(), "x.xlsx")
    check("不是價格本 → 400、講清楚要哪個欄位", r.status_code == 400 and "料號3" in r.get_json()["error"])
    check("歷程記下價格本更新", any(l["field"] == "shopee_conv" for l in c.get("/api/master/logs?scope=shopee&limit=200").get_json()["logs"]))

    print("\n【10】竹運採購進貨：除不盡要先改數量、整批分箱單位／小單位兩個檔")
    imp(purchase_xlsx([("80000000001_1", "6903148358214_CNN", "幫寶適 極上守護 拉拉褲(L+)30片", 30, "Pcs")],
                      po="POTWCTWX26100700B", pr="PRTWCOTWX202610070001", wh="TWX", date="2026-10-21"), "竹運採購進貨")
    pl = plan("竹運採購進貨")
    L, B = pl.get("POTWCTWA26100514L"), pl.get("POTWCTWX26100700B")
    p1 = item(L, "4987176232878_CNN") if L else {}
    check("計畫：兩張寶僑單；190 包除以轉換率 4 除不盡（餘 2），擋下來", L and B and L["blocking"] == 1 and p1["rem"] == 2 and p1["blocking"]
          and "除不盡" in p1["problems"][0] and not B["blocking"], str(p1)[:300])
    p2 = item(L, "4987176232878_RNN")
    check("轉換率 1 的放箱單位、不是 1 的放小單位；竹運進貨單位數量照蝦皮", p2["unit_group"] == "箱單位" and (p2["out_unit"], p2["out_qty"]) == ("箱", 26)
          and item(B, "6903148358214_CNN")["unit_group"] == "小單位" and (item(B, "6903148358214_CNN")["out_unit"], item(B, "6903148358214_CNN")["out_qty"]) == ("包", 30))
    r = gen("竹運採購進貨", ["POTWCTWA26100514L", "POTWCTWX26100700B"])
    check("有除不盡的品項 → 整批不產（400），列出哪一品、什麼原因", r.status_code == 400 and any("4987176232878_CNN" in d and "除不盡" in d for d in r.get_json()["details"]), str(r.get_json()))
    check("沒勾訂單 → 400；履約方式不對 → 400", gen("竹運採購進貨", []).status_code == 400 and gen("竹運出貨", ["POTWCTWA26100514L"]).status_code == 400)
    oid = p1["order_id"]
    r = c.post("/api/shopee/purchase/adjust", json={"order_id": oid, "qty": 191, "reason": "x"})
    check("改成 191 還是除不盡 → 400，提示 188 或 192", r.status_code == 400 and "188" in r.get_json()["error"] and "192" in r.get_json()["error"], r.get_json()["error"])
    check("沒寫原因 → 400", c.post("/api/shopee/purchase/adjust", json={"order_id": oid, "qty": 188, "reason": ""}).status_code == 400)
    r = c.post("/api/shopee/purchase/adjust", json={"order_id": oid, "qty": 188, "reason": "與業務確認改為整箱"})
    p1 = item(plan("竹運採購進貨")["POTWCTWA26100514L"], "4987176232878_CNN")
    logs = [l for l in c.get("/api/master/logs?scope=shopee&limit=200").get_json()["logs"] if l["field"] == "shopee_qty_adj"]
    check("改成 188：不再擋、下採數量 188、留原始數量、轉換率、人員、時間、原因", r.status_code == 200 and not p1["blocking"] and p1["use_qty"] == 188
          and p1["qty"] == 190 and p1["adj"]["qty_before"] == 190 and p1["adj"]["conv"] == 4 and p1["adj"]["operator"] == "Jerry" and p1["adj"]["updated_at"]
          and logs and logs[0]["old_value"] == "190" and logs[0]["new_value"] == "188" and logs[0]["reason"] == "與業務確認改為整箱", str(p1)[:300])
    r = gen("竹運採購進貨", ["POTWCTWA26100514L", "POTWCTWX26100700B"])
    z = zipfile.ZipFile(io.BytesIO(r.data)) if r.status_code == 200 else None
    names = sorted(z.namelist()) if z else []
    z = z if len(names) == 3 else None
    check("產出 zip：整批一組（跨 10/20、10/21 和兩個倉）＝箱單位、小單位兩個 EIP 檔＋一份下採明細；檔名寫該檔實際的日期和倉", names == sorted([
        "產品採購表上傳-寶僑竹運進貨_1020到觀音-箱單位.xls", "產品採購表上傳-寶僑竹運進貨_1020-1021到觀音、安南三-小單位.xls",
        "蝦皮特選-商品下採_寶僑竹運進貨_1020-1021到觀音、安南三.xlsx"]) and r.headers["X-Shopee-Eip-Files"] == "2", str(names))
    if z:
        head, rows = xls_rows(z.read("產品採購表上傳-寶僑竹運進貨_1020-1021到觀音、安南三-小單位.xls"))
        check("小單位檔：料號、蝦皮品名、單位包、數量用改過的 188、備註蝦皮特選、第 10 欄放轉換率且標題空白",
              head[:4] == ["項目", "料號", "品名", "單位"] and head[9] == "" and [r_[1] for r_ in rows] == ["4987176232878", "6903148358214"]
              and [(r_[3], r_[6], r_[8], r_[9]) for r_ in rows] == [("包", 188, "蝦皮特選", 4), ("包", 30, "蝦皮特選", 3)], str(rows))
        head, rows = xls_rows(z.read("產品採購表上傳-寶僑竹運進貨_1020到觀音-箱單位.xls"))
        check("箱單位檔：箱 26、轉換率 1；標題「採購數量」帶總數（跟酷澎一樣，Jerry 2026-10-28）", [(r_[1], r_[3], r_[6], r_[9]) for r_ in rows] == [("4987176232878", "箱", 26, 1)]
              and head[6] == "採購數量26", str((head[6], rows)))
        wb = openpyxl.load_workbook(io.BytesIO(z.read("蝦皮特選-商品下採_寶僑竹運進貨_1020-1021到觀音、安南三.xlsx")))
        small = [list(r_) for r_ in wb["小單位"].iter_rows(values_only=True)]
        check("下採明細：箱、小單位兩頁；欄位永豐料號／品名／單位／進貨／轉換率／數量每箱（公式）", wb.sheetnames == ["箱", "小單位"]
              and small[0] == ["永豐料號", "品名", "單位", "進貨", "轉換率", "數量/箱"] and small[1][3] == 188 and small[1][5] == "=D2/E2"
              and [list(r_) for r_ in wb["箱"].iter_rows(values_only=True)][1][3] == 26, str(small))
    pl = plan("竹運採購進貨")
    ov = next(g for g in orders()["orders"] if g["order_no"] == "POTWCTWA26100514L")
    check("產出後：計畫和訂單總覽的採購單產出狀態都是「已產出」", pl["POTWCTWA26100514L"]["po_status"] == "已產出" and ov["po_status"] == "已產出", ov["po_status"])
    eips = c.get("/api/shopee/eips").get_json()["eips"]
    small = next(e for e in eips if e["unit_group"] == "小單位")
    check("採購單產出清單：2 個檔、倉別兩個倉、對應的蝦皮訂單", len(eips) == 2 and small["warehouse_name"] == "觀音、安南三"
          and small["orders"] == ["POTWCTWA26100514L", "POTWCTWX26100700B"] and small["rows"] == 2 and small["qty"] == 218, str(small))
    r = gen("竹運採購進貨", ["POTWCTWA26100514L", "POTWCTWX26100700B"])
    check("還沒回填單號可以重產，舊的換掉不會變 4 個", r.status_code == 200 and len(c.get("/api/shopee/eips").get_json()["eips"]) == 2)
    eips = c.get("/api/shopee/eips").get_json()["eips"]
    small = next(e for e in eips if e["unit_group"] == "小單位"); box = next(e for e in eips if e["unit_group"] == "箱單位")
    bad = c.put(f"/api/shopee/eips/{small['id']}", json={"eip_po": "12345"})
    check("回填單號格式不對 → 400", bad.status_code == 400 and "PO" in bad.get_json()["error"])
    r = c.put(f"/api/shopee/eips/{small['id']}", json={"eip_po": "po202610003"})
    check("回填 PO202610003（小寫也行）", r.status_code == 200 and r.get_json()["eip_po"] == "PO202610003")
    check("同一個單號填到別張 → 400", c.put(f"/api/shopee/eips/{box['id']}", json={"eip_po": "PO202610003"}).status_code == 400)
    r = c.get(f"/api/shopee/eips/{small['id']}/file")
    check("重新下載：檔名帶 EIP 單號、內容一樣", r.status_code == 200 and "PO202610003" in r.headers["Content-Disposition"]
          and xls_rows(r.data)[1][0][6] == 188)
    pl = plan("竹運採購進貨")
    p1 = item(pl["POTWCTWA26100514L"], "4987176232878_CNN")
    check("回填後鎖住：那張檔裡的品項不能重產，說要先清單號", p1["blocking"] and "已回填" in p1["problems"][0]
          and gen("竹運採購進貨", ["POTWCTWA26100514L"]).status_code == 400, str(p1["problems"]))
    check("箱單位那張還沒回填 → 訂單狀態「已產出」；兩張都回填才算「已回填」", pl["POTWCTWA26100514L"]["po_status"] == "已產出" and pl["POTWCTWX26100700B"]["po_status"] == "已回填")
    c.put(f"/api/shopee/eips/{small['id']}", json={"eip_po": ""})
    check("清掉單號就能重產", gen("竹運採購進貨", ["POTWCTWA26100514L", "POTWCTWX26100700B"]).status_code == 200)
    logs = c.get("/api/master/logs?scope=shopee&limit=300").get_json()["logs"]
    check("歷程記下產出 EIP 和回填單號", any(l["field"] == "shopee_eip" and l["po_number"] == "POTWCTWA26100514L" for l in logs)
          and any(l["field"] == "shopee_eip_po" and l["new_value"] == "PO202610003" for l in logs))
    r = c.post("/api/shopee/purchase/adjust", json={"order_id": oid, "qty": ""})
    check("修改數量清空＝取消修改，又變回除不盡", r.get_json().get("cleared") and item(plan("竹運採購進貨")["POTWCTWA26100514L"], "4987176232878_CNN")["rem"] == 2)

    print("\n【11】供應商直送：一天一倉一個檔；寶僑小單位換箱、紙潔照蝦皮原本的單位（需求 10.2；紙潔 Jerry 2026-10-29 改回原單位）")
    imp(purchase_xlsx([("81_1", "1022257_PNN", "橘子工坊浴廁清潔劑480ml", 360, "Pcs"), ("81_2", "1022248_CNN", "橘子工坊洗碗精補充包500ml", 1392, "Pcs"),
                       ("81_3", "1021315_RNN", "橘子工坊洗衣精補充包 箱購", 42, "Carton")],
                      po="POTWCTWA26100800C", pr="PRTWCOTWA202610080001", date="2026-10-22"), "供應商直送")
    imp(purchase_xlsx([("82_1", "4987176232878_CNN", "幫寶適 清新幫 拉拉褲", 216, "Pcs"), ("82_2", "4987176232878_RNN", "幫寶適 清新幫 拉拉褲 箱購", 10, "Carton")],
                      po="POTWCTWG26100800D", pr="PRTWCOTWG202610080002", wh="TWG", date="2026-10-23"), "供應商直送")
    imp(purchase_xlsx([("83_1", "1088888_PNN", "價格本沒有的紙潔商品", 24, "Pcs"), ("83_2", "1099999_PNN", "箱入數空白的紙潔商品", 24, "Pcs")],
                      po="POTWCTWA26100800E", pr="PRTWCOTWA202610080003", date="2026-10-22"), "供應商直送")
    c.post("/api/shopee/fulfil", json={"order_no": "POTWCTWA26100514N", "fulfil": "供應商直送", "reason": "改回原廠直送"})
    pl = plan("供應商直送")
    E, N = pl["POTWCTWA26100800E"], pl["POTWCTWA26100514N"]
    check("價格本沒有的、箱入數空白的 → 擋下來並說要更新哪個價格本", E["blocking"] == 2 and "請更新紙潔價格本" in item(E, "1088888_PNN")["problems"][0]
          and "沒有箱入數" in item(E, "1099999_PNN")["problems"][0], str([p["problems"] for p in E["items"]]))
    check("瑪氏商品總表還沒有這些料號 → 擋下來、說要到瑪氏出貨頁更新商品總表", N["supported"] and N["blocking"] == 2
          and "瑪氏商品總表找不到" in N["items"][0]["problems"][0], str(N["items"][0]["problems"]))
    r = c.post("/api/shopee/direct_po/draft", json={"order_nos": ["POTWCTWG26100800D", "POTWCTWA26100800C"]})
    zd = zipfile.ZipFile(io.BytesIO(r.data)) if r.status_code == 200 else None
    dnames = sorted(zd.namelist()) if zd else []
    check("寶僑直送不先產 EIP 就能先下載直送採購單：箱單位、小單位兩份、單號留空；勾到的紙潔不算（Jerry 2026-10-29）", r.status_code == 200
          and dnames == ["蝦皮特選直送採購單_1023_高鐵南_小單位.xlsx", "蝦皮特選直送採購單_1023_高鐵南_箱單位.xlsx"], str(dnames) if zd else r.get_json())
    if zd:
        wsd = openpyxl.load_workbook(io.BytesIO(zd.read("蝦皮特選直送採購單_1023_高鐵南_小單位.xlsx"))).active
        check("先下載的直送採購單：D1 空白、一品一列（採購單ID＝PR、料號、數量 216、箱入數 4）", wsd["D1"].value in ("", None)
              and [wsd.cell(4, k).value for k in (2, 3, 5, 7)] == ["PRTWCOTWG202610080002", "4987176232878_CNN", 216, 4], str([wsd.cell(4, k).value for k in range(1, 9)]))
    check("只勾紙潔 → 400（紙潔還是要先回填 EIP 單號）", c.post("/api/shopee/direct_po/draft", json={"order_nos": ["POTWCTWA26100800C"]}).status_code == 400)
    r = gen("供應商直送", ["POTWCTWA26100800C", "POTWCTWA26100514N"])
    check("勾到被擋的瑪氏單 → 400", r.status_code == 400)
    r = gen("供應商直送", ["POTWCTWA26100800C", "POTWCTWG26100800D"])
    z = zipfile.ZipFile(io.BytesIO(r.data)) if r.status_code == 200 else None
    names = sorted(z.namelist()) if z else []
    z = z if len(names) == 5 else None
    check("產出：紙潔一個檔、寶僑分箱單位／小單位，各自一天一倉；zip 名稱帶到貨日", names == sorted([
        "產品採購表上傳-紙潔直送_1022到觀音.xls", "產品採購表上傳-寶僑直送_1023到高鐵南-箱單位.xls", "產品採購表上傳-寶僑直送_1023到高鐵南-小單位.xls",
        "蝦皮特選-商品下採_紙潔直送_1022到觀音.xlsx", "蝦皮特選-商品下採_寶僑直送_1023到高鐵南.xlsx"])
          and "1022-1023" in r.headers.get("Content-Disposition", ""), str(names))
    if z:
        head, rows = xls_rows(z.read("產品採購表上傳-紙潔直送_1022到觀音.xls"))
        check("紙潔：單位、數量照蝦皮原本的（360 瓶、1392 包、42 箱），不換箱，備註空白，第 10 欄標題不動（Jerry 2026-10-29）",
              [(r_[1], r_[3], r_[6], r_[8]) for r_ in rows] == [("1022257", "瓶", 360, ""), ("1022248", "包", 1392, ""), ("1021315", "箱", 42, "")]
              and head[9] != "", str(rows))
        wbd = openpyxl.load_workbook(io.BytesIO(z.read("蝦皮特選-商品下採_紙潔直送_1022到觀音.xlsx")))
        check("紙潔下採明細只有一頁（不分箱、小單位；Jerry 2026-10-28）", wbd.sheetnames == ["下採"] and wbd["下採"].max_row == 4, str(wbd.sheetnames))
        head, rows = xls_rows(z.read("產品採購表上傳-寶僑直送_1023到高鐵南-小單位.xls"))
        check("寶僑直送小單位：單位寫箱、數量 216÷4＝54、轉換率 4", [(r_[1], r_[3], r_[6], r_[9]) for r_ in rows] == [("4987176232878", "箱", 54, 4)], str(rows))
        head, rows = xls_rows(z.read("產品採購表上傳-寶僑直送_1023到高鐵南-箱單位.xls"))
        check("寶僑直送箱單位：箱 10", [(r_[3], r_[6]) for r_ in rows] == [("箱", 10)], str(rows))
    check("瑪氏出貨頁的歷程不會混進蝦皮特選的紀錄", not any(l["field"].startswith("shopee_") for l in c.get("/api/master/logs?scope=mars&limit=500").get_json()["logs"]))

    print("\n【12】瑪氏直送：EIP 單位數量照蝦皮、依品類／單位／中標分檔；回填後產瑪氏採購單（需求 10.4、第十一章）")
    cn = db.get_conn()
    for row in (("M60003196", "盒", "60003196", "Petcare", "希寶 誘惑泥 55入", 794.29, 24, 24, "Q", 24, "V", "", ""),
                ("M10281171", "箱", "10281171", "Petcare", "偉嘉 貓罐頭 箱購", 500.0, 1, 24, "R", 1, "", "11/22 轉貨號", "需加工貼小白標")):
        cn.execute("INSERT INTO mst_mars_products (yf_sku, unit, mars_code, category, name, price, inner_per_case, pcs_per_case, unit_code, box_qty, "
                   "inner_label, note, po_case_note) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", row)
    cn.commit(); cn.close()
    check("價格本區塊多一個瑪氏（用瑪氏出貨的商品總表）", c.get("/api/shopee/conv").get_json()["lines"].get("瑪氏", {}).get("count") == 2)
    imp(purchase_xlsx([("84_1", "M60003196_QMN", "希寶 誘惑泥 55入", 50, "Pcs")], po="POTWCTWX26101100F", pr="PRTWCOTWX202610110001", wh="TWX", date="2026-10-14"),
        "供應商直送")
    imp(purchase_xlsx([("85_1", "M60003196_QMN", "希寶 誘惑泥 55入", 48, "Pcs")], po="POTWCTWA26101100G", pr="PRTWCOTWA202610110002", date="2026-10-14"),
        "竹運採購進貨")
    pl = plan("供應商直送")
    N, F = pl["POTWCTWA26100514N"], pl["POTWCTWX26101100F"]
    a, b_ = item(N, "M60003196_QMN"), item(N, "M10281171_RMN")
    check("瑪氏：單位數量照蝦皮（48 盒、18 箱），箱數 2／18，依品類、單位、中標分檔，備註 MARS入倉＋到貨日＋倉",
          not N["blocking"] and (a["out_unit"], a["out_qty"], a["cases"], a["unit_group"]) == ("盒", 48, 2, "盒_PET_需貼中標")
          and (b_["out_unit"], b_["out_qty"], b_["cases"], b_["unit_group"]) == ("箱", 18, 18, "箱_PET_不貼中標")
          and a["remark"] == "MARS入倉20261013台灣瑪氏送蝦皮特選-觀音", str(a)[:300])
    f1 = item(F, "M60003196_QMN")
    check("50 盒除以箱入數 24 除不盡 → 擋下來", F["blocking"] == 1 and f1["rem"] == 2 and "除不盡" in f1["problems"][0], str(f1["problems"]))
    r = c.post("/api/shopee/purchase/adjust", json={"order_id": f1["order_id"], "qty": 48, "reason": "瑪氏整箱出貨"})
    check("瑪氏也能改數量（轉換率用商品總表的箱入數）", r.status_code == 200 and not plan("供應商直送")["POTWCTWX26101100F"]["blocking"], str(r.get_json()))
    g_ = plan("竹運採購進貨")["POTWCTWA26101100G"]
    check("瑪氏的竹運採購進貨 → 擋下來（目前只開放供應商直送）", g_["blocking"] and "只開放供應商直送" in g_["items"][0]["problems"][0], str(g_["items"][0]["problems"]))
    r = gen("供應商直送", ["POTWCTWA26100514N", "POTWCTWX26101100F"])
    z = zipfile.ZipFile(io.BytesIO(r.data)) if r.status_code == 200 else None
    names = sorted(z.namelist()) if z else []
    check("產出：一份＝一倉一到貨日一品類一單位一中標，共 3 個 EIP 檔、沒有採購下採明細", names == sorted([
        "產品採購表上傳-瑪氏直送_1013到觀音-盒_PET_需貼中標.xls", "產品採購表上傳-瑪氏直送_1013到觀音-箱_PET_不貼中標.xls",
        "產品採購表上傳-瑪氏直送_1014到安南三-盒_PET_需貼中標.xls"]), str(names))
    if len(names) == 3:
        wb_ = xlrd.open_workbook(file_contents=z.read("產品採購表上傳-瑪氏直送_1013到觀音-盒_PET_需貼中標.xls"))
        _, rows = xls_rows(z.read("產品採購表上傳-瑪氏直送_1013到觀音-盒_PET_需貼中標.xls"))
        check("瑪氏 EIP：套瑪氏範本（產品領用單）、料號、單位盒、數量 48、備註", wb_.sheet_names()[0] == "產品領用單"
              and [(r_[1], r_[3], r_[6], r_[8]) for r_ in rows] == [("M60003196", "盒", 48, "MARS入倉20261013台灣瑪氏送蝦皮特選-觀音")], str(rows))
    eips = {e["filename"]: e for e in c.get("/api/shopee/eips").get_json()["eips"] if e["line"] == "瑪氏"}
    box_ = eips.get("產品採購表上傳-瑪氏直送_1013到觀音-盒_PET_需貼中標.xls", {})
    case_ = eips.get("產品採購表上傳-瑪氏直送_1013到觀音-箱_PET_不貼中標.xls", {})
    check("採購單產出清單標出瑪氏", box_.get("mars") and case_.get("mars"))
    r = c.get(f"/api/shopee/eips/{box_['id']}/mars_po")
    check("還沒回填 EIP 採購單號 → 不能產瑪氏採購單（400）", r.status_code == 400 and "回填" in r.get_json()["error"])
    check("沒回填的不能打包下載（400）", c.get("/api/shopee/mars_po.zip").status_code == 400)
    c.put(f"/api/shopee/eips/{box_['id']}", json={"eip_po": "PO202610101"})
    c.put(f"/api/shopee/eips/{case_['id']}", json={"eip_po": "PO202610102"})
    r = c.get(f"/api/shopee/eips/{box_['id']}/mars_po")
    import datetime as _dt
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active if r.status_code == 200 else None
    check("瑪氏採購單檔名：EIP 採購單號＋單位＋品類＋中標＋到貨日（不寫倉別；Jerry 2026-10-28）",
          __import__("urllib.parse").parse.unquote(r.headers.get("Content-Disposition", "")).endswith("PO202610101永豐Mars採購單_盒_PET_需貼中標_20261013.xlsx"),
          r.headers.get("Content-Disposition"))
    if ws:
        check("表頭：下單日 10/8（往前 2 個工作天，跳過 10/9 補假和週末）、配送日 10/13、EIP 採購單號、觀音倉資料、1/2效期以上、品類",
              ws["C3"].value == _dt.datetime(2026, 10, 8) and ws["F3"].value == _dt.datetime(2026, 10, 13) and ws["C5"].value == "PO202610101"
              and ws["C6"].value == "永豐商店蝦皮特選-觀音倉(TWA)" and ws["F6"].value == "1/2效期以上" and ws["F7"].value == 17617040
              and ws["F8"].value == "進貨組" and ws["B9"].value == "Petcare寵物", str([ws[k].value for k in ("C3", "F3", "C5", "C6", "F6", "F7", "F8", "B9")]))
        check("作業備註：蝦皮特選嘜頭＋採購單、需貼中盒標（盒單位不寫貼板嘜）", ws["F5"].value == "蝦皮特選嘜頭(入庫通知單)+採購單\n需貼中盒標", repr(ws["F5"].value))
        check("明細：永豐料號、瑪氏貨號、單位需求 48、採購數量 2 箱、中標 V、每箱數", [ws.cell(12, k).value for k in (1, 2, 4, 5, 8, 10, 11)]
              == ["M60003196", "60003196", 48, 2, "V", 24, 24], str([ws.cell(12, k).value for k in range(1, 14)]))
    ws = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/shopee/eips/{case_['id']}/mars_po").data)).active
    check("箱單位那份：作業備註加貼板嘜、指定品需加工貼小白標；備註、小白標欄照商品總表",
          ws["F5"].value == "蝦皮特選嘜頭(入庫通知單)+採購單\n貼板嘜\n指定品需加工貼小白標" and ws.cell(12, 12).value == "11/22 轉貨號"
          and ws.cell(12, 13).value == "需加工貼小白標" and ws.cell(12, 5).value == 18, repr(ws["F5"].value))
    r = c.get("/api/shopee/mars_po.zip")
    zn = sorted(zipfile.ZipFile(io.BytesIO(r.data)).namelist()) if r.status_code == 200 else []
    check("打包下載只放已回填的 2 份", zn == ["PO202610101永豐Mars採購單_盒_PET_需貼中標_20261013.xlsx", "PO202610102永豐Mars採購單_箱_PET_不貼中標_20261013.xlsx"], str(zn))
    whs = {w["code"]: w for w in c.get("/api/shopee/warehouses").get_json()["warehouses"]}
    check("蝦皮倉資料：觀音、安南三照需求書，其他倉 ship-to 空著", whs["TWX"]["ship_to"] == "17647732" and whs["TWX"]["contact"] == "曾小姐"
          and whs["TWG"]["address"].startswith("桃園市中壢區") and whs["TWG"]["ship_to"] == "", str(whs["TWG"]))
    r = c.put("/api/shopee/warehouses/TWG", json={"ship_to": "17600001", "contact": "高鐵南進貨組"})
    whs = {w["code"]: w for w in c.get("/api/shopee/warehouses").get_json()["warehouses"]}
    logs = [l for l in c.get("/api/master/logs?scope=shopee&limit=500").get_json()["logs"] if l["field"] == "shopee_wh"]
    check("改高鐵南的 ship-to：存好、只記有改的那一欄", r.get_json()["changed"] == 1 and whs["TWG"]["ship_to"] == "17600001" and whs["TWG"]["edited"]
          and len(logs) == 1 and logs[0]["new_value"] == "17600001", str(logs))
    check("沒有的倉 → 404", c.put("/api/shopee/warehouses/XXX", json={"ship_to": "1"}).status_code == 404)

    print("\n【12b】寶僑、紙潔供應商直送：回填 EIP 採購單號後產給採購轉原廠的直送採購單（Alice 2026-10-21 的範本）")
    unq = __import__("urllib.parse").parse.unquote
    eips = {e["filename"]: e for e in c.get("/api/shopee/eips").get_json()["eips"]}
    pg_s = eips.get("產品採購表上傳-寶僑直送_1023到高鐵南-小單位.xls", {}); pg_b = eips.get("產品採購表上傳-寶僑直送_1023到高鐵南-箱單位.xls", {})
    zj = eips.get("產品採購表上傳-紙潔直送_1022到觀音.xls", {})
    check("清單標出哪些能產直送採購單：寶僑、紙潔的供應商直送；瑪氏不算", pg_s.get("direct") and pg_b.get("direct") and zj.get("direct") and not box_.get("direct"),
          str([(k, v.get("direct")) for k, v in eips.items()]))
    r = c.get(f"/api/shopee/eips/{pg_s['id']}/direct_po")
    check("還沒回填 EIP 採購單號 → 400；瑪氏那張 → 404；都沒回填 → 打包下載 400", r.status_code == 400 and "回填" in r.get_json()["error"]
          and c.get(f"/api/shopee/eips/{box_['id']}/direct_po").status_code == 404 and c.get("/api/shopee/direct_po.zip").status_code == 400)
    c.put(f"/api/shopee/eips/{pg_s['id']}", json={"eip_po": "PO202610201"}); c.put(f"/api/shopee/eips/{zj['id']}", json={"eip_po": "PO202610202"})
    r = c.get(f"/api/shopee/eips/{pg_s['id']}/direct_po")
    fn = unq(r.headers.get("Content-Disposition", ""))
    check("直送採購單檔名：EIP 採購單號＋到貨日＋倉＋單位", r.status_code == 200 and "PO202610201蝦皮特選直送採購單_1023_高鐵南_小單位.xlsx" in fn, fn)
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active if r.status_code == 200 else None
    if ws:
        check("表頭：A1 蝦皮特選、C1 採購單號、D1 EIP 採購單號、G1 下單 H1 今天、G2 到庫 H2 到貨日",
              ws["A1"].value == "蝦皮特選" and ws["C1"].value == "採購單號" and ws["D1"].value == "PO202610201" and ws["G1"].value == "下單"
              and ws["H1"].value.date() == _dt.date.today() and ws["G2"].value == "到庫" and ws["H2"].value == _dt.datetime(2026, 10, 23),
              str([ws[k].value for k in ("A1", "C1", "D1", "G1", "H1", "G2", "H2")]))
        check("第 3 列表頭；第 4 列一品：採購單ID 沒入庫單號就填 PR 單號、供應商料號、SKU名稱、數量照蝦皮 216（不換箱）、箱入數照價格本 4、轉換箱數是公式 =E/G、而且塞了算好的 54（受保護的檢視才看得到）",
              [ws.cell(3, k).value for k in range(2, 9)] == ["採購單ID", "供應商料號", "SKU名稱", "數量", "單位", "箱入數", "轉換箱數"]
              and [ws.cell(4, k).value for k in (1, 2, 3, 4, 5, 7, 8)] == [1, "PRTWCOTWG202610080002", "4987176232878_CNN", "幫寶適 清新幫 拉拉褲", 216, 4, "=E4/G4"]
              and ws["F4"].value and openpyxl.load_workbook(io.BytesIO(r.data), data_only=True).active["H4"].value == 54, str([ws.cell(4, k).value for k in range(1, 10)]))
        check("數量黃底、表頭和轉換箱數灰底；留 15 列之後是 Total 列、SUM 公式也塞了值",
              ws["E4"].fill.fgColor.rgb.endswith("FFFF00") and ws["B3"].fill.fgColor.rgb.endswith("D9D9D9") and ws["H5"].fill.fgColor.rgb.endswith("D9D9D9")
              and ws["D19"].value == "Total" and ws["E19"].value == "=SUM(E4:E18)" and ws["H19"].value == "=SUM(H4:H18)"
              and [openpyxl.load_workbook(io.BytesIO(r.data), data_only=True).active[k].value for k in ("E19", "H19")] == [216, 54], str([ws["D19"].value, ws["E19"].value, ws["H19"].value]))
        check("下面：請安排 10/23 到貨、蝦皮特選店-高鐵南、高鐵南的進倉說明、固定的注意事項",
              ws["B21"].value == "請安排10/23 到貨" and ws["A22"].value == "蝦皮特選店-高鐵南" and "高鐵南路四段 888 號" in (ws["C22"].value or "")
              and (ws["C23"].value or "").startswith("●同批進貨商品效期須一致"), str([ws["B21"].value, ws["A22"].value, (ws["C22"].value or "")[:40]]))
    ws = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/shopee/eips/{zj['id']}/direct_po").data)).active
    got = sorted((ws.cell(r_, 2).value, ws.cell(r_, 3).value, ws.cell(r_, 5).value, ws.cell(r_, 7).value) for r_ in (4, 5, 6))
    check("紙潔那份：三品三列、採購單ID 都是 PR、數量照蝦皮 360／1392／42、箱入數 12；第 7 列起空白", got == [("PRTWCOTWA202610080001", "1021315_RNN", 42, 1),
          ("PRTWCOTWA202610080001", "1022248_CNN", 1392, 12), ("PRTWCOTWA202610080001", "1022257_PNN", 360, 12)] and ws["B7"].value in ("", None) and ws["D19"].value == "Total", str(got))
    inb = inbound_xlsx([("INTWG0002610230001", "82_1", "4987176232878_CNN", "幫寶適 清新幫 拉拉褲", 216)], "POTWCTWG26100800D", wh="TWG", date="2026-10-23")
    pv = preview(inb, "InboundOrder_20261021.xlsx").get_json(); c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"]})
    ws = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/shopee/eips/{pg_s['id']}/direct_po").data)).active
    check("蝦皮入庫單上傳後再下載 → 採購單ID 從 PR 單號換成入庫單號（用下載當下的）", ws["B4"].value == "INTWG0002610230001", str(ws["B4"].value))
    r = c.get("/api/shopee/direct_po.zip")
    zn = sorted(zipfile.ZipFile(io.BytesIO(r.data)).namelist()) if r.status_code == 200 else []
    check("打包下載只放已回填的寶僑、紙潔 2 份（瑪氏的不混進來）、zip 名稱帶到貨日區間；到貨日範圍可以篩", zn == ["PO202610201蝦皮特選直送採購單_1023_高鐵南_小單位.xlsx",
          "PO202610202蝦皮特選直送採購單_1022_觀音.xlsx"] and "1022-1023" in unq(r.headers.get("Content-Disposition", ""))
          and c.get("/api/shopee/direct_po.zip?from=2026-10-23").headers.get("X-Shopee-Files") == "1", str(zn))
    sd = c.get("/api/shopee/order", query_string={"no": "POTWCTWG26100800D", "from": "2026-01-01"}).get_json()
    check("右側卡片的 EIP 區標出可下載直送採購單（小單位回填了、箱單位還沒）", sorted((e["unit_group"], bool(e["direct"]), bool(e["eip_po"])) for e in sd["eips"])
          == [("小單位", True, True), ("箱單位", True, False)], str(sd["eips"]))
    print("\n【12c】取消產出 EIP：回到待產採購單（Jerry 2026-10-21「產出 EIP 之後反悔」）")
    r = c.delete(f"/api/shopee/eips/{pg_s['id']}")
    check("已回填單號的不能取消 → 400、說要先清掉單號", r.status_code == 400 and "先清掉單號" in r.get_json()["error"] and c.delete("/api/shopee/eips/999999").status_code == 404, str(r.get_json()))
    r = c.delete(f"/api/shopee/eips/{pg_b['id']}")
    ids_left = {e["id"] for e in c.get("/api/shopee/eips").get_json()["eips"]}
    check("箱單位那份（沒回填）取消產出 → 檔刪掉、回 1 張單", r.status_code == 200 and r.get_json()["orders"] == ["POTWCTWG26100800D"] and pg_b["id"] not in ids_left and pg_s["id"] in ids_left, str(r.get_json()))
    pd_ = plan("供應商直送")["POTWCTWG26100800D"]
    st_ = {x["supplier_sku_id"]: (x["eip"] or {}).get("eip_po") if x["eip"] else None for x in pd_["items"]}
    check("那張單回到待產採購單：箱單位品項沒有 EIP 了（可以重產）、小單位那份還在（已回填鎖住）、整張單狀態「已產出」",
          st_ == {"4987176232878_RNN": None, "4987176232878_CNN": "PO202610201"} and pd_["po_status"] == "已產出", str(st_) + pd_["po_status"])
    logs = [l for l in c.get("/api/master/logs?scope=shopee&limit=300").get_json()["logs"] if l["field"] == "shopee_eip" and l["new_value"] == "未產出"]
    check("歷程記「取消產出 EIP 採購單」", len(logs) == 1 and logs[0]["po_number"] == "POTWCTWG26100800D" and "箱單位" in logs[0]["old_value"], str(logs))
    r = c.put("/api/shopee/warehouses/TWG", json={"notes": "高鐵南改過的進倉說明"})
    whs = {w["code"]: w for w in c.get("/api/shopee/warehouses").get_json()["warehouses"]}
    ws = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/shopee/eips/{pg_s['id']}/direct_po").data)).active
    check("蝦皮倉資料 9 個倉（加仁英、安南二）、多「進倉說明」：改高鐵南的 → 之後下載的直送採購單就換", len(whs) == 9 and "TWR" in whs and "TWC" in whs
          and r.get_json()["changed"] == 1 and whs["TWG"]["notes"] == "高鐵南改過的進倉說明" and "寶倉街" in whs["TWA"]["notes"]
          and ws["C22"].value == "高鐵南改過的進倉說明", str(whs["TWG"]))
    ship_tests(c, preview, orders, up)
    short_tests(c, orders)
    yx_tests(c, orders)
    flow_tests(c, orders)
    trash_tests(c, orders)
    po_pdf_tests(c, orders)
    mixed_tests(c, preview, orders)


def mixed_tests(c, preview, orders):
    print("\n【17】一個採購單檔三個線別混在一起：履約方式每個線別各選一次（Jerry 2026-10-13）")
    raw = purchase_xlsx([("MX1", "4987176232878_CNN", "幫寶適 混線測試", 24, "Pcs"),
                         ("MX2", "1022257_PNN", "五月花 混線測試", 12, "Pcs"),
                         ("MX3", "M10273409_QMN", "M&M'S 混線測試", 6, "Pcs")], pr="PRTWCOTWA202612010001", date="2026-12-01")
    pv = preview(raw, "PurchaseOrder_mixed.xlsx").get_json()
    check("預覽列出新單各線別幾項（寶僑、瑪氏、紙潔各 1）", pv["new_lines"] == {"寶僑": 1, "瑪氏": 1, "紙潔": 1}, str(pv.get("new_lines")))
    r = c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"], "fulfils": {"寶僑": "竹運出貨", "瑪氏": "供應商直送"}})
    check("紙潔沒選 → 400，講是哪個線別沒選", r.status_code == 400 and "紙潔" in r.get_json()["error"] and "寶僑" not in r.get_json()["error"], r.get_json()["error"])
    r = c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"],
                                                   "fulfils": {"寶僑": "竹運出貨", "紙潔": "竹運採購進貨", "瑪氏": "供應商直送"}})
    check("三個線別都選了 → 匯入 3 項", r.status_code == 200 and r.get_json()["counts"]["new"] == 3, str(r.get_json()))
    g = next(x for x in orders(**{"from": "2026-12-01", "to": "2026-12-01"})["orders"] if x["order_no"] == "PRTWCOTWA202612010001")
    by = {i["line"]: i["fulfil"] for i in g["items"]}
    check("同一張 PR 裡各品照自己線別的履約方式", by == {"寶僑": "竹運出貨", "紙潔": "竹運採購進貨", "瑪氏": "供應商直送"}, str(by))
    up = next(u for u in c.get("/api/shopee/uploads").get_json()["uploads"] if u["filename"] == "PurchaseOrder_mixed.xlsx")
    check("上傳紀錄寫各線別選了什麼", up["fulfil"] == "寶僑竹運出貨、瑪氏供應商直送、紙潔竹運採購進貨", up["fulfil"])
    pv = preview(purchase_xlsx([("MX4", "4987176232878_RNN", "幫寶適 混線測試 箱", 2, "Carton")], pr="PRTWCOTWA202612010002", date="2026-12-01"),
                 "PurchaseOrder_one.xlsx").get_json()
    r = c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"], "fulfil": "供應商直送"})
    check("舊的寫法（一個檔選一種）照樣能用", r.status_code == 200 and r.get_json()["counts"]["new"] == 1, str(r.get_json()))
    cn = db.get_conn()
    cn.execute("INSERT INTO mst_mars_products (yf_sku, unit, category, name, box_qty, inner_label) VALUES (?,?,?,?,?,?)",
               ("M10273409", "箱", "Chocolate", "M&M'S 箱購", 1, ""))
    cn.commit(); cn.close()
    pl = c.get("/api/shopee/purchase/plan", query_string={"fulfil": "供應商直送", "from": "2026-12-01", "to": "2026-12-01"}).get_json()["orders"]
    mi = next(i for g in pl for i in g["items"] if i["supplier_sku_id"] == "M10273409_QMN")
    check("瑪氏商品總表有這個料號、但單位不同 → 講清楚有哪個單位、少了哪個（不是只說找不到）",
          mi["blocking"] and mi["problems"] == ["瑪氏商品總表有 M10273409 的箱，沒有盒，請按「補商品資料」填寫"], str(mi["problems"]))

    print("\n【18】補商品資料：瑪氏新品在面板直接補進商品總表（Jerry 2026-10-14：不用改 Excel 再到瑪氏出貨頁重傳）")
    oid = mi["order_id"]
    g = c.get("/api/shopee/mars_product", query_string={"order_id": oid}).get_json()
    check("視窗資料：料號、單位、蝦皮品名、同料號的箱那列（可帶入品類）、三個品類", g["code"] == "M10273409" and g["unit"] == "盒" and not g["exists"]
          and [x["unit"] for x in g["siblings"]] == ["箱"] and g["siblings"][0]["category"] == "Chocolate" and len(g["categories"]) == 3, str(g)[:300])
    pg_item = next(i for gg in orders(**{"from": "2026-12-01", "to": "2026-12-01"})["orders"] for i in gg["items"] if i["line"] == "寶僑")
    check("不是瑪氏的品項 → 400", c.post("/api/shopee/mars_product", json={"order_id": pg_item["id"], "category": "Chocolate", "box_qty": 12}).status_code == 400)
    add = lambda **kw: c.post("/api/shopee/mars_product", json={"order_id": oid, **kw})  # noqa: E731
    check("沒選品類 → 400", add(box_qty=12).status_code == 400)
    check("箱入數 0、小數、文字 → 400 講清楚", all(add(category="Chocolate", box_qty=v).status_code == 400 for v in (0, 1.5, "abc", ""))
          and "箱入數" in add(category="Chocolate", box_qty="").get_json()["error"])
    r = add(category="Chocolate", box_qty=12, inner_label="V", mars_code="60003112", price="30.5", shelf_days=364)
    check("填好存進商品總表", r.status_code == 200 and r.get_json()["unit"] == "盒", str(r.get_json()))
    pl = c.get("/api/shopee/purchase/plan", query_string={"fulfil": "供應商直送", "from": "2026-12-01", "to": "2026-12-01"}).get_json()["orders"]
    mi = next(i for gg in pl for i in gg["items"] if i["supplier_sku_id"] == "M10273409_QMN")
    check("補完那品就不擋了：箱入數 12、6 盒除不盡（餘 6）改成要修改數量，品類、中標用補的",
          mi["conv"] == 12 and mi["problems"] and "除不盡" in mi["problems"][0] and mi["unit_group"] == "盒_CHO_需貼中標", str(mi)[:300])
    check("再補一次 → 409（已經有了）", add(category="Chocolate", box_qty=12, inner_label="V").status_code == 409)
    logs = [l for l in c.get("/api/master/logs?scope=shopee&limit=50").get_json()["logs"] if l["field"] == "shopee_mars_product"]
    check("歷程記下誰補了什麼", len(logs) == 1 and "箱入數 12" in logs[0]["new_value"] and "需貼中標" in logs[0]["new_value"], str(logs)[:200])
    st = c.get("/api/mars/status").get_json()
    check("瑪氏出貨頁狀態列出蝦皮特選補登的品項", st["manual"] == ["M10273409（盒）"], str(st.get("manual")))
    master = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples", "mars", "fake", "假_瑪氏商品總表.xlsx")
    up_m = lambda data, name: c.post("/api/mars/products/import", data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data").get_json()  # noqa: E731
    rd = c.post("/api/shopee/drop", data={"file": (io.BytesIO(open(master, "rb").read()), "假_瑪氏商品總表.xlsx")}, content_type="multipart/form-data")
    d = rd.get_json()
    check("瑪氏商品總表丟蝦皮特選的上傳入口 → 認成瑪氏商品總表、走瑪氏出貨頁同一支匯入（Jerry 2026-10-14）",
          rd.status_code == 200 and rd.headers.get("X-Shopee-Kind") == "mars" and d["codes"] == 8, str((rd.status_code, rd.headers.get("X-Shopee-Kind"), str(d)[:120])))
    sr = c.get("/api/mars/products", query_string={"q": "M10273409"}).get_json()["rows"]
    check("重傳總表 Excel（裡面還沒有這品）→ 補登的那品保留、Excel 沒有的其他列照樣蓋掉",
          d["kept_manual"] == ["M10273409（盒）"] and [(x["unit"], x["manual_by"]) for x in sr] == [("盒", "Jerry")], str(d.get("kept_manual")) + str(sr))
    wb = openpyxl.load_workbook(master); ws = wb["商品表"]
    ws.append(["60003112", "Chocolate", "M&M'S", "M10273409", "M&M'S 精選片裝", 31, 1, 12, 12, "", "", "盒", "Q", 12, "V", 364, "", ""])
    buf = io.BytesIO(); wb.save(buf)
    d = up_m(buf.getvalue(), "補進去的.xlsx")
    sr = c.get("/api/mars/products", query_string={"q": "M10273409"}).get_json()["rows"]
    check("Excel 補進去之後再傳 → 以 Excel 為準（不再標補登）", d["kept_manual"] == [] and [(x["unit"], x["manual_by"], x["name"]) for x in sr] == [("盒", "", "M&M'S 精選片裝")]
          and c.get("/api/mars/status").get_json()["manual"] == [], str(sr))

    print("\n【19】清除資料、刪除歷程紀錄（只有管理員；Jerry 2026-10-14）")
    users = app_module.get_users(); users.setdefault("Jerry", "pw"); users.setdefault("Alice", "pw")
    app_module._write_users(users, {"Jerry"})
    def as_user(name):
        with c.session_transaction() as s_:
            s_["user"] = name
    cn = db.get_conn()
    cn.execute("INSERT INTO edit_logs (order_id, po_number, sku_id, field, field_label, old_value, new_value, operator, source, note, changed_at) "
               "VALUES (?,?,?,?,?,?,?,?,?,?,?)", (1, "PO_TEST_DEL", "SKU1", "qty", "數量", "1", "2", "Jerry", "manual", "", "2026-10-14T09:00:00"))
    cn.commit(); cn.close()
    as_user("Alice")
    check("一般同事：歷程看得到但不能刪、蝦皮特選看不到清除資料", c.get("/api/master/logs?scope=shopee").get_json()["can_delete"] is False
          and c.post("/api/master/logs/delete", json={"items": [{"src": "mst", "id": 1}]}).status_code == 403
          and c.post("/api/shopee/reset", json={"confirm": "清空資料"}).status_code == 403 and 'id="reset-box"' not in c.get("/shopee").get_data(as_text=True))
    as_user("Jerry")
    lg = c.get("/api/master/logs?scope=shopee&limit=500").get_json()
    check("管理員：可以刪、每筆帶來源；蝦皮特選的資料與設定有清除資料", lg["can_delete"] is True and all(l["src"] == "mst" for l in lg["logs"])
          and 'id="reset-box"' in c.get("/shopee").get_data(as_text=True))
    two = lg["logs"][:2]
    cp = next(l for l in c.get("/api/master/logs?scope=all&q=PO_TEST_DEL").get_json()["logs"] if l["po_number"] == "PO_TEST_DEL")
    r = c.post("/api/master/logs/delete", json={"items": [{"src": "mst", "id": x["id"]} for x in two] + [{"src": "coupang", "id": cp["id"]}]})
    left = {l["id"] for l in c.get("/api/master/logs?scope=shopee&limit=500").get_json()["logs"]}
    check("勾兩筆蝦皮特選的、一筆酷澎訂單管理系統的 → 刪掉 3 筆、只刪勾的那幾筆", r.get_json().get("deleted") == 3 and not ({x["id"] for x in two} & left)
          and len(left) == len(lg["logs"]) - 2 and not c.get("/api/master/logs?scope=all&q=PO_TEST_DEL").get_json()["logs"], str(r.get_json()))
    check("什麼都沒勾 → 400", c.post("/api/master/logs/delete", json={"items": []}).status_code == 400)
    n_conv = c.get("/api/shopee/conv").get_json()["lines"]
    n_mars = db.get_conn().execute("SELECT COUNT(*) AS n FROM mst_mars_products").fetchone()["n"]
    check("沒照著打「清空資料」→ 400；一個範圍都沒勾 → 400", c.post("/api/shopee/reset", json={"confirm": "清空"}).status_code == 400
          and c.post("/api/shopee/reset", json={"confirm": "清空資料", "orders": False, "refs": False}).status_code == 400)
    r = c.post("/api/shopee/reset", json={"confirm": "清空資料", "orders": True, "refs": False, "logs": False})
    logs = c.get("/api/master/logs?scope=shopee&limit=500").get_json()["logs"]
    check("只清訂單、保留歷程：訂單沒了、價格本還在、多一筆「清除資料」紀錄", r.status_code == 200 and "歷程紀錄保留" in r.get_json()["message"]
          and orders(**{"from": "", "to": ""})["all_orders"]["count"] == 0 and c.get("/api/shopee/conv").get_json()["lines"] == n_conv
          and any(l["field"] == "shopee_reset" for l in logs) and len(logs) > 1, r.get_data(as_text=True)[:200])
    r = c.post("/api/shopee/reset", json={"confirm": "清空資料", "orders": False, "refs": True, "logs": True})
    check("再清參考資料、歷程一起刪：價格本、庫存資料沒了，蝦皮特選歷程 0 筆，瑪氏商品總表沒被動到",
          r.status_code == 200 and not c.get("/api/shopee/conv").get_json()["lines"].get("寶僑")
          and not c.get("/api/shopee/stock").get_json()["stock"] and c.get("/api/master/logs?scope=shopee").get_json()["logs"] == []
          and db.get_conn().execute("SELECT COUNT(*) AS n FROM mst_mars_products").fetchone()["n"] == n_mars, r.get_data(as_text=True)[:200])


def ship_tests(c, preview, orders, up):
    import xlwt
    from shopee.ship import ratio_for, DEFAULT_SET
    print("\n【13】竹運出貨：照 Alice 的核對檔判斷效期與庫存，逐列扣庫存，產拋單檔＋下修刪單清單")
    page = c.get("/shopee").get_data(as_text=True)
    check("待產拋單檔這一步的清單、庫存狀態與效期標準在資料與設定、批次視窗", all(x in page for x in ('id="st-ship"', 'id="stock-st"', 'id="s-rule"', 'id="z-tbl"', 'id="dlg-ship"')))
    st = dict(DEFAULT_SET)
    check("效期比例照蝦皮進倉規範：90 天內 2/3、365 天內 2/5、其他 1/3；自訂就用自訂", ratio_for(90, st)[1] == "2/3" and ratio_for(300, st)[1] == "2/5"
          and ratio_for(1095, st)[1] == "1/3" and ratio_for(1095, dict(st, rule="custom", ratio=0.66))[0] == 0.66)
    def imp_in(rows, po, date="2026-10-20"):
        pv = preview(inbound_xlsx(rows, po, date=date), "InboundOrder_z.xlsx").get_json()
        return c.post("/api/shopee/import/commit", json={"upload_id": pv["upload_id"], "fulfil": "竹運出貨"}).get_json()
    imp_in([("INZA01", "z1", "4987176377708_CNN", "幫寶適 安睡褲", 81), ("INZA02", "z2", "6903148224090_VNN", "好自在 衛生棉", 144),
            ("INZA03", "z3", "4987176372093_TNN", "好自在 雲枕組", 12), ("INZA04", "z4", "4902430056519_PNN", "歐蕾 護膚霜", 102),
            ("INZA05", "z5", "4987176386700_CNN", "幫寶適 新品", 12), ("INZA06", "z6", "4987176000000_CNN", "主檔沒有的品", 10)], "POTWCTWA26101500H")
    imp_in([("INZB01", "z7", "4987176377708_CNN", "幫寶適 安睡褲", 80)], "POTWCTWA26101500J")
    q = {"from": "2026-10-20", "to": "2026-10-20"}
    def plan():
        return c.get("/api/shopee/ship/plan", query_string=q).get_json()
    def it(pl, no, sup):
        return next(x for g in pl["orders"] if g["order_no"] == no for x in g["items"] if x["supplier_sku_id"] == sup)
    pl = plan()
    check("還沒上傳庫存 → 每品都擋下來，說要先上傳 EC 即時庫存、竹運庫存效期表", all(x["blocking"] for g in pl["orders"] for x in g["items"])
          and "EC 即時庫存" in it(pl, "POTWCTWA26101500H", "4987176377708_CNN")["problems"][0], str(it(pl, "POTWCTWA26101500H", "4987176377708_CNN")["problems"]))
    wbx = xlwt.Workbook(); sh = wbx.add_sheet("全商品")
    for j, h in enumerate(["結餘日期", "料號", "單位", "品名", "EC可用量"]):
        sh.write(0, j, h)
    for i, row in enumerate([("2026/10/19", "4987176377708", "包", "安睡褲", 132), ("2026/10/19", "6903148224090", "袋", "衛生棉", 187),
                             ("2026/10/19", "4987176372093", "組", "雲枕組", 1), ("2026/10/19", "4902430056519", "瓶", "護膚霜", 100),
                             ("2026/10/19", "4987176000000", "包", "沒主檔", 50), ("2026/10/19", "4987176377708", "箱", "安睡褲箱", 9)], start=1):
        for j, v in enumerate(row):
            sh.write(i, j, v)
    bx = io.BytesIO(); wbx.save(bx)
    r = up("/api/shopee/stock/import", bx.getvalue(), "全商品_20261019.xls").get_json()
    check("上傳 EC 即時庫存（舊版 .xls）：認得、記結餘日期", r.get("kind") == "ec" and r["stock"]["ec"]["file_date"] == "2026/10/19" and r["stock"]["ec"]["today"], str(r))
    wbl = openpyxl.Workbook(); ws = wbl.active; ws.title = "sheet1"
    ws.append(["儲區", "客戶品號", "儲位", "可出貨量", "庫存量", "效期", "製造日期"])
    for row in (("良品區", "4987176377708C", "A1", 100, 100, "20290522", "20260523"), ("良品區", "4987176377708C", "A2", 38, 38, "20290522", ""),
                ("良品區", "4987176377708C", "A3", 41, 41, "20290502", ""), ("不良品區", "4987176377708C", "X1", 500, 500, "20300101", ""),
                ("良品區", "6903148224090V", "B1", 96, 96, "20290724", ""), ("良品區", "6903148224090V", "B2", 91, 91, "20290601", ""),
                ("良品區", "4987176372093T", "C1", 2, 2, "20271217", ""), ("良品區", "4902430056519P", "D1", 50, 50, "20270601", ""),
                ("良品區", "4902430056519P", "D2", 20, 20, "20290101", ""), ("良品區", "4987176000000C", "E1", 50, 50, "20290101", "")):
        ws.append(list(row))
    bl = io.BytesIO(); wbl.save(bl)
    r = up("/api/shopee/stock/import", bl.getvalue(), "效期表.xlsx").get_json()
    check("上傳竹運庫存效期表：認得、不良品區不算", r.get("kind") == "lots" and "非良品區 1 列不計" in r["notes"], str(r))
    def master(extra=()):
        wbm = openpyxl.Workbook(); wbm.active.title = "全商品"; wbm.active.append(["料號", "單位", "EC可用量"])
        wsm = wbm.create_sheet("商品主檔"); wsm.append(["客戶品號", "商品名稱", "國際條碼", "供應商代碼", "單位", "長", "寬", "高", "重", "箱入數", "商品年限(日)"])
        for code in ("4987176377708C", "6903148224090V", "4987176372093T", "4902430056519P") + tuple(extra):
            wsm.append([code, "", "", "", "", "", "", "", "", 1, 1095])
        wsm.append(["9999999999999C", "年限空白", "", "", "", "", "", "", "", 1, ""])
        bm = io.BytesIO(); wbm.save(bm)
        return bm.getvalue()
    r = up("/api/shopee/stock/import", master(), "核對竹運出貨效期及庫存V5.xlsm").get_json()
    check("上傳 Alice 的核對檔：只匯入商品主檔、年限空白的列出來", r.get("kind") == "life" and r["items"] == 4
          and any("只匯入商品主檔" in n for n in r["notes"]) and any("空白" in n for n in r["notes"]), str(r))
    wb_ = openpyxl.Workbook(); wb_.active.append(["隨便", "一個"]); b_ = io.BytesIO(); wb_.save(b_)
    r = up("/api/shopee/stock/import", b_.getvalue(), "x.xlsx")
    check("認不出的檔 → 400、說要哪些欄位", r.status_code == 400 and "EC可用量" in r.get_json()["error"])
    pl = plan()
    H, J = "POTWCTWA26101500H", "POTWCTWA26101500J"
    a = it(pl, H, "4987176377708_CNN")
    check("同效期不同儲位加總（138）；由早到晚扣：2029-05-02 的 41 先出、再從 2029-05-22 出 40（以前單一效期會挑量夠的 05-22 出 81）；合格基準＝到貨日＋1095×1/3",
          a["final"] == 81 and a["alloc"] == [{"exp": "2029-05-02", "qty": 41}, {"exp": "2029-05-22", "qty": 40}] and a["threshold"] == "2027-10-20" and a["ratio_text"] == "1/3"
          and a["book"] == 179 and not a["blocking"], str({k: a[k] for k in ("final", "alloc", "threshold", "book")}))
    b_it = it(pl, H, "6903148224090_VNN")
    check("一批不夠就混合效期由早到晚扣（2029-06-01 出 91＋2029-07-24 出 53）→ 144、不缺、沒有要人確認的提醒（OP 經 Jerry 2026-10-20）",
          b_it["final"] == 144 and not b_it["short"] and b_it["alloc"] == [{"exp": "2029-06-01", "qty": 91}, {"exp": "2029-07-24", "qty": 53}]
          and not b_it["warnings"] and [l["rank"] for l in sorted(b_it["lots"], key=lambda l: l["exp"])] == [1, 2], str((b_it["alloc"], b_it["warnings"])))
    c_it = it(pl, H, "4987176372093_TNN")
    check("EC 只剩 1 → 出 1、提醒少於 12 件蝦皮會拒收", c_it["final"] == 1 and any("少於 12 件" in w for w in c_it["warnings"]), str(c_it["warnings"]))
    d = it(pl, H, "4902430056519_PNN")
    check("有一批效期不合格（2027-06-01）：只出合格的 20，缺貨原因效期不足＋庫存不足", d["final"] == 20 and d["qualified"] == 20 and d["book"] == 70
          and d["reasons"] == ["效期不足", "庫存不足"] and [l["ok"] for l in d["lots"]] == [False, True], str(d["reasons"]))
    e_ = it(pl, H, "4987176386700_CNN")
    check("完全沒貨、主檔也沒有 → 不擋，算庫存不足、出 0", e_["final"] == 0 and e_["reasons"] == ["庫存不足"] and not e_["blocking"], str(e_["problems"]))
    f_ = it(pl, H, "4987176000000_CNN")
    check("有貨但商品主檔沒有年限 → 擋下來要補主檔（不當成 0 天）", f_["blocking"] and "商品主檔" in f_["problems"][0], str(f_["problems"]))
    j = it(pl, J, "4987176377708_CNN")
    check("逐列扣庫存：後面那張單只剩 EC 51、2029-05-02 用完、2029-05-22 剩 98 → 出 51、缺 29", j["ec"] == 51 and j["final"] == 51 and j["short"] == 29
          and j["alloc"] == [{"exp": "2029-05-22", "qty": 51}], str({k: j[k] for k in ("ec", "final", "alloc")}))
    check("畫面要講得出來：EC 檔案 132、到貨日較早的 1 張單先用 81；竹運檔案 179、先用 81（Jerry 2026-10-20）；第一張單沒人先用",
          (j["ec_file"], j["ec_used"], j["ec_users"], j["book_file"], j["book_used"], j["book_users"]) == (132, 81, 1, 179, 81, 1) and a["ec_used"] == 0 and a["book_used"] == 0,
          str({k: j[k] for k in ("ec_file", "ec_used", "ec_users", "book_file", "book_used", "book_users")}))
    r = c.post("/api/shopee/ship/generate", json={"order_nos": [H, J], **q})
    check("有要處理的（主檔沒年限）→ 整批不產、列出原因", r.status_code == 400 and any("4987176000000_CNN" in x for x in r.get_json()["details"]), str(r.get_json()))
    up("/api/shopee/stock/import", master(("4987176000000C",)), "產品基本資料.xlsx")
    adj = lambda body: c.post("/api/shopee/ship/adjust", json={**q, **body})  # noqa: E731
    r = adj({"order_id": b_it["order_id"], "lots": [{"exp": "2029-07-24", "qty": 96}, {"exp": "2029-06-01", "qty": 48}], "reason": ""})
    check("改出貨量沒寫原因也能存（原因選填，之後在待產拋單檔一次填；Alice 2026-10-21）", r.status_code == 200 and r.get_json()["qty"] == 144, str(r.get_json()))
    r = adj({"order_id": b_it["order_id"], "lots": [{"exp": "2029-06-01", "qty": 92}], "reason": "x"})
    check("選的數量超過那批剩的 → 400", r.status_code == 400 and "只剩 91" in r.get_json()["error"], r.get_json()["error"])
    r = adj({"order_id": d["order_id"], "lots": [{"exp": "2027-06-01", "qty": 10}], "reason": "x"})
    check("選效期不合格的批次 → 400", r.status_code == 400 and "不符合" in r.get_json()["error"])
    r = adj({"order_id": b_it["order_id"], "lots": [{"exp": "2029-07-24", "qty": 96}, {"exp": "2029-06-01", "qty": 48}], "reason": "業務確認兩個效期分開入庫"})
    r2 = adj({"order_id": c_it["order_id"], "lots": [], "reason": "低於 12 件會被拒收，刪單"})
    pl = plan()
    b_it, c_it = it(pl, H, "6903148224090_VNN"), it(pl, H, "4987176372093_TNN")
    logs = [l for l in c.get("/api/master/logs?scope=shopee&limit=500").get_json()["logs"] if l["field"] == "shopee_ship_qty"]
    r_side = c.get("/api/shopee/order", query_string={**q, "no": H, "ship": "0"})
    check("右側面板帶 ship=0（畫面已經有同一組篩選的拋單檔計畫）→ 不再算一次、ship 回空（搜尋慢，2026-10-20）", r_side.status_code == 200 and r_side.get_json()["ship"] is None
          and c.get("/api/shopee/order", query_string={**q, "no": H}).get_json()["ship"]["order_no"] == H)
    check("人工改成 96＋48 → 出 144、不缺；記人員、原因；不再提醒入庫單要拆", r.status_code == 200 and b_it["final"] == 144 and not b_it["short"]
          and b_it["adj"]["reason"] == "業務確認兩個效期分開入庫" and b_it["adj"]["lots"] == [["2029-07-24", 96], ["2029-06-01", 48]] and not b_it["warnings"] and logs,
          str((b_it["adj"], b_it["warnings"])))
    check("改成 0 → 刪單，缺 12", r2.status_code == 200 and c_it["final"] == 0 and c_it["short"] == 12)
    # 一次填原因、少於 12 件一鍵改 0（Alice 2026-10-21「修改數量，使用批次確認原因」；Jerry：不擋產出）
    r = adj({"order_id": b_it["order_id"], "lots": [{"exp": "2029-07-24", "qty": 96}, {"exp": "2029-06-01", "qty": 48}], "reason": ""})
    pl = plan(); b_it = it(pl, H, "6903148224090_VNN")
    check("效期視窗不帶原因再存 → 沿用這品原本的原因（原因統一在右側卡片填，Jerry 2026-10-21）", r.status_code == 200 and b_it["adj"]["reason"] == "業務確認兩個效期分開入庫" and b_it["final"] == 144, str(b_it["adj"]))
    check("一次填原因：沒勾 → 400、沒原因 → 400", c.post("/api/shopee/ship/reasons", json={"order_ids": [], "reason": "x"}).status_code == 400
          and c.post("/api/shopee/ship/reasons", json={"order_ids": [b_it["order_id"]], "reason": ""}).status_code == 400)
    r = c.post("/api/shopee/ship/reasons", json={"order_ids": [b_it["order_id"], c_it["order_id"]], "reason": "竹運庫存不足"})
    pl = plan(); b_it, c_it = it(pl, H, "6903148224090_VNN"), it(pl, H, "4987176372093_TNN")
    logs_r = [l for l in c.get("/api/master/logs?scope=shopee&limit=500").get_json()["logs"] if l["field"] == "shopee_ship_reason"]
    check("兩項一起填「竹運庫存不足」→ 兩項的調整原因都換了、歷程記舊→新", r.status_code == 200 and r.get_json()["count"] == 2
          and b_it["adj"]["reason"] == "竹運庫存不足" and c_it["adj"]["reason"] == "竹運庫存不足" and len(logs_r) >= 2 and logs_r[0]["new_value"] == "竹運庫存不足", str(r.get_json()) + str(logs_r[:1]))
    r = c.post("/api/shopee/ship/zero", json={"order_ids": [b_it["order_id"]], "reason": "低於 12 件拒收，刪單", **q})
    pl = plan(); b_it = it(pl, H, "6903148224090_VNN")
    check("右側卡片「改成 0（刪單）」（Jerry 2026-10-21 做在卡片）→ 出 0、缺全部、原因記下；沒勾 → 400", r.status_code == 200 and r.get_json()["count"] == 1 and b_it["final"] == 0
          and b_it["short"] == b_it["qty"] and b_it["adj"]["reason"] == "低於 12 件拒收，刪單" and c.post("/api/shopee/ship/zero", json={"order_ids": [], **q}).status_code == 400, str(b_it["adj"]))
    r = adj({"order_id": b_it["order_id"], "lots": [{"exp": "2029-07-24", "qty": 96}, {"exp": "2029-06-01", "qty": 48}], "reason": "業務確認兩個效期分開入庫"})
    pl = plan(); b_it, c_it = it(pl, H, "6903148224090_VNN"), it(pl, H, "4987176372093_TNN")
    c.post("/api/shopee/ship/reasons", json={"order_ids": [c_it["order_id"]], "reason": "低於 12 件會被拒收，刪單"})
    pl = plan(); b_it, c_it = it(pl, H, "6903148224090_VNN"), it(pl, H, "4987176372093_TNN")
    check("改回 96＋48、刪單那項原因改回去（後面的測試照舊）", r.status_code == 200 and b_it["final"] == 144 and c_it["adj"]["reason"] == "低於 12 件會被拒收，刪單")
    r = c.put("/api/shopee/ship/settings", json={"rule": "custom", "ratio": "abc"})
    check("效期比例亂填 → 400", r.status_code == 400)
    c.put("/api/shopee/ship/settings", json={"rule": "custom", "ratio": "66"})
    d2 = it(plan(), H, "4902430056519_PNN")
    check("改成自訂比例 66（當 0.66）→ 基準日＝到貨日＋722 天", d2["ratio_text"] == "自訂 0.66" and d2["threshold"] == "2028-10-11", str((d2["ratio_text"], d2["threshold"])))
    c.put("/api/shopee/ship/settings", json={"rule": "shopee"})
    r = c.post("/api/shopee/ship/generate", json={"order_nos": [H, J], **q})
    z = zipfile.ZipFile(io.BytesIO(r.data)) if r.status_code == 200 else None
    names = sorted(z.namelist()) if z else []
    check("產出 zip：拋單檔（一天一倉一線別）＋下修刪單清單", names == ["指定到貨日20261020蝦皮(觀音)_PG.xlsx", "蝦皮特選下修刪單清單_1020.xlsx"], str(names))
    if len(names) == 2:
        rows = list(openpyxl.load_workbook(io.BytesIO(z.read(names[0]))).active.iter_rows(values_only=True))
        check("拋單檔照範例：表頭、訂單單號＝入庫單號、蝦皮-觀音、範例地址電話、料號＝國條、出貨量 0 的不放", rows[0] == ("訂單單號", "姓名", "地址", "電話", "手機", "出貨備註", "料號", "產品名稱", "單位", "數量")
              and rows[1][:4] == ("INZA01", "蝦皮-觀音", "桃園市觀音區寶倉街108號C棟1樓", "(03)262-6560") and rows[1][6:] == ("4987176377708", "幫寶適 安睡褲", "包", 81)
              and sorted((r_[0], r_[9]) for r_ in rows[1:]) == [("INZA01", 81), ("INZA02", 144), ("INZA04", 20), ("INZA06", 10), ("INZB01", 51)], str(rows[1:]))
        srows = list(openpyxl.load_workbook(io.BytesIO(z.read(names[1]))).active.iter_rows(values_only=True))
        acts = {(r_[5], r_[0], r_[10], r_[11], r_[12]) for r_ in srows[1:]}
        check("下修刪單清單：刪單／下修、下修數量、缺貨原因、調整原因", ("4987176372093_TNN", "刪單", 12, "庫存不足", "低於 12 件會被拒收，刪單") in acts
              and ("4902430056519_PNN", "下修", 82, "效期不足、庫存不足", None) in acts and ("4987176386700_CNN", "刪單", 12, "庫存不足", None) in acts
              and ("4987176377708_CNN", "下修", 29, "庫存不足", None) in acts and len(srows) == 5, str(acts))
    pl = plan()
    check("竹運採購進貨、入庫單到了的單（L）也在拋單檔清單裡，履約方式還是竹運採購進貨（Becky 2026-10-22）", "POTWCTWA26100514L" in {g["order_no"] for g in pl["orders"]}
          and next(g for g in pl["orders"] if g["order_no"] == "POTWCTWA26100514L")["dispatch_status"] == "未產出", str([g["order_no"] for g in pl["orders"]]))
    check("產出後兩張都「已產出」、品項照快照（不再分配庫存）", all(g["dispatch_status"] == "已產出" for g in pl["orders"] if g["order_no"] in (H, J))
          and it(pl, H, "6903148224090_VNN")["final"] == 144 and it(pl, H, "6903148224090_VNN")["dispatch"])
    ov = next(g for g in orders()["orders"] if g["order_no"] == H)
    check("訂單總覽的出貨量、缺貨量帶出來", ov["qty_ship"] == 255 and ov["qty_short"] == 106, str((ov["qty_ship"], ov["qty_short"])))
    check("再產一次同樣的單 → 400（要先取消產出）", c.post("/api/shopee/ship/generate", json={"order_nos": [H, J], **q}).status_code == 400)
    pv = preview(inbound_xlsx([("INZA01", "z1", "4987176377708_CNN", "幫寶適 安睡褲", 81)], H, date="2026-10-20"), "InboundOrder_again.xlsx").get_json()
    check("重傳已產過拋單檔那張的入庫單 → 預覽當場講「已產過拋單檔」、幾點產的（Jerry 2026-10-28）", [x["order_no"] for x in pv["dispatched"]] == [H]
          and pv["dispatched"][0]["at"][:10] == db.now()[:10], str(pv.get("dispatched")))
    c.post("/api/shopee/import/cancel", json={"upload_id": pv["upload_id"]})
    check("待產拋單檔有「顯示已產出」", 'id="z-done"' in c.get("/shopee").get_data(as_text=True))
    r = c.post("/api/shopee/ship/adjust", json={**q, "order_id": b_it["order_id"], "lots": [{"exp": "2029-07-24", "qty": 96}], "reason": "x"})
    check("已產出的不能改出貨量 → 400", r.status_code == 400 and "取消產出" in r.get_json()["error"])
    r = c.post("/api/shopee/ship/undo", json={"order_no": J})
    pl = plan()
    check("取消產出：那張回到未產出、重新照庫存判斷（前面那張已產出不再扣 → 這次可出 80）", r.get_json()["removed"] == 1
          and next(g for g in pl["orders"] if g["order_no"] == J)["dispatch_status"] == "未產出" and it(pl, J, "4987176377708_CNN")["final"] == 80,
          str(it(pl, J, "4987176377708_CNN")["final"]))
    logs = c.get("/api/master/logs?scope=shopee&limit=500").get_json()["logs"]
    check("歷程記下上傳庫存、產出拋單檔、取消產出", any(l["field"] == "shopee_stock" for l in logs)
          and any(l["field"] == "shopee_ship" and "出貨" in (l["new_value"] or "") for l in logs) and any(l["field"] == "shopee_ship" and l["new_value"] == "未產出" for l in logs))
    w = {x["code"]: x for x in c.get("/api/shopee/dispatch_wh").get_json()["warehouses"]}
    r = c.put("/api/shopee/dispatch_wh/TWA", json={"phone": "03-2600500#5924"})
    check("拋單檔收件資料：7 個倉、觀音照範例，可以改", len(w) == 7 and w["TWA"]["phone"] == "(03)262-6560" and r.get_json()["changed"] == 1
          and {x["code"]: x for x in c.get("/api/shopee/dispatch_wh").get_json()["warehouses"]}["TWA"]["phone"] == "03-2600500#5924")


def short_tests(c, orders):
    print("\n【14】缺貨維護與通知：寶僑兩次、紙潔一次、竹運出貨自動帶入，產 LINE 文字＋Excel，通知過鎖住")
    page = c.get("/shopee").get_data(as_text=True)
    check("待通知蝦皮這一步的清單、填缺貨視窗、通知視窗", all(x in page for x in ('id="st-notify"', 'id="k-tbl"', 'id="dlg-short"', 'id="dlg-notify"')))
    def plan():
        return {g["order_no"]: g for g in c.get("/api/shopee/short/plan", query_string={"from": "2026-01-01"}).get_json()["orders"]}
    def it(pl, no, sup):
        return next(x for x in pl[no]["items"] if x["supplier_sku_id"] == sup)
    save = lambda body: c.post("/api/shopee/short/save", json=body)  # noqa: E731
    L, H, C, D = "POTWCTWA26100514L", "POTWCTWA26101500H", "POTWCTWA26100800C", "POTWCTWG26100800D"
    pl = plan()
    check("清單有寶僑、紙潔、瑪氏；瑪氏只列供應商直送（竹運採購進貨那張 G 不列）",
          {g["line"] for g in pl.values()} == {"寶僑", "紙潔", "瑪氏"} and "POTWCTWA26100514N" in pl and "POTWCTWA26101100G" not in pl, str(sorted(pl)))
    nn = it(pl, "POTWCTWA26100514N", "M60003196_QMN")
    check("瑪氏不能人工填缺貨（由勇信核對），最多一次", nn["edit_round"] is None and nn["max_rounds"] == 1
          and save({"order_id": nn["order_id"], "short": 1, "reasons": ["庫存不足"]}).status_code == 400)
    h = it(pl, H, "4902430056519_PNN")
    check("竹運出貨產過拋單檔的缺貨自動帶進來、待通知（不用再填、這裡不能改）", pl[H]["status"] == "待通知" and h["rounds"][0]["source"] == "竹運出貨"
          and h["final"] == 20 and h["cut"] == 82 and h["rounds"][0]["reasons"] == ["效期不足", "庫存不足"] and h["edit_round"] is None, str(h["rounds"]))
    a = it(pl, L, "4987176232878_CNN")
    check("寶僑竹運採購進貨：還沒缺貨 → 原始訂購量就是目前蝦皮數量 190、可以填第 1 次", a["orig"] == 190 and a["final"] == 190 and a["edit_round"] == 1
          and a["max_rounds"] == 2 and pl[L]["status"] == "未通知", str(a)[:200])
    oid = a["order_id"]
    check("缺貨數量超過可出貨量 → 400", save({"order_id": oid, "short": 999, "reasons": ["庫存不足"]}).status_code == 400)
    check("沒勾原因 → 400", save({"order_id": oid, "short": 40, "reasons": []}).status_code == 400)
    check("竹運出貨的品項不能在這裡填 → 400", save({"order_id": h["order_id"], "short": 1, "reasons": ["庫存不足"]}).status_code == 400)
    r = save({"order_id": oid, "short": 40, "reasons": ["庫存不足"], "note": "採購 10/12 來信缺貨"})
    r2 = save({"order_id": oid, "short": 50, "reasons": ["庫存不足", "效期不足"], "note": "採購 10/12 來信缺貨（更正）"})
    a = it(plan(), L, "4987176232878_CNN")
    check("填第 1 次缺 40，還沒通知前改成缺 50 → 還是第 1 次、190 → 140、兩個原因", r.status_code == 200 and r2.get_json()["round"] == 1 and len(a["rounds"]) == 1
          and (a["rounds"][0]["before"], a["rounds"][0]["after"]) == (190, 140) and a["rounds"][0]["reasons"] == ["庫存不足", "效期不足"] and a["pending"], str(a["rounds"]))
    ov = next(g for g in orders()["orders"] if g["order_no"] == L)
    check("訂單總覽：通知狀態「待通知」、可出貨量、缺貨量帶出來", ov["notify_status"] == "待通知" and ov["qty_short"] == 50, str((ov["notify_status"], ov["qty_ship"], ov["qty_short"])))
    r = c.post("/api/shopee/short/notify", json={"order_nos": [L, H]})
    n = c.get(f"/api/shopee/short/notify/{r.get_json()['id']}").get_json() if r.status_code == 200 else {}
    ms = n.get("messages", [])
    mL = next((m for m in ms if m.startswith("【蝦皮缺貨下修通知")), "")
    check("合成兩則（Penny 2026-10-22「刪單的合成一則，缺貨的合成一則」）：同一個賣場的刪單一則、下修一則，刪單在前；L、H 的下修都在下修那則",
          r.status_code == 200 and [m.split("\n")[0] for m in ms] == ["【蝦皮缺貨刪單通知｜首次通知】", "【蝦皮缺貨下修通知｜首次通知】"]
          and "PR ID：PRTWCOTWA202610010001" in mL and "POTWCTWA26101500H" in mL, str([m.split("\n")[:3] for m in ms]))
    check("下修那則：開頭賣場 P&G、共幾張 PR 幾項；裡面照 PR 分段（■ PR ID、預計到貨日、到貨倉別、下修SKU）、PO ID、蝦皮SKU ID、料號品名、訂購／可出貨／缺貨、處理方式",
          mL.split("\n")[:3] == ["【蝦皮缺貨下修通知｜首次通知】", "賣場：P&G", f"共 2 張 PR、{mL.count('處理方式：')} 項"]
          and "■ PR ID：PRTWCOTWA202610010001\n預計到貨日：2026-10-20\n到貨倉別：TWA\n下修SKU（1項）：" in mL
          and "1. PO ID：POTWCTWA26100514L" in mL and "蝦皮SKU ID：41511463369_285962283475" in mL and "4987176232878_CNN 幫寶適 清新幫 拉拉褲(XXL)26片" in mL
          and "訂購190／可出貨140／缺貨50" in mL and "處理方式：下修至140" in mL and mL.endswith("請協助下修訂單數量，完成後請回覆，謝謝。"), mL)
    mH = next((m for m in ms if m.startswith("【蝦皮缺貨刪單通知")), "")
    check("刪單那則：H 那段刪單SKU 2 項、PO ID、不寫數量、結尾請協助處理刪單", "刪單SKU（2項）：" in mH and "4987176372093_TNN" in mH and "4987176386700_CNN" in mH
          and "訂購" not in mH and mH.endswith("請協助處理刪單，完成後請回覆，謝謝。"), mH)
    check("取得通知時帶刪單幾項、下修幾項（Email 主旨用）", n.get("del_items") == 2 and n.get("cut_items", 0) >= 1, str((n.get("del_items"), n.get("cut_items"))))
    import shopee.short as _sh
    big = [{"line": "寶僑", "order_no": f"PO{i}", "expected_date": "2026-10-20", "warehouse": "TWA",
            "items": [{"pending": True, "final": 1, "pr_id": f"PRBIG{i:03d}", "po_id": f"PO{i}", "shopee_sku_id": "1" * 20, "supplier_sku_id": "4987176232878_CNN",
                       "name": "很長的品名" * 8, "orig": 10, "cut": 9, "rounds": [{"round": 1}]} for _ in range(4)]} for i in range(30)]
    bm = _sh.line_messages(big)
    check("超過 LINE 上限（5000 字）→ 照 PR 切成（1/2）（2/2）…，每則都不超過、每則都有開頭和結尾", len(bm) > 1 and all(len(m) <= 5000 for m in bm)
          and bm[0].split("\n")[0].endswith(f"（1/{len(bm)}）") and all(m.endswith("完成後請回覆，謝謝。") for m in bm)
          and sum(m.count("■ PR ID：") for m in bm) == 30, str([len(m) for m in bm]))
    r_m = c.put("/api/shopee/notify_mail", json={"to": "a@shopee.tw；b@shopee.tw"})
    check("Email 收件人：分號、逗號都可以，存成逗號分開；格式不對 → 400", r_m.status_code == 200 and c.get("/api/shopee/notify_mail").get_json()["to"] == "a@shopee.tw, b@shopee.tw"
          and c.put("/api/shopee/notify_mail", json={"to": "不是信箱"}).status_code == 400, str(r_m.get_json()))
    pl = plan()
    a, h = it(pl, L, "4987176232878_CNN"), it(pl, H, "4902430056519_PNN")
    check("通知後：鎖住、狀態「已通知（第 1 次）」、竹運出貨那幾筆也存下來了", not a["pending"] and a["notified_round"] == 1 and a["rounds"][0]["notified"]
          and pl[L]["status"] == "已通知（第 1 次）" and h["rounds"][0]["saved"] and not h["pending"], str(pl[L]["status"]))
    check("沒有待通知的單 → 400", c.post("/api/shopee/short/notify", json={"order_nos": [L]}).status_code == 400)
    check("寶僑通知過第 1 次 → 可以填第 2 次（拋單時又缺）", a["edit_round"] == 2 and a["prev_final"] == 140)
    r = save({"order_id": oid, "short": 20, "reasons": ["庫存不足"], "note": "拋單時又缺"})
    a = it(plan(), L, "4987176232878_CNN")
    check("第 2 次：140 → 120，第 1 次的紀錄還在、累計下修 70", r.get_json()["round"] == 2 and [(x["round"], x["before"], x["after"]) for x in a["rounds"]] == [(1, 190, 140), (2, 140, 120)]
          and a["cut"] == 70 and a["rounds"][0]["notified"], str(a["rounds"]))
    rn = it(plan(), L, "4987176232878_RNN")
    save({"order_id": rn["order_id"], "short": 2, "reasons": ["庫存不足"], "note": "同一張單另一個 SKU 第一次缺"})
    r = c.post("/api/shopee/short/notify", json={"order_nos": [L]})
    ms2 = c.get(f"/api/shopee/short/notify/{r.get_json()['id']}").get_json()["messages"]; t = ms2[0]
    check("同一張 PR 一個 SKU 第 2 次、另一個第 1 次 → 下修合成一則（Jerry 2026-10-14），標題用第二次通知、每品標第幾次、缺貨寫累計、下修至 120",
          len(ms2) == 1 and t.startswith("【蝦皮缺貨下修通知｜第二次通知】") and "下修SKU（2項）：" in t and "訂購190／可出貨120／缺貨70" in t
          and "處理方式：下修至120" in t and "拉拉褲(XXL)26片（第2次）" in t and "4987176232878_RNN" in t and "箱購(XXL)26片X4包（第1次）" in t, t)
    r = save({"order_id": oid, "short": 1, "reasons": ["庫存不足"]})
    check("寶僑兩次都通知過 → 不能再填（400）", r.status_code == 400 and "2 次" in r.get_json()["error"], r.get_json()["error"])
    cc = it(plan(), C, "1022257_PNN")
    save({"order_id": cc["order_id"], "short": cc["final"], "reasons": ["效期不足"], "note": "Teams 通知"})
    r = c.post("/api/shopee/short/notify", json={"order_nos": [C]})
    nid = r.get_json()["id"]
    x = c.get(f"/api/shopee/short/notify/{nid}/xlsx")
    wb_ = openpyxl.load_workbook(io.BytesIO(x.data)) if x.status_code == 200 else None
    check("明細 Excel：下修明細、刪單明細、最終可出貨量清單三頁；紙潔整筆缺在刪單明細", wb_ and wb_.sheetnames == ["下修明細", "刪單明細", "最終可出貨量清單"]
          and [r_[6] for r_ in wb_["刪單明細"].iter_rows(min_row=2, values_only=True)] == ["1022257_PNN"]
          and list(wb_["刪單明細"].iter_rows(values_only=True))[0][10:15] == ("原始訂購量", "最終可出貨量", "下修數量", "缺貨原因", "下修次數"))
    check("紙潔通知過一次 → 不能再填（400）", save({"order_id": cc["order_id"], "short": 1, "reasons": ["庫存不足"]}).status_code == 400)
    dd = it(plan(), D, "4987176232878_CNN")
    save({"order_id": dd["order_id"], "short": 10, "reasons": ["庫存不足"]})
    r = save({"order_id": dd["order_id"], "short": 0})
    check("還沒通知的那次填 0 → 取消", r.get_json().get("cleared") and not it(plan(), D, "4987176232878_CNN")["rounds"])
    ov = next(g for g in orders()["orders"] if g["order_no"] == L)
    check("訂單總覽：已通知（第 2 次）", ov["notify_status"] == "已通知（第 2 次）", ov["notify_status"])
    notes = c.get("/api/shopee/short/plan").get_json()["notifies"]
    logs = c.get("/api/master/logs?scope=shopee&limit=800").get_json()["logs"]
    check("通知紀錄 3 筆、歷程記下填缺貨和通知", len(notes) == 3 and any(l["field"] == "shopee_short" for l in logs) and any(l["field"] == "shopee_notify" for l in logs))



def yx_tests(c, orders):
    print("\n【15】瑪氏缺貨：勇信配送明細表對回瑪氏 EIP，逐品比箱數，寫成缺貨紀錄，再產 LINE 通知（需求 8.1）")
    N = "POTWCTWA26100514N"
    page = c.get("/shopee").get_data(as_text=True)
    check("勇信配送明細表從同一個上傳入口進來（收 PDF），比對結果在視窗裡看", 'accept=".xlsx,.xls,.xlsm,.pdf"' in page and 'id="dlg-yx"' in page and 'id="yx-out"' in page)
    eips = {e["eip_po"]: e for e in c.get("/api/shopee/eips").get_json()["eips"] if e["line"] == "瑪氏" and e["eip_po"]}
    e101, e102 = eips["PO202610101"]["id"], eips["PO202610102"]["id"]
    cn = db.get_conn()
    cn.execute("UPDATE mst_mars_products SET shelf_days = 365 WHERE yf_sku = 'M60003196'")
    cn.commit(); cn.close()

    def cmp(pdf=None, **body):
        if pdf is not None:
            return c.post("/api/shopee/yx/compare", data={"file": (io.BytesIO(pdf), "勇信.pdf"), "skip_eips": body.get("skip", "")},
                          content_type="multipart/form-data")
        return c.post("/api/shopee/yx/compare", json=body)

    def rows(res):
        return {(b["eip_po"], r["supplier_sku_id"]): (b, r) for b in res["eips"] for r in b["items"]}

    def plan():
        return {g["order_no"]: g for g in c.get("/api/shopee/short/plan", query_string={"from": "2026-01-01", "line": "瑪氏"}).get_json()["orders"]}

    def item(pl, sup):
        return next(x for x in pl[N]["items"] if x["supplier_sku_id"] == sup)

    check("沒傳檔 → 400", c.post("/api/shopee/yx/compare", json={}).status_code == 400)
    bad = c.post("/api/shopee/yx/compare", data={"file": (io.BytesIO(b"not a pdf"), "x.pdf")}, content_type="multipart/form-data")
    check("不是 PDF → 400、說打不開", bad.status_code == 400 and "無法開啟" in bad.get_json()["error"], str(bad.get_json()))
    pdf = yx_pdf([("PO202610101", "2026/10/13", [("60003196", [(1, "2027/06/01"), (1, "2027/01/01")])]),
                  ("PO209999999", "2026/10/13", [("12345678", 3, "2027/05/01")])])
    r = cmp(pdf)
    d = r.get_json()
    rw = rows(d) if r.status_code == 200 else {}
    b1, a = rw.get(("PO202610101", "M60003196_QMN"), ({}, {}))
    b2, b_ = rw.get(("PO202610102", "M10281171_RMN"), ({}, {}))
    check("收貨單號對到 EIP PO202610101：訂 2 箱、勇信出 2 箱，兩批效期分開讀", r.status_code == 200 and b1.get("in_pdf") and (a["ordered"], a["shipped"]) == (2, 2)
          and [(x["qty"], x["expiry"]) for x in a["batches"]] == [(1, "2027-06-01"), (1, "2027-01-01")], str(a)[:300])
    check("效期 2027-01-01 那箱早於蝦皮允收（到貨日 10/13＋效期天 365×2/5＝2027-03-08）→ 不算可出貨：效期不足、48 盒下修至 24（需求 5.3）",
          a["need"] == "2027-03-08" and [x["ok"] for x in a["batches"]] == [True, False] and (a["good"], a["short"], a["status"], a["final"], a["action"], a["reasons"])
          == (1, 1, "partial", 24, "下修至 24", ["效期不足"]), str(a)[:400])
    check("同一天同一倉已回填、表上卻沒有的 PO202610102 → 整份沒出：18 箱刪單", not b2.get("in_pdf")
          and (b_["ordered"], b_["shipped"], b_["status"], b_["final"], b_["action"], b_["reasons"]) == (18, 0, "none", 0, "刪單", ["庫存不足"]), str(b_)[:300])
    check("別天、還沒回填單號的安南三那份不在範圍內", len(d["eips"]) == 2 and not d["no_eip"], str([(b["eip_po"], b["filename"]) for b in d["eips"]]))
    check("表上的號碼系統裡找不到（可能是酷澎的單）→ 列出來", [u["eip_po"] for u in d["unknown"]] == ["PO209999999"], str(d["unknown"]))
    check("倉別讀不到的提醒不顯示（蝦皮看系統裡 EIP 的倉）、日期、倉、統計", not any("倉別" in w for w in d["page_warnings"]) and d["dates"] == ["2026-10-13"]
          and d["warehouses"] == ["TWA"] and (d["summary"]["partial"], d["summary"]["none"], d["summary"]["absent"], d["summary"]["changes"], d["summary"]["expiry_short"]) == (1, 1, 1, 2, 1),
          str(d["summary"]))
    r = cmp(pages=d["pages"], skip_eips=[e102])
    d2 = r.get_json()
    check("取消勾選 PO202610102（這次不核）→ 那份不算缺", r.status_code == 200 and next(b for b in d2["eips"] if b["eip_id"] == e102)["skipped"]
          and d2["summary"]["absent"] == 0 and d2["summary"]["changes"] == 1, str(d2["summary"]))
    r = c.post("/api/shopee/yx/apply", json={"pages": d["pages"], "skip_eips": [e102]})
    pl = plan()
    a = item(pl, "M60003196_QMN")
    check("確認（跳過 102）：寫入 1 筆第一次下修，來源勇信核對、48 → 24、效期不足、待通知，說明寫哪一天以前的幾箱不算", r.status_code == 200 and r.get_json()["written"] == 1
          and [(x["round"], x["source"], x["before"], x["after"], x["reasons"]) for x in a["rounds"]] == [(1, "勇信核對", 48, 24, ["效期不足"])]
          and a["pending"] and pl[N]["status"] == "待通知" and "勇信 PO202610101 出 2 箱，訂 2 箱；效期早於 2027-03-08 的 1 箱不算" in a["rounds"][0]["note"], str(a["rounds"]))
    check("跳過的那份沒寫", not item(pl, "M10281171_RMN")["rounds"])
    r = c.post("/api/shopee/yx/apply", json={"pages": d["pages"]})
    pl = plan()
    check("再確認一次（不跳過）：只多寫 102 那筆（18 → 0），101 一樣的不重寫", r.get_json()["written"] == 1
          and [(x["before"], x["after"]) for x in item(pl, "M10281171_RMN")["rounds"]] == [(18, 0)] and len(item(pl, "M60003196_QMN")["rounds"]) == 1, str(r.get_json()))
    both = rows(cmp(yx_pdf([("PO202610101", "2026/10/13", [("60003196", 1, "2027/01/01")]),
                            ("PO202610102", "2026/10/13", [("10281171", 18, "2026/11/01")])])).get_json())
    x1, x2 = both[("PO202610101", "M60003196_QMN")][1], both[("PO202610102", "M10281171_RMN")][1]
    check("勇信少出、出的那箱效期又不夠 → 兩個原因一起標（庫存不足、效期不足），刪單（需求 5.4）",
          (x1["shipped"], x1["good"], x1["status"], x1["final"], x1["reasons"]) == (1, 0, "none", 0, ["庫存不足", "效期不足"]), str(x1)[:300])
    check("商品總表沒有效期天（M10281171）→ 無法判斷效期，提醒、不扣，18 箱照出", (x2["status"], x2["final"], x2["reasons"]) == ("ok", None, [])
          and "沒有效期天" in x2["warn"] and x2["need"] == "", str(x2)[:300])
    over = cmp(yx_pdf([("PO202610101", "2026/10/13", [("60003196", 3, "2027/06/01")])])).get_json()
    o_ = rows(over)[("PO202610101", "M60003196_QMN")][1]
    check("勇信出得比訂得多（3 箱）→ 數量異常、不修改", (o_["status"], o_["action"], o_["final"]) == ("over", "數量異常，不修改", None), str(o_)[:200])
    fixed = yx_pdf([("PO202610101", "2026/10/13", [("60003196", 2, "2027/06/01")])])
    fx = cmp(fixed).get_json()
    check("重傳更正過的 PDF（出 2 箱）→ 101 不缺，顯示要取消前次缺貨；效期夠了不提醒",
          rows(fx)[("PO202610101", "M60003196_QMN")][1]["action"] == "取消前次缺貨" and not rows(fx)[("PO202610101", "M60003196_QMN")][1]["warn"])
    r = c.post("/api/shopee/yx/apply", json={"pages": fx["pages"]})
    pl = plan()
    check("確認後：101 那次缺貨取消、102 照舊刪單", r.get_json()["cleared"] == 1 and not item(pl, "M60003196_QMN")["rounds"]
          and item(pl, "M10281171_RMN")["final"] == 0, str(r.get_json()))
    r = c.post("/api/shopee/short/notify", json={"order_nos": [N]})
    msg = c.get(f"/api/shopee/short/notify/{r.get_json()['id']}").get_json()["messages"] if r.status_code == 200 else []
    check("通知蝦皮：賣場寫瑪氏、刪單那則有 M10281171", len(msg) == 1 and msg[0].startswith("【蝦皮缺貨刪單通知｜首次通知】") and "賣場：瑪氏" in msg[0]
          and "M10281171_RMN" in msg[0], str(msg))
    again = cmp(pdf).get_json()
    lk = rows(again)[("PO202610102", "M10281171_RMN")][1]
    r = c.post("/api/shopee/yx/apply", json={"pages": again["pages"]})
    pl = plan()
    check("通知過的鎖住：再核對顯示「已通知，不再修改」，確認也不動它", lk["locked"] and lk["action"] == "已通知，不再修改"
          and [(x["before"], x["after"], x["notified"]) for x in item(pl, "M10281171_RMN")["rounds"]] == [(18, 0, True)], str(lk)[:200])
    ov = next(g for g in orders(line="瑪氏")["orders"] if g["order_no"] == N)
    logs = c.get("/api/master/logs?scope=shopee&limit=800").get_json()["logs"]
    check("訂單總覽看得到瑪氏的通知狀態、歷程記下勇信核對", ov["notify_status"].startswith(("已通知", "待通知"))
          and any(l["field"] == "shopee_yx" for l in logs) and any(l["field"] == "shopee_short" and "勇信核對" in (l.get("note") or "") for l in logs),
          ov["notify_status"])



def flow_tests(c, orders):
    print("\n【16】改版：作業流程格子（每張單的下一步）、一個上傳入口自動認檔（Jerry 2026-10-13）")
    o = orders()
    by = {g["order_no"]: g for g in o["orders"]}
    st = {x["key"]: x["count"] for x in o["stages"]}
    check("瑪氏竹運採購進貨（擋下來的）→ 待產採購單", by["POTWCTWA26101100G"]["todo"] == ["buy"] and by["POTWCTWA26101100G"]["next"] == "產出 EIP 採購單",
          str(by["POTWCTWA26101100G"]["todo"]))
    check("竹運出貨還沒產拋單檔 → 待產拋單檔", by["POTWCTWA26101500J"]["todo"] == ["ship"], str(by["POTWCTWA26101500J"]["todo"]))
    check("竹運出貨產過拋單檔、缺貨也通知過 → 沒有待辦（已處理）", by["POTWCTWA26101500H"]["todo"] == [] and by["POTWCTWA26101500H"]["next"] == "",
          str(by["POTWCTWA26101500H"]["todo"]))
    check("瑪氏直送產過 EIP、還沒回填 → 待回填 EIP 採購單號", by["POTWCTWX26101100F"]["todo"] == ["eip"], str(by["POTWCTWX26101100F"]["todo"]))
    check("格子數字＝有這項待辦的訂單數；已處理＝沒有待辦的", all(st[k] == sum(1 for g in o["orders"] if k in g["todo"]) for k in ("diff", "buy", "ship", "eip", "notify"))
          and st["done"] == sum(1 for g in o["orders"] if not g["todo"]), str(st))

    def drop(data, name):
        return c.post("/api/shopee/drop", data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")
    r = drop(purchase_xlsx([("91_1", "4987176232878_CNN", "幫寶適 清新幫 拉拉褲", 30, "Pcs")], pr="PRTWCOTWA202610300001", date="2026-10-30"), "PurchaseOrder_x.xlsx")
    check("上傳入口：蝦皮採購單 → 認成蝦皮訂單，回預覽（要選履約方式）", r.status_code == 200 and r.headers.get("X-Shopee-Kind") == "orders"
          and r.get_json()["needs_fulfil"] and r.get_json()["kind_text"] == "採購單", str(r.headers.get("X-Shopee-Kind")))
    c.post("/api/shopee/import/cancel", json={"upload_id": r.get_json()["upload_id"]})
    r = drop(price_xlsx({"價格本整合": [("4987176232878C", "包", "幫寶適 清新幫 拉拉褲", 3), ("4987176232878R", "箱", "幫寶適 清新幫 箱購", 1)]}), "PG價格本整合.xlsx")
    check("上傳入口：價格本 → 認成價格本，直接更新", r.status_code == 200 and r.headers.get("X-Shopee-Kind") == "conv" and "寶僑" in r.get_json()["lines"],
          str((r.status_code, r.headers.get("X-Shopee-Kind"), r.get_json())))
    r = drop(yx_pdf([("PO202610101", "2026/10/13", [("60003196", 2, "2027/06/01")])]), "勇信.pdf")
    check("上傳入口：PDF → 認成勇信配送明細表，回比對結果", r.status_code == 200 and r.headers.get("X-Shopee-Kind") == "yx" and r.get_json()["eips"],
          str(r.headers.get("X-Shopee-Kind")))
    wb = openpyxl.Workbook(); wb.active.append(["隨便", "一個", "表"]); buf = io.BytesIO(); wb.save(buf)
    r = drop(buf.getvalue(), "不知道是什麼.xlsx")
    check("上傳入口：認不出來的 → 400，說可以上傳哪幾種", r.status_code == 400 and "認不出" in r.get_json()["error"] and "勇信配送明細表" in r.get_json()["error"],
          str(r.get_json()))
    check("通知紀錄可以單獨讀（資料與設定頁用）", isinstance(c.get("/api/shopee/short/notifies").get_json()["notifies"], list))

    def side(no, **q):
        return c.get("/api/shopee/order", query_string={"no": no, "from": "2026-01-01", **q})
    L = side("POTWCTWA26100514L").get_json()
    check("右側面板：竹運採購進貨那張 → 品項、採購單計畫（轉換率、下採）、兩份 EIP、差異紀錄、歷程都在", L["order"]["order_no"] == "POTWCTWA26100514L"
          and L["buy"] and L["buy"]["fulfil"] == "竹運採購進貨" and {p["supplier_sku_id"] for p in L["buy"]["items"]} >= {"4987176232878_CNN"}
          and len(L["eips"]) == 2 and L["diffs"] and L["logs"] and L["ship"] and L["short"]["items"], str({k: bool(v) for k, v in L.items()}))
    check("竹運採購進貨、入庫單到了 → 右側卡片也有拋單檔那段；下一步同時有回填 EIP（或已回填）和產出拋單檔", L["ship"]["order_no"] == "POTWCTWA26100514L"
          and "ship" in L["order"]["todo"], str(L["order"]["todo"]))
    check("右側面板：缺貨紀錄在（寶僑兩次都通知過）", [r["round"] for r in next(x for x in L["short"]["items"] if x["supplier_sku_id"] == "4987176232878_CNN")["rounds"]] == [1, 2])
    J = side("POTWCTWA26101500J").get_json()
    check("右側面板：竹運出貨那張 → 出貨量、效期批次（跟清單同一組篩選算）", J["ship"] and J["ship"]["order_no"] == "POTWCTWA26101500J" and J["buy"] is None
          and "alloc" in J["ship"]["items"][0] and not J["ship"].get("alone"), str(J["ship"])[:200])
    J2 = side("POTWCTWA26101500J", **{"from": "2026-12-01"}).get_json()
    check("不在目前篩選範圍裡的竹運出貨單 → 只就這張單試算、標出來", J2["ship"] and J2["ship"].get("alone"))
    N_ = side("POTWCTWA26100514N").get_json()
    check("右側面板：瑪氏直送那張 → 兩份 EIP 都回填了、勇信核對的缺貨在", all(e["eip_po"] and e["mars"] for e in N_["eips"]) and len(N_["eips"]) == 2
          and any(r["source"] == "勇信核對" for x in N_["short"]["items"] for r in x["rounds"]), str(N_["eips"]))
    check("找不到的單 → 404；沒給單號 → 400", side("POXXXXXXXX").status_code == 404 and c.get("/api/shopee/order").status_code == 400)

    print("\n【20】上傳價格本後直接講：還在等價格的品項找到了沒有（Becky 2026-10-16「更新價格本了 是不是得重拋採購單？」）")

    def up(url, data, name):
        return c.post(url, data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")

    def imp(data, fulfil):
        uid = up("/api/shopee/import/preview", data, "PurchaseOrder_t.xlsx").get_json()["upload_id"]
        return c.post("/api/shopee/import/commit", json={"upload_id": uid, "fulfil": fulfil}).get_json()

    def plan(fulfil):
        return {g["order_no"]: g for g in c.get("/api/shopee/purchase/plan", query_string={"fulfil": fulfil}).get_json()["orders"]}

    def item(g, sup):
        return next(p for p in g["items"] if p["supplier_sku_id"] == sup)
    imp(purchase_xlsx([("90000000001_1", "1077777_PNN", "新品 五月花 甲", 24, "Pcs"), ("90000000001_2", "1055555_PNN", "新品 五月花 乙", 12, "Pcs")],
                      pr="PRTWCOTWA202610169001", po="POTWCTWA26101690Z", date="2026-10-28"), "供應商直送")
    old = price_xlsx({"紙品價格本": [("1022257P", "瓶", "橘子工坊浴廁清潔劑", 12)]})
    r = up("/api/shopee/conv/import", old, "紙潔價格本整合.xlsx").get_json(); pz = r["pending"]["紙潔"]
    miss = {x["code3"]: x["name"] for x in pz["missing"]}
    check("價格本還沒有新品 → 列出這兩個找不到、帶品名，沒有找到的", miss.get("1077777P") == "新品 五月花 甲" and miss.get("1055555P") == "新品 五月花 乙"
          and not pz["fixed"] and pz["waiting"] >= 2, str(pz))
    new = price_xlsx({"紙品價格本": [("1022257P", "瓶", "橘子工坊浴廁清潔劑", 12), ("1077777P", "包", "新品 五月花 甲", 12), ("1055555P", "包", "新品 五月花 乙", None)]})
    r = up("/api/shopee/conv/import", new, "紙潔價格本整合.xlsx").get_json(); pz = r["pending"]["紙潔"]
    check("重新整理後再傳（PM 補上了）→ 原本找不到 2 個、這次找到 1 個；另 1 個有了但沒箱入數，分開講", pz["was_blocked"] >= 2 and [x["code3"] for x in pz["fixed"]] == ["1077777P"]
          and [x["code3"] for x in pz["no_box"]] == ["1055555P"] and not any(x["code3"] in ("1077777P", "1055555P") for x in pz["missing"]), str(pz))
    pl = plan("供應商直送").get("POTWCTWA26101690Z")
    check("不用重拋：訂單那邊直接對到新的價格本（甲有轉換率、乙寫沒有箱入數）", pl is not None and item(pl, "1077777_PNN")["conv"] == 12
          and any("沒有箱入數" in x for x in item(pl, "1055555_PNN")["problems"]), str(pl)[:300] if pl else "沒有這張單")
    r = up("/api/shopee/conv/import", price_xlsx({"價格本整合": [("4987176232878C", "包", "幫寶適", 4)]}), "PG價格本整合.xlsx").get_json()
    check("別的線別（寶僑）上傳不影響紙潔那兩個的狀態，只講寶僑自己的", set(r["pending"]) == {"寶僑"}, str(r["pending"]))

    print("\n【21】改履約方式時，已做好的 EIP、拋單檔系統一起處理（Jerry 2026-10-22「要有按鈕讓他們修改」）")
    A, B = "PRTWCOTWA202701050001", "PRTWCOTWA202701050002"
    imp(purchase_xlsx([("95_1", "4987176232878_CNN", "幫寶適 清新幫 拉拉褲", 40, "Pcs")], pr=A, date="2027-01-05"), "竹運採購進貨")
    imp(purchase_xlsx([("95_1", "4987176232878_CNN", "幫寶適 清新幫 拉拉褲", 40, "Pcs")], pr=B, wh="TWX", date="2027-01-06"), "竹運採購進貨")
    r = c.post("/api/shopee/purchase/generate", json={"fulfil": "竹運採購進貨", "order_nos": [A, B]})
    eip = next((e for e in c.get("/api/shopee/eips", query_string={"from": "2027-01-01"}).get_json()["eips"] if A in e["orders"]), {})
    check("兩張單一起產出：同一份 EIP（小單位，竹運採購進貨不換箱）、共 80 包", r.status_code == 200 and sorted(eip.get("orders", [])) == [A, B] and eip.get("qty") == 80, str((eip.get("qty"), eip.get("rows"))))
    import json
    cn = db.get_conn()          # 模擬 2026-10-21 以前產的 EIP：沒記每張單各多少（parts）
    rows_ = json.loads(cn.execute("SELECT rows_json FROM shp_eip WHERE id = ?", (eip["id"],)).fetchone()["rows_json"])
    for x in rows_:
        x.pop("parts", None)
    cn.execute("UPDATE shp_eip SET rows_json = ? WHERE id = ?", (json.dumps(rows_, ensure_ascii=False), eip["id"])); cn.commit(); cn.close()
    chk = c.post("/api/shopee/fulfil/check", json={"order_nos": [A], "fulfil": "竹運出貨"}).get_json()["orders"]
    check("改之前先問：A 在那份 EIP 裡、還有另外 1 張單、還沒回填", len(chk) == 1 and chk[0]["eips"][0]["id"] == eip["id"] and chk[0]["eips"][0]["others"] == 1
          and chk[0]["eips"][0]["eip_po"] == "" and chk[0]["dispatch"] is None, str(chk))
    check("改成同一個履約方式、或沒產過東西的單 → 不用提醒", c.post("/api/shopee/fulfil/check", json={"order_nos": [A], "fulfil": "竹運採購進貨"}).get_json()["orders"] == [])
    r = c.post("/api/shopee/fulfil", json={"order_nos": [A], "fulfil": "竹運出貨"}).get_json()
    eips2 = c.get("/api/shopee/eips", query_string={"from": "2027-01-01"}).get_json()["eips"]
    e2 = next((e for e in eips2 if e["id"] == eip["id"]), {})
    check("A 改成竹運出貨 → 從那份 EIP 拿掉；那份只留 B 的 40 包、檔名照 B 重算（安南三）、標「重新做過」",
          r["count"] == 1 and "已從 EIP" in r["handled"][0] and e2.get("orders") == [B] and e2.get("qty") == 40 and "安南三" in e2.get("filename", "")
          and "觀音" not in e2.get("filename", "") and "拿掉 " + A in e2.get("note", "") and "重新下載再上傳" in e2.get("note", ""), str((r, e2))[:500])
    ov = {g["order_no"]: g for g in orders(**{"from": "2027-01-01"})["orders"]}
    check("A 移到待產拋單檔（不用產採購單）、B 還在待回填 EIP", ov[A]["todo"] == ["ship"] and ov[A]["po_status"] == "不需產出" and ov[B]["todo"] == ["eip"],
          str((ov[A]["todo"], ov[B]["todo"])))
    lg = [l for l in c.get("/api/master/logs?scope=shopee&limit=300").get_json()["logs"] if l["field"] == "shopee_fulfil" and l["po_number"] == A]
    check("歷程：履約方式竹運採購進貨 → 竹運出貨，備註寫從哪份 EIP 拿掉", lg and lg[0]["new_value"] == "竹運出貨" and "已從 EIP" in (lg[0]["note"] or ""), str(lg[:1]))
    c.put(f"/api/shopee/eips/{eip['id']}", json={"eip_po": "PO202701001"})
    chk = c.post("/api/shopee/fulfil/check", json={"order_nos": [B], "fulfil": "供應商直送"}).get_json()["orders"]
    check("B 已回填單號再改：提醒裡帶 EIP 採購單號、這份只有 B", chk and chk[0]["eips"][0]["eip_po"] == "PO202701001" and chk[0]["eips"][0]["others"] == 0, str(chk))
    r = c.post("/api/shopee/fulfil", json={"order_nos": [B], "fulfil": "供應商直送"}).get_json()
    check("B 改成供應商直送 → 那份 EIP 只有它，整份取消；B 回到待產採購單", "整份" in r["handled"][0]
          and not any(e["id"] == eip["id"] for e in c.get("/api/shopee/eips", query_string={"from": "2027-01-01"}).get_json()["eips"])
          and {g["order_no"]: g for g in orders(**{"from": "2027-01-01"})["orders"]}[B]["todo"] == ["buy"], str(r))
    H = "POTWCTWA26101500H"
    chk = c.post("/api/shopee/fulfil/check", json={"order_nos": [H], "fulfil": "竹運採購進貨"}).get_json()["orders"]
    check("竹運出貨產過拋單檔、有入庫單，改成竹運採購進貨 → 還是要拋單，不取消拋單檔、不用提醒", chk == [], str(chk))
    chk = c.post("/api/shopee/fulfil/check", json={"order_nos": [H], "fulfil": "供應商直送"}).get_json()["orders"]
    check("改成供應商直送（不走竹運了）→ 提醒帶拋單檔產出時間", chk and chk[0]["dispatch"] and chk[0]["dispatch"]["created_at"], str(chk))
    r = c.post("/api/shopee/fulfil", json={"order_nos": [H], "fulfil": "供應商直送"}).get_json()
    check("改了 → 這張的拋單檔取消、移到待產採購單", "取消這張單的竹運拋單檔" in "".join(r["handled"])
          and c.post("/api/shopee/fulfil/check", json={"order_nos": [H], "fulfil": "供應商直送"}).get_json()["orders"] == []
          and "buy" in {g["order_no"]: g for g in orders(**{"from": "2026-01-01"})["orders"]}[H]["todo"], str(r))
    chk = c.post("/api/shopee/fulfil/check", json={"order_nos": ["POTWCTWA26100514L"], "fulfil": "竹運出貨"}).get_json()["orders"]
    check("竹運採購進貨改竹運出貨 → 提醒「入庫單到了會自動進待產拋單檔，不用改」", chk and chk[0]["buy_to_ship"], str(chk))



def trash_tests(c, orders):
    print("\n【22】刪除／還原訂單、匯出訂單明細（Alice 2026-10-30；Jerry：大家都能刪、要能還原、一張工作表、缺貨整列黃底）")

    def up(data, name):
        return c.post("/api/shopee/import/preview", data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")
    X, XP = "POTWCTWA27020190X", "PRTWCOTWA202702019001"
    raw = purchase_xlsx([("90000000201_1", "1078888_PNN", "贈品 五月花 試用包", 1500, "Pcs"), ("90000000201_2", "1078889_PNN", "贈品 五月花 小包", 12, "Pcs")],
                        pr=XP, po=X, date="2027-02-03")
    c.post("/api/shopee/import/commit", json={"upload_id": up(raw, "PurchaseOrder_gift.xlsx").get_json()["upload_id"], "fulfil": "供應商直送"})
    check("先匯入一張贈品單：在清單、待產採購單", X in {g["order_no"] for g in orders()["orders"]}
          and "buy" in {g["order_no"]: g for g in orders()["orders"]}[X]["todo"])
    r = c.post("/api/shopee/order/delete", json={"order_no": X, "reason": ""})
    gx = {g["order_no"]: g for g in orders(deleted="1")["orders"]}.get(X)
    check("原因選填（Jerry 2026-10-30）：沒填也能刪，記成空白", r.status_code == 200 and gx and gx["trash"]["reason"] == "", str(r.get_json()))
    c.post("/api/shopee/order/restore", json={"order_no": X})
    n_buy = next(s_["count"] for s_ in orders()["stages"] if s_["key"] == "buy")
    r = c.post("/api/shopee/order/delete", json={"order_no": X, "reason": "贈品不用出"})
    check("填原因刪除 → 成功，訊息講怎麼還原", r.status_code == 200 and "顯示已刪除" in r.get_json()["message"], str(r.get_json()))
    v = orders()
    check("刪掉的單：清單沒有、待產採購單少一張、回傳有幾張被刪", X not in {g["order_no"] for g in v["orders"]}
          and next(s_["count"] for s_ in v["stages"] if s_["key"] == "buy") == n_buy - 1 and v["deleted_count"] == 1, str(v["deleted_count"]))
    plan_ = c.get("/api/shopee/purchase/plan", query_string={"fulfil": "供應商直送", "from": "2026-01-01"}).get_json()["orders"]
    check("產 EIP 的計畫也跳過刪掉的單", X not in {g["order_no"] for g in plan_})
    vd = orders(deleted="1")
    gx = {g["order_no"]: g for g in vd["orders"]}.get(X)
    check("「顯示已刪除」才列出來：帶誰刪的、原因，沒有待辦；數字列不算它", gx and gx["trash"]["reason"] == "贈品不用出" and gx["trash"]["deleted_by"] == "Jerry"
          and gx["todo"] == [] and vd["summary"]["orders"] == v["summary"]["orders"], str(gx and gx["trash"]))
    sd = c.get("/api/shopee/order", query_string={"no": X}).get_json()
    check("刪掉的單還是打得開訂單視窗（看得到是誰刪的）", sd.get("order", {}).get("trash", {}).get("reason") == "贈品不用出", str(sd)[:200])
    c.post("/api/shopee/import/commit", json={"upload_id": up(raw, "PurchaseOrder_gift.xlsx").get_json()["upload_id"], "fulfil": "供應商直送"})
    check("同一份蝦皮檔再上傳一次 → 刪掉的單不會跑回來", X not in {g["order_no"] for g in orders()["orders"]})
    r = c.post("/api/shopee/order/delete", json={"order_no": X, "reason": "重複的單"})
    check("已經刪掉的再刪 → 講已經刪除了", r.status_code == 400 and "已經刪除" in r.get_json()["error"], str(r.get_json()))
    made = next((g for g in orders()["orders"] if g["po_status"] in ("已產出", "已回填")), None)
    r = c.post("/api/shopee/order/delete", json={"order_no": made["order_no"], "reason": "測試"})
    check("已產出 EIP 的單不能刪，要先取消產出", r.status_code == 400 and "取消產出" in r.get_json()["error"], str(r.get_json()))
    r = c.post("/api/shopee/order/restore", json={"order_no": X})
    v = orders()
    check("還原 → 回到清單、回到待產採購單、刪除數歸零", r.status_code == 200 and "buy" in {g["order_no"]: g for g in v["orders"]}[X]["todo"] and v["deleted_count"] == 0,
          str(r.get_json()))
    lg = [l for l in c.get("/api/master/logs?scope=shopee&limit=300").get_json()["logs"] if l["po_number"] == X and l["field"] in ("shopee_delete", "shopee_restore")]
    check("歷程：記刪除（原因）和還原", {l["field"] for l in lg} == {"shopee_delete", "shopee_restore"}
          and any(l["field"] == "shopee_delete" and l["reason"] == "贈品不用出" for l in lg), str(lg)[:300])

    r = c.post("/api/shopee/orders/export", json={"order_nos": []})
    check("沒勾就匯出 → 擋下來", r.status_code == 400)
    sh = next((g for g in orders()["orders"] if (g["qty_short"] or 0) > 0), None)
    nos = [X] + ([sh["order_no"]] if sh else [])
    r = c.post("/api/shopee/orders/export", json={"order_nos": nos, "from": "2026-01-01"})
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    head = [x.value for x in ws[1]]
    rows = [[x.value for x in row] for row in ws.iter_rows(min_row=2)]
    col = {h: k for k, h in enumerate(head)}
    check("匯出訂單明細：一張工作表「訂單明細」、檔名寫幾張、一品一列、每列帶訂單編號",
          r.status_code == 200 and len(openpyxl.load_workbook(io.BytesIO(r.data)).sheetnames) == 1 and ws.title == "訂單明細"
          and f"_{len(nos)}" in r.headers.get("Content-Disposition", "") and sum(1 for x in rows if x[col["訂單編號"]] == X) == 2,
          str((head[:8], len(rows))))
    check("欄位照訂單視窗的品項表", all(h in head for h in ("訂單編號", "線別", "倉別", "到貨日", "履約方式", "永豐料號", "品名", "入庫單號", "蝦皮數量",
                                                          "轉換率", "下採數量", "EIP 數量", "出貨量", "缺貨", "效期", "調整原因")), str(head))
    xr = next(x for x in rows if x[col["永豐料號"]] == "1078888_PNN")
    check("那品的數字：蝦皮數量 1500、倉別寫中文", xr[col["蝦皮數量"]] == 1500 and "觀音" in (xr[col["倉別"]] or ""), str(xr))
    if sh:
        fills = [(row[col["缺貨"]].value, row[0].fill.fgColor.rgb) for row in ws.iter_rows(min_row=2) if row[col["訂單編號"]].value == sh["order_no"]]
        check("有缺貨的那列整列黃底，沒缺貨的不塗", any((q or 0) > 0 for q, _f in fills) and all(f.endswith("FFF2A8") == bool(q) for q, f in fills)
              and int(r.headers.get("X-Shopee-Short", 0)) == sum(1 for q, _f in fills if q), str(fills))
    else:
        check("（這份測試資料裡要有一張有缺貨的單）", False)



def po_pdf_tests(c, orders):
    print("\n【23】蝦皮採購單 PDF 畫缺貨（Alice 2026-10-30；Jerry：不出貨整列畫線，短出劃掉數量、寫實際數量、蓋章）")
    import pymupdf

    def up(*files):
        return c.post("/api/shopee/po_pdf", data={"file": [(io.BytesIO(d), n) for d, n in files]}, content_type="multipart/form-data")
    g = next((x for x in orders()["orders"] if (x["qty_short"] or 0) > 0 and x["po_id"]), None)
    check("（測試資料裡要有一張有缺貨、有 PO 的單）", g is not None)
    if g is None:
        return
    sd = c.get("/api/shopee/order", query_string={"no": g["order_no"], "from": "2026-01-01"}).get_json()
    by = lambda grp: {x["order_id"]: x for x in (grp or {}).get("items", [])}
    ship, short = by(sd["ship"]), by(sd["short"])
    alive = [i for i in sd["order"]["items"] if not i["removed"]]
    final = {i["supplier_sku_id"]: (ship[i["id"]]["final"] if i["id"] in ship else short[i["id"]]["final"] if i["id"] in short else None) for i in alive}
    want_no = sum(1 for i in alive if final[i["supplier_sku_id"]] == 0)
    want_short = sum(1 for i in alive if final[i["supplier_sku_id"]] not in (None, 0) and final[i["supplier_sku_id"]] < i["qty"])
    rows = [(i["supplier_sku_id"], i["qty"]) for i in alive]
    r = up((po_pdf(g["po_id"], rows), f"678050_{g['po_id']}_PO.pdf"), (yx_pdf([("PO209901001", "2026/10/20", [("10001", 1, "2027/01/01")])]), "勇信.pdf"),
           (po_pdf(g["po_id"], rows, text=False), "印出來的.pdf"), (po_pdf("POTWCTWA9999999Z9", rows), "別張.pdf"))
    d = r.get_json()
    o = (d.get("orders") or [{}])[0]
    check("上傳區丟四個 PDF：蝦皮採購單收下、對到那張單，照缺貨紀錄算出不出貨幾品、短出幾品", r.status_code == 200 and len(d["orders"]) == 1
          and o["order_no"] == g["order_no"] and o["rows"] == len(rows) and (o["noship"], o["short"]) == (want_no, want_short) and o["noship"] + o["short"] > 0,
          str((o, want_no, want_short))[:300])
    check("勇信配送明細表不收、交回去給勇信核對", d["others"] == ["勇信.pdf"], str(d["others"]))
    check("自己印成 PDF 的（讀不到字）講清楚要上傳後台下載的原檔；系統沒有的 PO 也講", any("印出來的.pdf" in e and "原檔" in e for e in d["errors"])
          and any("POTWCTWA9999999Z9" in e for e in d["errors"]), str(d["errors"]))
    sd = c.get("/api/shopee/order", query_string={"no": g["order_no"], "from": "2026-01-01"}).get_json()
    check("訂單視窗帶採購單 PDF：誰上傳、要畫哪幾列", sd["po_pdf"] and sd["po_pdf"]["uploaded_by"] == "Jerry" and len(sd["po_pdf"]["marks"]) == want_no + want_short,
          str(sd["po_pdf"])[:300])
    r = c.get(f"/api/shopee/po_pdf/{g['order_no']}/file", query_string={"from": "2026-01-01"})
    doc = pymupdf.open(stream=r.data, filetype="pdf")
    reds = sum(1 for p in doc for dr in p.get_drawings() if dr.get("color") and dr["color"][0] > 0.8 and dr["color"][1] < 0.2)
    text = "".join(p.get_text() for p in doc)
    acts = [str(final[s_]) for s_, q in rows if final[s_] not in (None, 0) and final[s_] < q]
    check("下載畫好的採購單：檔名加「_已畫」、不出貨整列紅線、短出兩格劃掉（各一條紅線）並寫實際數量、沒上傳印章就畫紅框「業務專用章」",
          r.status_code == 200 and "_%E5%B7%B2%E7%95%AB.pdf" in r.headers.get("Content-Disposition", "") and int(r.headers["X-Shopee-Drawn"]) == want_no + want_short
          and reds >= want_no + 3 * want_short and all(a in text for a in acts) and ("業務專用章" in text) == (want_short > 0), str((reds, acts, r.headers.get("X-Shopee-Drawn"))))
    sd = c.get("/api/shopee/order", query_string={"no": g["order_no"], "from": "2026-01-01"}).get_json()
    lg = [l for l in sd["logs"] if l["field"] in ("shopee_po_pdf", "shopee_po_draw")]
    check("記下誰下載過、歷程有上傳和畫採購單", sd["po_pdf"]["drawn_by"] == "Jerry" and {l["field"] for l in lg} == {"shopee_po_pdf", "shopee_po_draw"}, str(lg)[:200])
    r = c.get("/api/shopee/po_pdf/POTWCTWA0000000X0/file")
    check("沒上傳過的單按下載 → 講要先上傳", r.status_code == 404 and "還沒上傳" in r.get_json()["error"])

    print("\n【24】待回填 EIP 採購單號勾選 → 一次下載勾的那幾份（Jerry 2026-10-30「就跟其他部份一樣」）")
    eips = c.get("/api/shopee/eips", query_string={"from": "2026-01-01"}).get_json()["eips"]
    pick = eips[:2]
    r = c.get("/api/shopee/eips.zip", query_string={"ids": ",".join(str(e["id"]) for e in pick)})
    names = zipfile.ZipFile(io.BytesIO(r.data)).namelist() if r.status_code == 200 else []
    check("勾兩份 → EIP 檔壓縮檔裡就這兩份、檔名寫勾選幾份", len(names) == 2 and "%E5%8B%BE%E9%81%B82%E4%BB%BD" in r.headers.get("Content-Disposition", ""),
          str((r.status_code, names, r.headers.get("Content-Disposition"))))
    check("沒勾就下載 EIP 檔 → 擋下來", c.get("/api/shopee/eips.zip").status_code == 400)
    mf = [e for e in eips if e["line"] == "瑪氏" and e["eip_po"]]
    other = next((e for e in eips if e["line"] != "瑪氏"), None)
    if mf and other:
        r = c.get("/api/shopee/mars_po.zip", query_string={"ids": f"{mf[0]['id']},{other['id']}"})
        check("勾一份已回填的瑪氏＋一份寶僑 → 瑪氏採購單只有那 1 份（不管到貨日）", r.status_code == 200
              and len(zipfile.ZipFile(io.BytesIO(r.data)).namelist()) == 1 and r.headers["X-Shopee-Files"] == "1", str(r.status_code))
        r = c.get("/api/shopee/mars_po.zip", query_string={"ids": str(other["id"])})
        check("勾的裡面沒有已回填的瑪氏 → 講「勾選的檔裡沒有」", r.status_code == 400 and "勾選的檔裡沒有" in r.get_json()["error"], str(r.get_json()))
    else:
        check("（測試資料裡要有已回填的瑪氏 EIP 和別的線別）", False, str([(e["line"], e["eip_po"]) for e in eips]))


if __name__ == "__main__":
    main()
