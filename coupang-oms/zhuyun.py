"""竹運出貨拋檔及料號對照表轉換（網址 /zhuyun）。

Chloe 2026-09 的需求：傳一份 Kate 系統匯出的整合表（Coupang_PO_Integration），只抓線別是寶僑的列，
依「到貨日＋到貨倉」一組，每組產兩個檔：
  1. 指定到貨日20260918,酷澎_PG.xlsx          訂單系統拋檔用（給竹運）
     （同一天有兩個以上的倉會撞名，那時才加 (倉)：指定到貨日20260823,酷澎_PG(TXRC29).xlsx）
     A 訂單編號＝酷澎 PO  B 收件人、C 收件人手機 固定字  D 收件地址＝整合表地址  E 料號＝永豐料號
     F 數量＝出貨數量  G 單價＝酷澎下單單價(含稅)  H 金額小計＝F×G  I 出貨備註、J 運費 空白
  2. 永豐料號對照表_0918交貨(TAO5).xlsx       倉庫對不到資料時用
     A PO No.  B SKU No.  C 條碼  D 料號（永豐料號）  E 品名  F 數量＝出貨數量  G 單位  H 轉換率＝箱入數
     I 箱數＝F÷H 無條件進位（Jerry：範例寫 12 是手打的，照進位算 20）  J 允收效期 空白
跟採購表轉換一樣是獨立工具：檔案傳上來直接轉，不進訂單資料庫。
收件人固定「酷澎股份有限公司」。手機**寶僑自己一張倉別→手機表**（Jerry 2026-10-05：瑪氏跟寶僑能去的倉不同，不跟瑪氏的倉庫資料表共用），
頁面上直接填、也能上傳 PG_Ship-to 那種表（FC／TEL 兩欄）；寫法統一成拋檔範例的 02-55927598、0911556291
（+886-02-55927598 → 02-55927598、+886-986368794 → 0986368794、886-988705486 → 0988705486）。沒填的倉用固定字。
固定字與倉別手機都放 config.json 的 zhuyun 區。
"""
import collections
import datetime as _dt
import io
import json
import math
import os
import re
import zipfile

import openpyxl
from flask import Blueprint, jsonify, render_template, request, send_file, session

import db
import importer
from master.common import _log as _mst_log, _operator as _mst_operator
from normalize import norm_text

zhuyun_bp = Blueprint("zhuyun", __name__)

LINE = "寶僑"
DEFAULTS = {"recipient": "酷澎股份有限公司", "phone": "02-55927598"}
WH_CODE_RE = re.compile(r"^[A-Z0-9]{2,12}$")
# 設定走 app.py 的 load_json／save_json（SQLite 放 config.json、PostgreSQL 存資料庫）；app 匯入這個模組時那兩個
# 函式還沒定義，所以用到時才去 app 模組拿
def _app_json():
    import app as _app
    return _app.load_json, _app.save_json


def _zy_cfg():
    load_json, _ = _app_json()
    cfg = load_json("config.json", {}) or {}
    return cfg, (cfg.get("zhuyun") or {})


def _load_settings():
    _cfg, z = _zy_cfg()
    s = dict(DEFAULTS); s.update({k: v for k, v in z.items() if k in DEFAULTS and isinstance(v, str)})
    return s


def _save_settings(new):
    load_json, save_json = _app_json()
    cfg, z = _zy_cfg()
    z = dict(z); z.update(new); cfg["zhuyun"] = z
    save_json("config.json", cfg)


def load_phones():
    """寶僑各倉的收件人手機：{倉別: 手機}，已經是拋檔要的寫法。"""
    _cfg, z = _zy_cfg()
    ph = z.get("phones") if isinstance(z.get("phones"), dict) else {}
    return {str(k).upper(): str(v) for k, v in ph.items() if str(v).strip()}


def save_phones(phones):
    _load, save_json = _app_json()
    cfg, z = _zy_cfg()
    z = dict(z); z["phones"] = dict(sorted(phones.items())); cfg["zhuyun"] = z
    save_json("config.json", cfg)


def pg_phone(raw):
    """統一成拋檔範例的寫法：+886-02-55927598 → 02-55927598、+886-0911556291 → 0911556291、
    +886-986368794 → 0986368794、886-988705486 → 0988705486、02-5592-7598 → 02-55927598。看不懂（有分機、文字）照原樣。"""
    t = str(raw or "").strip().replace("＋", "+")
    if not t:
        return ""
    t = re.sub(r"^\+?886[\s\-]?", "", t)                       # 去國碼
    parts = [x for x in re.split(r"[\s\-().／（）]+", t) if x]
    digits = "".join(parts)
    if not digits.isdigit():
        return str(raw).strip()
    if not digits.startswith("0"):
        digits = "0" + digits; parts = [("0" + parts[0])] + parts[1:]
    if digits.startswith("09"):
        return digits
    if len(parts) >= 2:
        return f"{parts[0]}-{''.join(parts[1:])}"
    if len(digits) == 10 and digits.startswith("02"):            # 0255927598 → 02-55927598
        return f"02-{digits[2:]}"
    if len(digits) == 9:
        return f"{digits[:2]}-{digits[2:]}"
    return digits


PH_HEADERS = {"fc": "code", "倉別": "code", "倉": "code", "倉庫": "code", "到貨倉別": "code", "warehouse": "code",
              "tel": "phone", "電話": "phone", "手機": "phone", "收件人手機": "phone", "phone": "phone", "聯絡電話": "phone"}


def parse_phone_sheet(fileobj):
    """PG_Ship-to 那種表：一列一個倉，表頭 FC／TEL（或 倉別／電話）。回 {倉別: 手機}。"""
    try:
        wb = openpyxl.load_workbook(fileobj, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"打不開這個 Excel：{exc}") from exc
    for ws in wb.worksheets:
        for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            cols = {}
            for i, v in enumerate(row):
                f = PH_HEADERS.get(norm_text(v).lower().replace(" ", ""))
                if f and f not in cols:
                    cols[f] = i
            if "code" in cols and "phone" in cols:
                out = {}
                for raw in ws.iter_rows(min_row=idx + 1, values_only=True):
                    code = norm_text(raw[cols["code"]] if cols["code"] < len(raw) else "").upper()
                    ph = pg_phone(raw[cols["phone"]] if cols["phone"] < len(raw) else "")
                    if code and ph:
                        out[code] = ph
                return out
    raise ValueError("找不到倉別手機表：第一列要有「FC」（或倉別）和「TEL」（或電話）兩欄。")


def warehouse_phones():
    return load_phones()


def _is_pg(line):
    return norm_text(line) == LINE


def group_rows(rows):
    """只留寶僑，依 (到貨日, 倉) 分組。回傳 [{date, warehouse, rows}]，到貨日舊的在前。"""
    groups = collections.OrderedDict()
    for r in rows:
        if not _is_pg(r.get("line")):
            continue
        key = (r.get("delivery_date") or "", r.get("warehouse") or "")
        groups.setdefault(key, []).append(r)
    out = []
    for (d, wh), rs in sorted(groups.items()):
        rs.sort(key=lambda x: x.get("row_no") or 0)          # 照整合表原本的順序（Chloe 的範例就是這樣）
        out.append({"date": d, "warehouse": wh, "rows": rs})
    return out


def _qty(r):
    q = r.get("qty_file_ship")
    return q if q is not None else r.get("qty_coupang")


def _mmdd(d):
    return f"{int(d[5:7]):02d}{int(d[8:10]):02d}" if d else "日期不明"


def throw_filename(g, groups=None):
    """照 Chloe 的範例：指定到貨日20260918,酷澎_PG.xlsx。同一天有兩個以上的倉時檔名會撞，才在後面加 (倉)。"""
    same_day = [x for x in (groups or []) if x["date"] == g["date"]]
    tail = f"({g['warehouse'] or '倉不明'})" if len(same_day) > 1 else ""
    return f"指定到貨日{(g['date'] or '').replace('-', '') or '日期不明'},酷澎_PG{tail}.xlsx"


def mapping_filename(g):
    return f"永豐料號對照表_{_mmdd(g['date'])}交貨({g['warehouse'] or '倉不明'}).xlsx"


def throw_phone(g, settings, phones):
    return phones.get((g["warehouse"] or "").upper()) or settings["phone"]


def throw_workbook(g, settings, phones=None):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    phone = throw_phone(g, settings, phones or {})
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "PO"
    ws.append(["訂單編號", "收件人", "收件人手機", "收件地址", "料號", "數量", "單價", "金額小計", "出貨備註", "運費"])
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="F4B183"); c.font = Font(name="新細明體", size=12)
    for r in g["rows"]:
        q = _qty(r) or 0; p = r.get("unit_price")
        sub = round(q * p, 2) if p is not None else None
        if sub is not None and float(sub).is_integer():
            sub = int(sub)
        if p is not None and float(p).is_integer():
            p = int(p)
        ws.append([r["po_number"], settings["recipient"], phone, r.get("address") or "", r.get("yf_sku") or "", q, p, sub, None, None])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = Font(name="新細明體", size=12)
        row[0].number_format = "@"; row[4].number_format = "@"; row[7].number_format = "0_);[Red]\\(0\\)"
    for i, w in enumerate([18.5, 19.5, 22, 35.1, 21.1, 9.2, 9.3, 13.7, 47.1, 5.6], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def mapping_workbook(g):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "工作表1"
    ws.append(["PO No.", "SKU No.", "條碼", "料號", "品名", "數量", "單位", "轉換率", "箱數", "允收效期"])
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="DDEBF7"); c.font = Font(name="微軟正黑體", size=12, bold=True)
    for r in g["rows"]:
        q = _qty(r) or 0; box = r.get("box_size")
        cases = math.ceil(q / box) if box else None
        ws.append([r["po_number"], r["sku_id"], r.get("barcode") or "", r.get("yf_sku") or "", r.get("product_name") or "",
                   q, r.get("unit") or "", box, cases, None])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = Font(name="微軟正黑體", size=12)
        for i in (0, 1, 2, 3):
            row[i].number_format = "@"
        row[8].number_format = "0.0_);[Red]\\(0.0\\)"
    for i, w in enumerate([23.3, 23.3, 25.7, 23.9, 73.8, 12.1, 9.5, 12.1, 9.0, 21.4], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _summary(groups, rows, warnings, settings, phones):
    skipped = collections.Counter(norm_text(r.get("line")) or "空白" for r in rows if not _is_pg(r.get("line")))
    out = []
    for g in groups:
        rs = g["rows"]
        issues = []
        if any(_qty(r) in (None, 0) for r in rs):
            issues.append(f"{sum(1 for r in rs if _qty(r) in (None, 0))} 列出貨數量是 0 或空的")
        if any(not r.get("box_size") for r in rs):
            issues.append(f"{sum(1 for r in rs if not r.get('box_size'))} 列沒有箱入數，箱數算不出來")
        if any(r.get("unit_price") is None for r in rs):
            issues.append(f"{sum(1 for r in rs if r.get('unit_price') is None)} 列沒有單價，金額小計空白")
        if any(not r.get("yf_sku") for r in rs):
            issues.append(f"{sum(1 for r in rs if not r.get('yf_sku'))} 列沒有永豐料號")
        out.append({"date": g["date"], "warehouse": g["warehouse"], "rows": len(rs), "pos": len({r["po_number"] for r in rs}),
                    "qty": sum(_qty(r) or 0 for r in rs), "cases": sum(math.ceil((_qty(r) or 0) / r["box_size"]) for r in rs if r.get("box_size")),
                    "amount": round(sum((_qty(r) or 0) * (r.get("unit_price") or 0) for r in rs), 2),
                    "throw_file": throw_filename(g, groups), "mapping_file": mapping_filename(g), "issues": issues,
                    "phone": throw_phone(g, settings, phones), "phone_from": "倉別手機表" if phones.get((g["warehouse"] or "").upper()) else "固定字",
                    "recipient": settings["recipient"],
                    "preview": [{"po": r["po_number"], "sku": r["sku_id"], "yf_sku": r.get("yf_sku"), "name": (r.get("product_name") or "")[:40],
                                 "qty": _qty(r), "unit": r.get("unit"), "box": r.get("box_size"),
                                 "cases": math.ceil((_qty(r) or 0) / r["box_size"]) if r.get("box_size") else None,
                                 "price": r.get("unit_price"), "address": r.get("address")} for r in rs[:60]]})
    return {"groups": out, "skipped_lines": dict(skipped), "warnings": warnings, "total_rows": len(rows),
            "pg_rows": sum(len(g["rows"]) for g in groups)}


def _parse_files(files):
    rows, warnings, names = [], [], []
    for f in files:
        try:
            rs, ws = importer.parse_workbook(f.stream, f.filename)
        except importer.ImportError_ as exc:
            raise ValueError(f"{f.filename}：{exc}") from exc
        rows += rs; warnings += [f"{f.filename}：{w}" for w in ws]; names.append(f.filename)
    return rows, warnings, names


@zhuyun_bp.route("/zhuyun")
def zhuyun_page():
    from flask import current_app
    return render_template("zhuyun.html", logged_in_user=session.get("user", ""), settings=_load_settings(), phones=load_phones(),
                           build_version=current_app.config.get("BUILD_VERSION", ""))


@zhuyun_bp.route("/api/zhuyun/settings", methods=["GET", "PUT"])
def api_zhuyun_settings():
    if request.method == "GET":
        return jsonify(_load_settings())
    payload = request.get_json(silent=True) or {}
    new = _load_settings()
    for k in DEFAULTS:
        if k in payload:
            new[k] = str(payload.get(k) or "").strip()[:100]
    if not new["recipient"]:
        return jsonify({"error": "收件人不能空白。"}), 400
    _save_settings(new)
    return jsonify({"ok": True, **new})


@zhuyun_bp.route("/api/zhuyun/phones")
def api_zhuyun_phones():
    return jsonify({"phones": load_phones(), "default": _load_settings()["phone"]})


@zhuyun_bp.route("/api/zhuyun/phones", methods=["PUT"])
def api_zhuyun_phone_set():
    """填一個倉的手機（空白＝刪掉那個倉）。倉別大寫英數。"""
    payload = request.get_json(silent=True) or {}
    code = norm_text(payload.get("code")).upper().replace(" ", "")
    if not WH_CODE_RE.match(code):
        return jsonify({"error": f"倉別要像 TAO5、TXRC17 這樣的英數（2～12 碼），你填的是「{code or '空白'}」。"}), 400
    ph = pg_phone(payload.get("phone"))
    phones = load_phones(); old = phones.get(code, "")
    if ph:
        phones[code] = ph
    else:
        phones.pop(code, None)
    save_phones(phones)
    if old != ph:
        _log_phone(code, old, ph)
    return jsonify({"ok": True, "code": code, "phone": ph, "phones": phones})


@zhuyun_bp.route("/api/zhuyun/phones/import", methods=["POST"])
def api_zhuyun_phones_import():
    f = request.files.get("file")
    if f is None or not f.filename:
        return jsonify({"error": "沒有收到檔案。"}), 400
    try:
        new = parse_phone_sheet(f.stream)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    if not new:
        return jsonify({"error": "檔案裡沒有任何倉別＋手機。"}), 400
    phones = load_phones()
    added = sum(1 for k in new if k not in phones); updated = sum(1 for k, v in new.items() if k in phones and phones[k] != v)
    phones.update(new); save_phones(phones)
    _log_phone("", "", f"上傳 {f.filename}：{len(new)} 個倉，新增 {added}、更新 {updated}")
    return jsonify({"ok": True, "rows": len(new), "added": added, "updated": updated, "same": len(new) - added - updated, "phones": phones})


def _log_phone(code, old, new):
    try:
        conn = db.get_conn()
        try:
            _mst_log(conn, LINE, "", "", "", "zhuyun_phone", f"寶僑倉別手機 {code}".strip(), old, new, _mst_operator(), "manual")
            conn.commit()
        finally:
            conn.close()
    except Exception:  # noqa: BLE001
        pass


@zhuyun_bp.route("/api/zhuyun/parse", methods=["POST"])
def api_zhuyun_parse():
    files = [f for f in request.files.getlist("file") if f and f.filename]
    if not files:
        return jsonify({"error": "沒有收到檔案。"}), 400
    try:
        rows, warnings, names = _parse_files(files)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    groups = group_rows(rows)
    if not groups:
        lines = collections.Counter(norm_text(r.get("line")) or "空白" for r in rows)
        return jsonify({"error": "這份檔裡沒有線別是寶僑的列（" + "、".join(f"{k} {v} 列" for k, v in lines.items()) + "），竹運拋檔只做寶僑。"}), 400
    out = _summary(groups, rows, warnings, _load_settings(), warehouse_phones()); out["files"] = names
    return jsonify(out)


@zhuyun_bp.route("/api/zhuyun/export", methods=["POST"])
def api_zhuyun_export():
    """同一份檔再傳一次、直接產 zip：每組兩個檔。只有一組就各自一個檔也放 zip（檔名一致，人不用猜）。"""
    files = [f for f in request.files.getlist("file") if f and f.filename]
    if not files:
        return jsonify({"error": "沒有收到檔案。"}), 400
    try:
        rows, _warnings, _names = _parse_files(files)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    groups = group_rows(rows)
    if not groups:
        return jsonify({"error": "這份檔裡沒有線別是寶僑的列。"}), 400
    only = norm_text(request.form.get("only"))      # "YYYY-MM-DD|倉"：只產這一組（檔名還是照全部算，同一天多倉才帶倉）
    all_groups = groups
    if only:
        groups = [g for g in groups if f"{g['date']}|{g['warehouse']}" == only]
        if not groups:
            return jsonify({"error": "找不到這一組到貨日＋倉。"}), 400
    settings = _load_settings(); phones = warehouse_phones()
    try:
        conn = db.get_conn()
        try:
            for g in groups:
                _mst_log(conn, LINE, "", "", "", "zhuyun_export", "匯出竹運拋檔", "", f"{g['date']} {g['warehouse']}：{len(g['rows'])} 列、{len({r['po_number'] for r in g['rows']})} 張 PO",
                         _mst_operator(), "manual", f"{throw_filename(g, all_groups)}；{mapping_filename(g)}")
            conn.commit()
        finally:
            conn.close()
    except Exception:  # noqa: BLE001
        pass
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for g in groups:
            zf.writestr(throw_filename(g, all_groups), throw_workbook(g, settings, phones))
            zf.writestr(mapping_filename(g), mapping_workbook(g))
    buf.seek(0)
    dates = sorted({g["date"] for g in groups})
    name = f"竹運拋檔_{_mmdd(dates[0])}" + ("" if len(dates) == 1 else f"-{_mmdd(dates[-1])}") + "交貨.zip"
    resp = send_file(buf, as_attachment=True, download_name=name, mimetype="application/zip")
    resp.headers["X-Zhuyun-Groups"] = str(len(groups))
    return resp


__all__ = ["zhuyun_bp", "group_rows", "throw_workbook", "mapping_workbook", "throw_filename", "mapping_filename", "throw_phone",
           "warehouse_phones", "load_phones", "save_phones", "pg_phone", "parse_phone_sheet", "DEFAULTS"]
